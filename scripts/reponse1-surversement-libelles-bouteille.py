#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REPONSE 1 - Sur-versement au verre : deux libelles de caisse qui sont des VENTES
DE BOUTEILLE et non des ventes au verre.

POURQUOI. Le libelle de la caisse est tronque a vingt caracteres. Deux libelles
ressemblent a un article au verre alors que la vente porte sur une bouteille
entiere de 75 cl :

  - « Arbois Chardonnay Le », reference 2019, vendu 24,90 puis 29,00 euros ;
  - « Beaujolais Moulin a », sous la SEULE reference 1831, vendu 29,00 puis
    38,00 euros. La reference 1832 du MEME libelle, vendue 5,80 puis 7,60 euros,
    est bien le verre de 15 cl : les deux references coexistent.

Les prix se lisent sur les cartes des vins versees au dossier
(public/documents/vins-boissons/carte_vins-boissons_2021-07-20.pdf et
carte_vins-boissons_2023-04-26_actuelle.pdf) :
  - Arbois Chardonnay : une seule colonne de prix, la colonne 75 cl, a 24,90
    puis 29,00 euros. Aucun prix au verre n'existe pour ce vin.
  - Moulin a Vent : trois colonnes, verre 15 cl a 5,80 puis 7,60 euros, pichet
    50 cl a 21,00 puis 26,00 euros, bouteille 75 cl a 29,00 puis 38,00 euros.
Le service ecrit lui-meme, page 36 de sa proposition de rectifications du
18/05/2026 : « Hormis pour des ventes de bouteilles specifiques, comme l'Arbois
Blanc Bethanie, le Chardonnay et les bouteilles bouchees de rose, la societe
propose comme article a la vente : - vin au verre (doses de 15cl) - vin au
pichet de 50 cl - vin au pichet de 75 cl. »

CE QUE LE SCRIPT ETABLIT, en lecture seule et de facon reproductible :
 1. le releve ligne a ligne des deux libelles dans les tickets de caisse
    (ANNEXE-C1 a C3), par reference, par exercice, avec les prix pratiques ;
 2. le decompte des VERRES DE VIN NOMME, avant et apres retrait des ventes de
    bouteille : 5 557,5 unites annoncees, 307,5 qui sont des bouteilles, donc
    5 250,0 verres reels ; et 4 681,5 verres dans le perimetre de la degustation
    (hors les deux vins achetes en BIB de 10 L) ;
 3. l'effet du reclassement sur l'assiette du sur-versement : les 307,5 unites
    passent de 15 cl a 75 cl, ce qui retire 46,13 L au verre, ajoute 230,63 L a
    la bouteille, augmente de 184,50 L le volume vendu aux doses de la carte et
    reduit d'autant l'ecart de bilan matiere.

Le scenario « avant » reproduit a l'identique les totaux publies dans
src/data/reponse1Calculs/sur-versement-fourchette.json : c'est le controle qui
garantit que le scenario « apres » ne differe que par la correction.

Sources, toutes en lecture seule :
  public/documents/caisse-enregistreuse/ANNEXE-C{1,2,3}_detail-tickets_*.xls
  src/data/calculsBoissons/itemsCaisse.json
  src/data/calculsBoissons/boissonsHorsCocktail.json
  src/data/calculsBoissons/consoTotaleParBoisson.json
  src/data/reponse1Calculs/sur-versement-fourchette.json (controle)
  scripts/reponse1-degustation-verre-pichet.py (mapping des vins nommes, repris
  a l'identique, aucun ajout ni retrait)

Sortie : public/documents/pieces-reponse-1/R1-sur-versement-libelles-bouteille.xlsx
"""
import os
import json
import importlib.util
import collections
import xlrd

ICI = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(ICI, ".."))
CAISSE = os.path.join(ROOT, "public/documents/caisse-enregistreuse")
CALC = os.path.join(ROOT, "src/data/calculsBoissons")
DATA = os.path.join(ROOT, "src/data/reponse1Calculs")
PIECES = os.path.join(ROOT, "public/documents/pieces-reponse-1")
OUT = os.path.join(PIECES, "R1-sur-versement-libelles-bouteille.xlsx")

EXOS = ["2022-2023", "2023-2024", "2024-2025"]

# Les deux libelles, et les references qui sont des ventes de bouteille.
# None = toutes les references du libelle.
VENTE_BOUTEILLE = {"Arbois Chardonnay Le": None, "Beaujolais Moulin à": {"1831"}}
TAILLE_BOUTEILLE_CL = 75.0

# Taux, regimes : repris a l'identique de reponse1-sur-versement-fourchette.py
VIN = {"vin_blanc", "vin_rouge", "vin", "vin_rose", "vin_de_liqueur"}
SPIRIT = {"aperitif", "liqueur", "eau_de_vie", "spiritueux", "digestif"}
INSCOPE = VIN | SPIRIT
TAUX_VIN, TAUX_SPIRIT, TAUX_COCKTAIL = 0.236, 0.20, 0.42
MAIN, PREMESURE, SCELLE = "verre", "pichet", "bouteille"


def regime(fmt):
    f = (fmt or "").lower()
    if "pichet" in f:
        return PREMESURE
    if "bouteille" in f:
        return SCELLE
    return MAIN


def charger_module_degustation():
    spec = importlib.util.spec_from_file_location(
        "deg_vp", os.path.join(ICI, "reponse1-degustation-verre-pichet.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# --------------------------------------------------------------------------
# 1. Releve des deux libelles dans les tickets
# --------------------------------------------------------------------------
def releve_tickets():
    """[(libelle, reference, exercice, quantite, prix pratiques, CA TTC)]"""
    agg = collections.defaultdict(lambda: {"q": 0.0, "ca": 0.0, "prix": set(), "n": 0})
    for i, e in enumerate(EXOS, 1):
        wb = xlrd.open_workbook(os.path.join(
            CAISSE, f"ANNEXE-C{i}_detail-tickets_{e}.xls"))
        sh = wb.sheet_by_index(0)
        for r in range(1, sh.nrows):
            lib = str(sh.cell_value(r, 10)).strip()
            if lib not in VENTE_BOUTEILLE:
                continue
            ref = str(sh.cell_value(r, 9)).strip()
            try:
                ref = str(int(float(ref)))
            except ValueError:
                pass
            d = agg[(lib, ref, e)]
            d["q"] += float(sh.cell_value(r, 11) or 0)
            d["ca"] += float(sh.cell_value(r, 16) or 0)
            pu = sh.cell_value(r, 13)
            if pu not in ("", None):
                d["prix"].add(round(float(pu), 2))
            d["n"] += 1
    out = []
    for (lib, ref, e), d in sorted(agg.items()):
        vb = VENTE_BOUTEILLE[lib]
        est_bouteille = (vb is None) or (ref in vb)
        out.append({"libelle": lib, "reference": ref, "exercice": e,
                    "lignes": d["n"], "quantite": round(d["q"], 2),
                    "ca_ttc": round(d["ca"], 2),
                    "prix": sorted(d["prix"]), "vente_bouteille": est_bouteille})
    return out


# --------------------------------------------------------------------------
# 2. Decompte des verres de vin nomme
# --------------------------------------------------------------------------
def verres_de_vin_nomme(mod, quantites_bouteille):
    items = {it["produit"]: it for it in json.load(
        open(os.path.join(CALC, "itemsCaisse.json"), encoding="utf-8"))["items"]}
    lignes = []
    for lib, vin in sorted(mod.TASTE.items()):
        it = items.get(lib)
        if not it or it.get("format_service") != "Verre":
            continue
        q = {e: float(it["quantite"].get(e, 0) or 0) for e in EXOS}
        bout = {e: quantites_bouteille.get((lib, e), 0.0) for e in EXOS}
        lignes.append({
            "libelle": lib, "vin": vin,
            "canonique": mod.VIN_SOURCE[vin],
            "conditionnement": mod.COND[vin],
            "q": q, "total": round(sum(q.values()), 2),
            "bouteille": bout, "total_bouteille": round(sum(bout.values()), 2),
            "verres": round(sum(q.values()) - sum(bout.values()), 2),
        })
    return lignes


# --------------------------------------------------------------------------
# 3. Effet du reclassement sur l'assiette du sur-versement
# --------------------------------------------------------------------------
def assiette(correction):
    """Reproduit le calcul de reponse1-sur-versement-fourchette.py. Si
    correction est vrai, les 307,5 unites passent de 15 cl a 75 cl : elles
    quittent le regime du verre pour celui de la bouteille, et le volume vendu
    aux doses de la carte augmente d'autant."""
    hors = json.load(open(os.path.join(CALC, "boissonsHorsCocktail.json"),
                          encoding="utf-8"))["boissons"]
    conso = json.load(open(os.path.join(CALC, "consoTotaleParBoisson.json"),
                           encoding="utf-8"))["boissons"]

    # Correction en volume, par vin canonique et par exercice.
    delta = collections.defaultdict(lambda: collections.defaultdict(float))
    if correction:
        for (lib, ref, e), q in QTE_BOUTEILLE_REF.items():
            canon = LIB_CANON[lib]
            delta[canon][e] += q * (TAILLE_BOUTEILLE_CL - 15.0) / 100.0

    vent = collections.defaultdict(
        lambda: {e: collections.defaultdict(float) for e in EXOS})
    for b in hors:
        if b.get("categorie") not in INSCOPE:
            continue
        r = regime(b.get("format_service"))
        for e in EXOS:
            pp = (b.get("par_periode", {}).get(e, {}) or {})
            vent[b["nom_canonique"]][e][r] += pp.get("volume_l", 0.0) or 0.0
    if correction:
        for canon, par_exo in delta.items():
            for e, d in par_exo.items():
                # le volume au verre des unites reclassees est retire, et le
                # volume reel de la bouteille entiere est porte au regime scelle
                q = d / ((TAILLE_BOUTEILLE_CL - 15.0) / 100.0)
                vent[canon][e][MAIN] -= q * 0.15
                vent[canon][e][SCELLE] += q * TAILLE_BOUTEILLE_CL / 100.0

    tot = collections.Counter()
    borne = {"dispo": 0.0, "vendu": 0.0, "base_main": 0.0,
             "retenu": 0.0, "plafonne": 0.0, "boissons": 0}
    for b in conso:
        if b.get("categorie") not in INSCOPE:
            continue
        nom = b["nom_canonique"]
        taille = b.get("taille_achat_cl") or 0
        inv = b.get("inventaire_fin_contenants_par_periode") or {}
        ach = b.get("achats_litres_par_periode") or {}
        stock_fin = {e: (inv.get(e) or 0) * taille / 100.0 for e in EXOS}
        b_dispo = b_vendu = b_main = b_retenu = b_plaf = 0.0
        for i, e in enumerate(EXOS):
            pp = (b.get("par_periode", {}) or {}).get(e, {}) or {}
            de = pp.get("detail_exact_l", {}) or {}
            seches = (de.get("boissons_seches", 0.0) or 0.0) + delta[nom][e]
            cocktails = de.get("ingredients_cocktails", 0.0) or 0.0
            cuisine = de.get("plats_desserts", 0.0) or 0.0
            menus = ((pp.get("menu_estime_l") or {}).get("moyen", 0.0)) or 0.0
            achats = ach.get(e) or 0.0
            stock_ini = stock_fin[EXOS[i - 1]] if i > 0 else stock_fin[e]
            dispo_service = achats + stock_ini - stock_fin[e] - cuisine
            vv = vent.get(nom, {}).get(e, {})
            v_verre = vv.get(MAIN, 0.0)
            vendu = seches + cocktails + menus
            base_main = v_verre + cocktails
            taux_sec = TAUX_VIN if b["categorie"] in VIN else TAUX_SPIRIT
            retenu = v_verre * taux_sec + cocktails * TAUX_COCKTAIL
            ecart = dispo_service - vendu
            est_borne = ecart > 0 and base_main > 0
            plafonne = min(retenu, ecart) if est_borne else retenu
            tot["verre"] += v_verre
            tot["pichet"] += vv.get(PREMESURE, 0.0)
            tot["bouteille"] += vv.get(SCELLE, 0.0)
            tot["cocktails"] += cocktails
            tot["base_main"] += base_main
            tot["retenu"] += retenu
            tot["vendu"] += vendu
            # Le perimetre borne se decide au niveau de la BOISSON, pas de
            # l'exercice : une boisson versee a la main y entre avec tous ses
            # exercices (cf. total_borne de reponse1-sur-versement-fourchette.py).
            b_dispo += dispo_service
            b_vendu += vendu
            b_main += base_main
            b_retenu += retenu
            b_plaf += plafonne
        if b_main > 0:
            borne["boissons"] += 1
            borne["dispo"] += b_dispo
            borne["vendu"] += b_vendu
            borne["base_main"] += b_main
            borne["retenu"] += b_retenu
            borne["plafonne"] += b_plaf
    borne["ecart"] = borne["dispo"] - borne["vendu"]
    return {k: round(v, 2) for k, v in tot.items()}, \
           {k: (round(v, 2) if isinstance(v, float) else v) for k, v in borne.items()}


def controle(av, bo):
    """Le scenario « avant » doit reproduire les totaux publies."""
    ref = json.load(open(os.path.join(DATA, "sur-versement-fourchette.json"),
                         encoding="utf-8"))
    rt, tb = ref["regimes_total"], ref["total_borne"]
    verifs = [("vin au verre", av["verre"], rt["vin_verre_l"] + rt["spiritueux_verre_l"]),
              ("pichet", av["pichet"], rt["vin_pichet_l"]),
              ("bouteille", av["bouteille"], rt["vin_bouteille_l"]),
              ("cocktails", av["cocktails"], rt["cocktails_l"]),
              ("base versee a la main", av["base_main"], rt["base_main_l"]),
              ("sur-versement retenu", av["retenu"], rt["sv_total_l"]),
              ("ecart borne", bo["ecart"], tb["ecart_l"]),
              ("boissons bornees", bo["boissons"], tb["boissons"])]
    ok = True
    for lib, a, b in verifs:
        if abs(a - b) > 0.05:
            ok = False
            print(f"  ECART {lib} : {a} vs {b} publie")
        else:
            print(f"  ok {lib} : {a}")
    assert ok, "le scenario avant ne reproduit pas les valeurs publiees"
    return verifs


def ecrire(rel, nom, av, bo_av, ap, bo_ap):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    blanc = Font(bold=True, color="FFFFFF")
    tete = PatternFill("solid", fgColor="0F766E")
    surligne = PatternFill("solid", fgColor="CCE7E2")
    trait = Side(style="thin", color="D8DEE4")
    bord = Border(left=trait, right=trait, top=trait, bottom=trait)
    droite = Alignment(horizontal="right")

    def titre(ws, t, st):
        ws.append([t]); ws["A1"].font = Font(bold=True, size=13)
        ws.append([st]); ws["A2"].font = Font(italic=True, size=9, color="64748B")
        ws.append([])

    def entetes(ws, cols):
        ws.append(cols)
        r = ws.max_row
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.fill, cell.border = blanc, tete, bord
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)

    def fin(ws, depuis=2, gras=False):
        r = ws.max_row
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).border = bord
            if c >= depuis:
                ws.cell(row=r, column=c).alignment = droite
            if gras:
                ws.cell(row=r, column=c).font = Font(bold=True)
                ws.cell(row=r, column=c).fill = surligne

    def larg(ws, ws_w):
        for i, w in enumerate(ws_w, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # Feuille 1 : le releve de caisse
    ws = wb.active
    ws.title = "Releve de caisse"
    titre(ws, "Les deux libelles de caisse qui sont des ventes de bouteille",
          "Releve ligne a ligne des tickets ANNEXE-C1 a C3. Le prix pratique "
          "identifie le contenant : la carte des vins ne donne qu'une colonne "
          "75 cl pour l'Arbois Chardonnay, et trois colonnes pour le Moulin a "
          "Vent (verre 15 cl, pichet 50 cl, bouteille 75 cl).")
    entetes(ws, ["Libelle de caisse", "Reference", "Exercice", "Lignes de ticket",
                 "Quantite", "Prix pratiques (euros TTC)", "CA TTC (euros)",
                 "Nature de la vente"])
    for r in rel:
        ws.append([r["libelle"], r["reference"], r["exercice"], r["lignes"],
                   r["quantite"], " ; ".join(f"{p:.2f}" for p in r["prix"]),
                   r["ca_ttc"],
                   "Bouteille 75 cl" if r["vente_bouteille"] else "Verre 15 cl"])
        fin(ws, 4)
    tb = sum(r["quantite"] for r in rel if r["vente_bouteille"])
    tv = sum(r["quantite"] for r in rel if not r["vente_bouteille"])
    ws.append([])
    ws.append(["Total des unites vendues en BOUTEILLE entiere", "", "", "",
               round(tb, 2), "", "", ""])
    fin(ws, 4, gras=True)
    ws.append(["Total des unites reellement vendues au VERRE", "", "", "",
               round(tv, 2), "", "", ""])
    fin(ws, 4)
    larg(ws, [24, 12, 14, 15, 12, 26, 14, 18])

    # Feuille 2 : les verres de vin nomme
    ws = wb.create_sheet("Verres de vin nomme")
    titre(ws, "Verres de vin nomme : 5 557,5 annonces, 307,5 bouteilles, 5 250 verres",
          "Un vin nomme est un vin identifie sur la note par son appellation, par "
          "opposition aux generiques. Mapping repris a l'identique de "
          "scripts/reponse1-degustation-verre-pichet.py.")
    cols = ["Libelle de caisse", "Vin", "Conditionnement d'achat"]
    cols += [f"Unites {e}" for e in EXOS]
    cols += ["Total unites", "dont ventes de bouteille", "Verres reels"]
    entetes(ws, cols)
    for l in nom:
        ws.append([l["libelle"], l["canonique"], l["conditionnement"]]
                  + [l["q"][e] for e in EXOS]
                  + [l["total"], l["total_bouteille"], l["verres"]])
        fin(ws, 4)
    t_tot = round(sum(l["total"] for l in nom), 2)
    t_bou = round(sum(l["total_bouteille"] for l in nom), 2)
    t_ver = round(sum(l["verres"] for l in nom), 2)
    ws.append(["TOTAL", "", ""] + ["" for _ in EXOS] + [t_tot, t_bou, t_ver])
    fin(ws, 4, gras=True)
    bib = round(sum(l["verres"] for l in nom if l["conditionnement"].startswith("BIB")), 2)
    ws.append([])
    ws.append(["dont vins achetes en BIB de 10 L, hors demande de degustation",
               "", ""] + ["" for _ in EXOS] + ["", "", bib])
    fin(ws, 4)
    ws.append(["Perimetre de la degustation : verres de vin nomme hors BIB",
               "", ""] + ["" for _ in EXOS] + ["", "", round(t_ver - bib, 2)])
    fin(ws, 4, gras=True)
    larg(ws, [24, 32, 22, 14, 14, 14, 13, 20, 13])

    # Feuille 3 : l'effet sur l'assiette
    ws = wb.create_sheet("Effet sur l'assiette")
    titre(ws, "Effet du reclassement sur l'assiette du sur-versement",
          "Les 307,5 unites passent de 15 cl a 75 cl. Le scenario « avant » "
          "reproduit a l'identique les totaux publies dans "
          "src/data/reponse1Calculs/sur-versement-fourchette.json.")
    entetes(ws, ["Grandeur (3 exercices)", "Avant correction",
                 "Apres correction", "Ecart"])
    def l3(lib, a, b):
        ws.append([lib, a, b, round(b - a, 2)]); fin(ws, 2)
    l3("Volume verse a la main au verre (L)", av["verre"], ap["verre"])
    l3("Volume au pichet 50 et 75 cl (L)", av["pichet"], ap["pichet"])
    l3("Volume a la bouteille et demi-bouteille (L)", av["bouteille"], ap["bouteille"])
    l3("Alcool des cocktails (L)", av["cocktails"], ap["cocktails"])
    l3("Base reellement versee a la main (L)", av["base_main"], ap["base_main"])
    l3("Volume vendu aux doses de la carte (L)", av["vendu"], ap["vendu"])
    ws.append([])
    l3("Sur-versement retenu (L)", av["retenu"], ap["retenu"])
    fin(ws, 2, gras=True)
    ws.append([])
    l3("Disponible au service, boissons bornees (L)", bo_av["dispo"], bo_ap["dispo"])
    l3("Vendu aux doses de la carte, boissons bornees (L)", bo_av["vendu"], bo_ap["vendu"])
    l3("Ecart de bilan matiere (L)", bo_av["ecart"], bo_ap["ecart"])
    l3("Nombre de boissons bornees", bo_av["boissons"], bo_ap["boissons"])
    ws.append([])
    pa = round(100 * av["retenu"] / bo_av["ecart"], 2)
    pb = round(100 * ap["retenu"] / bo_ap["ecart"], 2)
    l3("Part de l'ecart couverte par le sur-versement (%)", pa, pb)
    fin(ws, 2, gras=True)
    larg(ws, [46, 18, 18, 14])

    os.makedirs(PIECES, exist_ok=True)
    wb.save(OUT)


if __name__ == "__main__":
    mod = charger_module_degustation()
    rel = releve_tickets()

    QTE_BOUTEILLE_REF = {}
    for r in rel:
        if r["vente_bouteille"]:
            QTE_BOUTEILLE_REF[(r["libelle"], r["reference"], r["exercice"])] = r["quantite"]
    LIB_CANON = {lib: mod.VIN_SOURCE[mod.TASTE[lib]] for lib in VENTE_BOUTEILLE}
    globals()["QTE_BOUTEILLE_REF"] = QTE_BOUTEILLE_REF
    globals()["LIB_CANON"] = LIB_CANON

    par_lib_exo = collections.defaultdict(float)
    for (lib, ref, e), q in QTE_BOUTEILLE_REF.items():
        par_lib_exo[(lib, e)] += q
    nom = verres_de_vin_nomme(mod, par_lib_exo)

    print("Controle du scenario avant correction :")
    av, bo_av = assiette(False)
    controle(av, bo_av)
    ap, bo_ap = assiette(True)

    print("\nUnites vendues en bouteille entiere :",
          round(sum(QTE_BOUTEILLE_REF.values()), 2))
    print("Verres de vin nomme :", round(sum(l["total"] for l in nom), 2),
          "->", round(sum(l["verres"] for l in nom), 2))
    print("Perimetre degustation (hors BIB) :",
          round(sum(l["verres"] for l in nom
                    if not l["conditionnement"].startswith("BIB")), 2))
    print("\nAssiette avant :", av)
    print("Assiette apres :", ap)
    print("Borne avant :", bo_av)
    print("Borne apres :", bo_ap)
    ecrire(rel, nom, av, bo_av, ap, bo_ap)
    print("\nEcrit :", OUT)
