#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REPONSE 1 - Abattements 5 % + 5 % + 5 % : le meme 15 % oppose a chaque poste.

Le service (reponse du 04/09/2026) oppose UN SEUL abattement forfaitaire de
15 % (5 % offerts + 5 % pertes + 5 % consommation du personnel, applique en fin
de methode de reconstitution) a SIX postes de perte differents, l'un apres
l'autre, dans huit pages distinctes : freinte biere (p. 61 et 83), degustation
au verre (p. 64 et 86), offerts / pertes / consommation du personnel (p. 72, 73
et 89) et coefficient de revente (p. 93, deja p. 59).

Ce script produit la piece R1-abattements-double-emploi.xlsx :
  1. le releve page par page des postes auxquels le meme 15 % est oppose ;
  2. l'itemisation de nos postes de perte, en litres ET en euros, par exercice ;
  3. la comparaison forfait (service) / itemisation (defense), en volume et en
     euros, avec la decomposition exacte des 35 133,26 € et 105 399,78 €.

Sources (lecture seule, reproductible) :
  - src/data/reponse1Calculs/cascade-10622.json        SOURCE UNIQUE des volumes
    et des euros de l'itemisation. Produite par scripts/reponse1-cascade-valeurs.py.
    Ce script ne recalcule plus aucun poste : il met en forme. Les valeurs de
    juillet 2026 (sur-versement 720 L, cremant 347 L, freinte 129 L, itemisation
    1 505 L et 95 976 EUR) ne doivent plus apparaitre ici.
  - proposition de rectifications du 18/05/2026, p. 52 (recapitulation par
    exercice), reprise dans
    public/documents/rapports-des-finances-publiques/synthese/06-methode-reconstitution-2.md

Sortie :
  public/documents/pieces-reponse-1/R1-abattements-double-emploi.xlsx
"""
import json
import os

ICI = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(ICI, ".."))
DATA = os.path.join(ROOT, "src", "data")
PIECES = os.path.join(ROOT, "public", "documents", "pieces-reponse-1")
EXOS = ["2022-2023", "2023-2024", "2024-2025"]

# --- RECAPITULATION PAR EXERCICE DU SERVICE (proposition p. 52) --------------
# CA TTC. « avant » = C-A LIQUIDES avant remise, pertes et conso. personnel ;
# « cinq_pct » = chacun des trois abattements de 5 % tel que chiffre par le
# service lui-meme ; « coef » = rapport liquides/solides ; « total » = TOTAL
# GENERAL DU CA reconstitue. Aucun de ces chiffres n'est de nous.
RECAP = [
    {"ex": "2022-2023", "avant": 178341.42, "cinq_pct": 8917.07,
     "apres": 151590.21, "coef": 2.94, "total": 597265.42, "declare": 404030.87},
    {"ex": "2023-2024", "avant": 169014.34, "cinq_pct": 8450.72,
     "apres": 143662.19, "coef": 3.02, "total": 577522.00, "declare": 438658.43},
    # 455 249,20 EUR etait un artefact d'OCR : l'OCR brut porte HT 389 519,47 et
    # TVA 46 005,45, dont la somme fait 435 524,92 EUR. Controle decisif : le
    # courrier ecrit lui-meme « - 139 727,76 EUR » sur cet exercice, et
    # 575 252,68 - 435 524,92 = 139 727,76 exactement.
    {"ex": "2024-2025", "avant": 165065.33, "cinq_pct": 8253.27,
     "apres": 140305.53, "coef": 3.10, "total": 575252.68, "declare": 435524.92},
]

# --- RELEVE PAGE PAR PAGE (reponse du service du 04/09/2026) ----------------
# (page, poste auquel le 15 % est oppose, citation courte du courrier)
RELEVE = [
    (61, "Freinte technique de la biere pression",
     "« le service, compte tenu de la methode de reconstitution employee, applique, en fin de "
     "methode, 15 % supplementaires en taux de perte/offert/consommation du personnel. De ce "
     "fait, le service retient donc en fait, un taux total de 30 % sur la biere »"),
    (64, "Degustation offerte au verre (gouter du vin)",
     "« le service a deja retenu un taux offert/perte/consommation du personnel de 15 % (en fin "
     "de methode de reconstitution) » ; le meme 15 % est ensuite converti en litres de vin : "
     "« 45 433 centilitres x 15 % », « 1 169 litres x 15 % = 175,35 litres »"),
    (72, "Offerts (aperitifs et cafes)",
     "« avec la proportion de 5 % pour l'offert, le service a minore le montant du chiffre "
     "d'affaires TTC reconstitue, sur le premier exercice a hauteur de 35 133,26 € [...] "
     "multiplie par 3 (105 399,78 €), ce qui est plus de 13 fois ce que propose Maitre THIVEND »"),
    (72, "Pertes (casse, freinte, gaspillage)",
     "« avec la proportion de 5 % pour les pertes, le service minore le montant du chiffre "
     "d'affaires TTC reconstitue a hauteur de 35 133,26 € [...] ce montant equivaut a 8,70 % du "
     "chiffre d'affaires TTC declare »"),
    (73, "Consommation du personnel",
     "« avec la proportion de 5 % pour la consommation du personnel, service a minore le montant "
     "du chiffre d'affaires TTC reconstitue, sur le premier exercice a hauteur de 35 133,26 € "
     "[...] multiplie par 3 (105 399,78 €), ce qui est pres de 3 fois ce que propose Maitre "
     "THIVEND »"),
    (83, "Freinte technique de la biere pression (reprise)",
     "« en fin de reconstitution, le service applique bien un taux de 15 % pour pertes, offerts "
     "et consommations du personnel. Mais, a l'interieur de la methode [...] un taux specifique "
     "sur perte de biere de 15 %. Par consequent, en cumulant les deux taux [...] le service "
     "avait deja applique un taux de pertes, offerts, consommation du personnel de 30 % »"),
    (86, "Degustation offerte au verre (reprise)",
     "« il retenait comme fraction, non prise en compte dans le calcul de la reconstitution du "
     "chiffre d'affaires [...] 45 433 centilitres x 15 % » ; « Le service a egalement retranche "
     "une fraction de 15 % supplementaire sur les bouteilles de vin de 75 cl bouchees »"),
    (89, "Offerts, pertes et consommation du personnel (reprise)",
     "« en retenant un taux de 5 % d'offerts [...] 35 133,26 € [...] multiplie par 3 "
     "(105 399,78 €) » ; « avec une proportion de 5 % pour les pertes [...] 35 133,26 € » ; "
     "« Avec la proportion de 5 % pour la consommation du personnel [...] 35 133,26 € »"),
    (93, "Coefficient de revente (deja soutenu p. 59)",
     "« grace aux taux de perte/offerts/consommation du personnel de 15 %, retenu par le "
     "service, le coefficient moyen de l'exploitation de la societe devrait etre de : "
     "3,85 x (1-0,15) = 3,2725 »"),
]


# --- Lecture des donnees internes -------------------------------------------
def charger():
    """Lit la SOURCE UNIQUE de la cascade et de l'itemisation."""
    chemin = os.path.join(DATA, "reponse1Calculs", "cascade-10622.json")
    if not os.path.exists(chemin):
        raise SystemExit("Lancer d'abord : python3 scripts/reponse1-cascade-valeurs.py")
    with open(chemin, encoding="utf-8") as f:
        return json.load(f)


# --- Construction des postes itemises, par exercice --------------------------
# Aucun calcul ici : les litres, les euros et les prix de revente viennent tous
# de src/data/reponse1Calculs/cascade-10622.json, produit par
# scripts/reponse1-cascade-valeurs.py.
def postes(casc):
    P = []
    for it in casc["itemisation"]:
        P.append({
            "poste": it["poste"],
            "litres": [it["litres_par_exercice"][e] for e in EXOS],
            "eur": [it["euros_par_exercice"][e] for e in EXOS],
            "base": it["base"],
        })
    return P, casc["prix_de_revente_eur_par_l"]


# --- Ecriture de la piece ----------------------------------------------------
def ecrire(casc, P, prix):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    blanc = Font(bold=True, color="FFFFFF")
    gras = Font(bold=True)
    tete = PatternFill("solid", fgColor="0F766E")
    surligne = PatternFill("solid", fgColor="CCE7E2")
    alerte = PatternFill("solid", fgColor="FDE2E1")
    trait = Side(style="thin", color="D8DEE4")
    bord = Border(left=trait, right=trait, top=trait, bottom=trait)

    def titre(ws, txt, sous):
        ws.append([txt])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([sous])
        ws["A2"].font = Font(italic=True, size=9, color="64748B")
        ws.append([])

    def entetes(ws, cols):
        ws.append(cols)
        r = ws.max_row
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.fill, cell.border = blanc, tete, bord
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def largeurs(ws, l):
        for i, w in enumerate(l, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ---- Feuille 1 : le releve page par page -------------------------------
    ws = wb.active
    ws.title = "1. Le meme 15 pct"
    titre(ws, "Le meme abattement de 15 % oppose successivement a six postes de perte",
          "Reponse de la DDFiP du Jura du 04/09/2026. Le 15 % (5 % offerts + 5 % pertes + 5 % "
          "consommation du personnel) est applique UNE SEULE FOIS, en fin de methode de "
          "reconstitution (proposition du 18/05/2026, p. 51 et 52). Citations reproduites "
          "sans accents typographiques.")
    entetes(ws, ["Page", "Poste auquel le 15 % est oppose", "Citation du courrier"])
    for pg, poste, cit in RELEVE:
        ws.append([f"p. {pg}", poste, cit])
        for c in range(1, 4):
            ws.cell(row=ws.max_row, column=c).border = bord
            ws.cell(row=ws.max_row, column=c).alignment = Alignment(vertical="top", wrap_text=True)
    ws.append([])
    ws.append(["Total", "6 postes distincts sur 8 pages", "Un forfait unique ne peut pas couvrir "
               "la somme des postes qui lui sont successivement opposes."])
    for c in range(1, 4):
        ws.cell(row=ws.max_row, column=c).font = gras
        ws.cell(row=ws.max_row, column=c).fill = alerte
    largeurs(ws, [8, 46, 96])

    # ---- Feuille 2 : itemisation par exercice ------------------------------
    ws = wb.create_sheet("2. Itemisation par exercice")
    titre(ws, "Nos postes de consommation sans vente, en litres et en euros, par exercice",
          "Perimetre : hors cuisine et hors alcool des menus. Litres : mesures par exercice "
          "quand la source le permet, sinon repartis selon la cle indiquee. Euros : prix de "
          "revente au litre issus de nos propres donnees, sauf consommation du chef et "
          "offerts, repris du fichier 56 remis au service. Volumes et euros lus dans "
          "src/data/reponse1Calculs/cascade-10622.json, source unique de la cascade. "
          "Les offerts deja enregistres a 0,00 EUR (25,75 L) sont deja compris dans les "
          "postes de vente de la cascade : ils ne sont pas reclames une seconde fois ici.")
    entetes(ws, ["Poste"] + [f"{e} (L)" for e in EXOS] + ["Total (L)"]
            + [f"{e} (€)" for e in EXOS] + ["Total (€)", "Base de calcul"])
    tl = [0.0, 0.0, 0.0]
    te = [0.0, 0.0, 0.0]
    for p in P:
        ws.append([p["poste"]] + [round(v, 1) for v in p["litres"]] + [round(sum(p["litres"]), 1)]
                  + [round(v) for v in p["eur"]] + [round(sum(p["eur"]))] + [p["base"]])
        for i in range(3):
            tl[i] += p["litres"][i]
            te[i] += p["eur"][i]
        for c in range(1, 11):
            ws.cell(row=ws.max_row, column=c).border = bord
    ws.append(["TOTAL itemise (hors cuisine et menus)"] + [round(v, 1) for v in tl]
              + [round(sum(tl), 1)] + [round(v) for v in te] + [round(sum(te))], )
    for c in range(1, 10):
        ws.cell(row=ws.max_row, column=c).font = gras
        ws.cell(row=ws.max_row, column=c).fill = surligne
    largeurs(ws, [42, 13, 13, 13, 13, 13, 13, 13, 14, 70])

    # ---- Feuille 3 : forfait contre itemisation ----------------------------
    achat_l = casc["achats_l"]
    forfait_l = casc["forfait_15pct"]["litres"]
    forfait_l_r = int(round(forfait_l))
    item_l = sum(tl)
    ASS = {a["cle"]: a for a in casc["assiettes"]}
    exploit_l = ASS["C"]["litres"]
    cuisine_l = casc["cuisine"]["plafonne_l"]
    residu_l = casc["residu_arrondi_l"]

    ws = wb.create_sheet("3. Forfait vs itemisation")
    titre(ws, "Ce que le forfait de 15 % represente reellement, et ce qu'il devrait couvrir",
          "Colonne de gauche : les chiffres du service (proposition p. 52, reponse p. 72, 73 "
          "et 89). Colonne de droite : nos postes itemises. Les deux grandeurs comparees par "
          "le service (un abattement en volume d'un cote, une perte de chiffre d'affaires en "
          "euros de l'autre) ne sont pas homogenes.")

    entetes(ws, ["A. Decomposition des montants avances par le service (CA TTC)", "Exercice 1",
                 "Exercice 2", "Exercice 3", "Total 3 exercices"])
    lignes_a = [
        ("CA liquides avant abattements (proposition p. 52)", [r["avant"] for r in RECAP]),
        ("Un abattement de 5 % sur les liquides", [r["cinq_pct"] for r in RECAP]),
        ("Les trois abattements (15 %) sur les liquides", [3 * r["cinq_pct"] for r in RECAP]),
        ("Effet d'un abattement de 5 % sur le CA total reconstitue (x 1 + coefficient)",
         [r["cinq_pct"] * (1 + r["coef"]) for r in RECAP]),
        ("Effet des trois abattements (15 %) sur le CA total reconstitue",
         [3 * r["cinq_pct"] * (1 + r["coef"]) for r in RECAP]),
        ("dont part boissons", [3 * r["cinq_pct"] for r in RECAP]),
        ("dont part cuisine, obtenue par extrapolation au coefficient",
         [3 * r["cinq_pct"] * r["coef"] for r in RECAP]),
        ("CA total reconstitue (proposition p. 52)", [r["total"] for r in RECAP]),
        ("CA declare (proposition p. 52)", [r["declare"] for r in RECAP]),
    ]
    for lbl, vals in lignes_a:
        ws.append([lbl] + [round(v, 2) for v in vals] + [round(sum(vals), 2)])
        for c in range(1, 6):
            ws.cell(row=ws.max_row, column=c).border = bord
    ws.append([])
    for lbl, val in [
        ("Verification : 8 917,07 € x (1 + 2,94) = le « 35 133,26 € » de la p. 72",
         RECAP[0]["cinq_pct"] * (1 + RECAP[0]["coef"])),
        ("Verification : 35 133,26 € x 3 = le « 105 399,78 € » de la p. 72, qui est aussi la "
         "TOTALITE du 15 % du seul exercice 1", 3 * RECAP[0]["cinq_pct"] * (1 + RECAP[0]["coef"])),
        ("Verification : CA reconstitue exercice 1 / 0,85 - CA reconstitue exercice 1",
         RECAP[0]["total"] / 0.85 - RECAP[0]["total"]),
        ("Un abattement de 5 % cumule sur les 3 exercices (et non x 3 fois l'exercice 1)",
         sum(r["cinq_pct"] * (1 + r["coef"]) for r in RECAP)),
        ("Ecart introduit par la multiplication par 3 de l'exercice 1",
         3 * RECAP[0]["cinq_pct"] * (1 + RECAP[0]["coef"])
         - sum(r["cinq_pct"] * (1 + r["coef"]) for r in RECAP)),
        ("35 133,26 € rapportes au CA declare de l'exercice 1 (le « 8,70 % » de la p. 72, en %)",
         100 * RECAP[0]["cinq_pct"] * (1 + RECAP[0]["coef"]) / RECAP[0]["declare"]),
    ]:
        ws.append([lbl, round(val, 2)])
        ws.cell(row=ws.max_row, column=1).font = gras
    ws.append([])

    entetes(ws, ["B. Comparaison en volume d'alcool (3 exercices)", "Litres", "% des achats", ""
                 , ""])
    for lbl, v in [
        ("Alcool achete sur les 3 exercices", achat_l),
        ("Ce que represente le forfait de 15 % du service", forfait_l_r),
        ("A. Nos postes itemises, hors cuisine et menus", ASS["A"]["litres_arrondi"]),
        ("Alcool passe en cuisine et dans les menus, plafonne aux achats factures",
         round(cuisine_l)),
        ("B. Nos postes itemises, cuisine et alcool des menus compris", ASS["B"]["litres_arrondi"]),
        ("Residu de perte pure (solde du bilan matiere)", residu_l),
        ("C. Perte d'exploitation totale, hors cuisine et menus", ASS["C"]["litres_arrondi"]),
        ("D. Alcool achete et non vendu, cuisine et menus compris", ASS["D"]["litres_arrondi"]),
    ]:
        ws.append([lbl, round(v), round(100 * v / achat_l, 1)])
        for c in range(1, 4):
            ws.cell(row=ws.max_row, column=c).border = bord
    ws.append(["Part du forfait de 15 % que chaque assiette represente, en volume",
               "A : %.1f %%" % ASS["A"]["pct_du_forfait_en_volume"],
               "B : %.1f %%" % ASS["B"]["pct_du_forfait_en_volume"],
               "C : %.1f %%" % ASS["C"]["pct_du_forfait_en_volume"],
               "D : %.1f %%" % ASS["D"]["pct_du_forfait_en_volume"]])
    for c in range(1, 6):
        ws.cell(row=ws.max_row, column=c).font = gras
        ws.cell(row=ws.max_row, column=c).fill = alerte
    ws.append(["Les seuls postes itemises hors cuisine et menus ne suffisent plus a absorber "
               "le forfait : ils en representent %.1f %%. C'est l'assiette B, cuisine et alcool "
               "des menus compris, qui le depasse la premiere."
               % ASS["A"]["pct_du_forfait_en_volume"]])
    ws.append([])

    entetes(ws, ["C. Comparaison en euros de chiffre d'affaires boissons (3 exercices)",
                 "Euros", "", "", ""])
    part_boissons = sum(3 * r["cinq_pct"] for r in RECAP)
    for lbl, v in [
        ("Ce que le forfait de 15 % retire au CA boissons reconstitue", part_boissons),
        ("A. Nos postes itemises, hors cuisine et menus, valorises au prix de revente",
         ASS["A"]["euros"]),
        ("    Ecart de l'assiette A au forfait", ASS["A"]["ecart_euros_au_forfait"]),
        ("B. Cuisine et alcool des menus compris", ASS["B"]["euros"]),
        ("    Ecart de l'assiette B au forfait", ASS["B"]["ecart_euros_au_forfait"]),
        ("C. Perte d'exploitation totale, hors cuisine et menus", ASS["C"]["euros"]),
        ("    Ecart de l'assiette C au forfait", ASS["C"]["ecart_euros_au_forfait"]),
        ("D. Alcool achete et non vendu, cuisine et menus compris", ASS["D"]["euros"]),
        ("    Ecart de l'assiette D au forfait", ASS["D"]["ecart_euros_au_forfait"]),
    ]:
        ws.append([lbl, round(v, 2)])
        for c in range(1, 3):
            ws.cell(row=ws.max_row, column=c).border = bord
    ws.append([])
    ws.append(["Le residu de perte pure des assiettes C et D est valorise a %.2f EUR/L, "
               "moyenne des prix de revente ponderee par les litres achetes : c'est une "
               "moyenne, non une mesure."
               % casc["prix_de_revente_eur_par_l"]["residu_moyenne_des_achats"]])
    largeurs(ws, [78, 18, 18, 18, 18])

    os.makedirs(PIECES, exist_ok=True)
    out = os.path.join(PIECES, "R1-abattements-double-emploi.xlsx")
    wb.save(out)
    return out, {
        "forfait_l": forfait_l, "item_l": item_l, "exploit_l": exploit_l,
        "achat_l": achat_l, "cuisine_l": cuisine_l, "assiettes": ASS,
        "item_eur": sum(te), "part_boissons": part_boissons,
        "eur_par_exo": te, "l_par_exo": tl, "prix": prix,
    }


def main():
    casc = charger()
    P, prix = postes(casc)
    out, r = ecrire(casc, P, prix)
    print("REPONSE 1 - Abattements : le meme 15 % oppose a six postes")
    print("-" * 66)
    for pg, poste, _ in RELEVE:
        print(f"  p. {pg:>3} : {poste}")
    print("-" * 66)
    print(f"  Prix de revente retenus : vins {prix['vins']:.2f} €/L, spiritueux "
          f"{prix['spiritueux']:.2f} €/L, cremant {prix['cremant']:.2f} €/L, "
          f"biere {prix['biere']:.2f} €/L")
    print(f"  Litres par exercice     : {[round(v) for v in r['l_par_exo']]}  "
          f"(total {round(r['item_l'])} L)")
    print(f"  Euros par exercice      : {[round(v) for v in r['eur_par_exo']]}  "
          f"(total {round(r['item_eur'])} €)")
    print(f"  Forfait 15 % en volume  : {round(r['forfait_l'])} L")
    print(f"  Forfait 15 % en euros (part boissons) : {r['part_boissons']:.2f} €")
    for cle in ("A", "B", "C", "D"):
        a = r["assiettes"][cle]
        print(f"  {cle}. {a['libelle'][:52]:<52} {a['litres']:>9.1f} L "
              f"({a['pct_du_forfait_en_volume']:>5.1f} % du forfait)  "
              f"{a['euros']:>12.2f} € ({a['pct_du_forfait_en_euros']:>5.1f} %)")
    print(f"  Piece : {out}")


if __name__ == "__main__":
    main()
