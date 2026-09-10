#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
reponse1-conso-personnel.py
Piece R1-consommation-personnel.xlsx : consommation du personnel et du chef,
poste par poste et PAR EXERCICE, en litres et en euros, avec methode et sources.

Sources (lecture seule, aucun chiffre saisi a la main) :
  - src/data/incertitudeDisparu/base_disparu_ajuste.json : achats, stock final,
    consommations tracees en caisse, prix d'achat et de revente au litre.
  - src/data/incertitudeDisparu/06_conso_staff.py : methode de bornage (Picon
    borne par l'achat ; Macvin borne par un taux de 2 a 3 verres de 6 cl/jour).
  - src/data/calculsBoissons/cocktailsComposition.json : dose de Picon dans un
    Picon biere (4 cl pour 25 cl, 8 cl pour 50 cl).
  - Jours d'ouverture verifies en caisse (1 ticket Z par jour, annexe H).
  - Reponse DDFiP du 04/09/2026, p. 70 et 71 : volumes et montants du service.
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(RACINE, "src", "data")
SORTIE = os.path.join(RACINE, "public", "documents", "pieces-reponse-1",
                      "R1-consommation-personnel.xlsx")

base = json.load(open(os.path.join(SRC, "incertitudeDisparu", "base_disparu_ajuste.json"), encoding="utf-8"))
B = {x["nom"]: x for x in base["boissons"]}

JOURS = {"2022-2023": 222, "2023-2024": 221, "2024-2025": 219}
EXOS = list(JOURS)
J = sum(JOURS.values())

# ---------------------------------------------------------------- postes
picon = B["Picon"]
picon_trace = sum(picon["conso"][k] for k in ("seches_l", "cocktails_l", "plats_l", "menu_moyen_l"))
picon_chef = round(picon["achats_l"] - picon["stock_final_l"] - picon_trace, 2)   # 43,70 L

MACVIN_VERRE_CL, MACVIN_VERRES_J = 6.0, 2.5
macvin_chef_j = MACVIN_VERRES_J * MACVIN_VERRE_CL / 100.0                          # 0,15 L/j
macvin_chef = round(macvin_chef_j * J, 2)
macvin_bas = round(2.0 * MACVIN_VERRE_CL / 100.0 * J, 2)
macvin_haut = round(3.0 * MACVIN_VERRE_CL / 100.0 * J, 2)

COCA_CANETTES_J, COCA_CL = 6, 33.0
coca_j = COCA_CANETTES_J * COCA_CL / 100.0                                        # 1,98 L/j
coca_staff = round(coca_j * J, 2)

def par_jour(total):
    return {e: total * JOURS[e] / J for e in EXOS}

POSTES = [
    {"poste": "Picon, consommation du chef", "cat": "Alcool",
     "l": par_jour(picon_chef),
     "pa": picon["prix_achat_l"], "pv": picon["prix_revente_l"],
     "methode": "Achats 72,00 L, stock final 0, seul usage trace en caisse = Picon biere (28,30 L a 4 cl de Picon par verre). Le residu est borne par l'achat.",
     "source": "base_disparu_ajuste.json (Picon) ; cocktailsComposition.json (dose 4 cl / 8 cl)"},
    {"poste": "Macvin, consommation du chef", "cat": "Alcool",
     "l": par_jour(macvin_chef),
     "pa": B["Macvin"]["prix_achat_l"], "pv": B["Macvin"]["prix_revente_l"],
     "methode": "Taux d'ouverture borne : 2 a 3 verres de 6 cl par jour d'ouverture, valeur centrale 2,5 verres. Fourchette 3 exercices : %s L a %s L." % (
         f"{macvin_bas:.2f}".replace(".", ","), f"{macvin_haut:.2f}".replace(".", ",")),
     "source": "06_conso_staff.py ; base_disparu_ajuste.json (Macvin)"},
    {"poste": "Coca, consommation du personnel de salle", "cat": "Soft",
     "l": par_jour(coca_staff),
     "pa": B["Coca"]["prix_achat_l"], "pv": B["Coca"]["prix_revente_l"],
     "methode": "6 canettes de 33 cl par jour d'ouverture (4 pour le serveur, 2 pour la gerante), soit 1,98 L par jour.",
     "source": "Declaration cliente du 09/06/2026 ; base_disparu_ajuste.json (Coca)"},
]

# ---------------------------------------------------- chiffres du service (p. 70-71)
VIN_DISPO, AUTRES_DISPO = 1676.83, 1229.23
BIERE_DISPO = 754.80
TAUX = 0.05
MINORATION_EUR, PART_CA_DECLARE = 35133.26, 0.0870
forfait_vin, forfait_autres = VIN_DISPO * TAUX, AUTRES_DISPO * TAUX
forfait_total = forfait_vin + forfait_autres                       # 145,30 L (1 exercice)
assiette_ca = MINORATION_EUR / TAUX                                # CA TTC reconstitue implique
ca_declare = MINORATION_EUR / PART_CA_DECLARE                      # CA TTC declare implique
prix_implicite = MINORATION_EUR / forfait_total                    # euros par litre
part_vin_biere = (VIN_DISPO + BIERE_DISPO) / (VIN_DISPO + AUTRES_DISPO)

# ---------------------------------------------------------------- mise en forme
GRAS = Font(bold=True)
BLANC_GRAS = Font(bold=True, color="FFFFFF")
ENTETE = PatternFill("solid", fgColor="1F4E5F")
TOTAL = PatternFill("solid", fgColor="E8EEF1")
BORD = Border(*[Side(style="thin", color="BFCBD1")] * 4)
HAUT = Alignment(vertical="top", wrap_text=True)

def entete(ws, cols, largeurs):
    ws.append(cols)
    for i, (c, w) in enumerate(zip(cols, largeurs), start=1):
        cel = ws.cell(row=ws.max_row, column=i)
        cel.font, cel.fill, cel.alignment, cel.border = BLANC_GRAS, ENTETE, HAUT, BORD
        ws.column_dimensions[cel.column_letter].width = w

def ligne(ws, vals, gras=False, fond=False, fmt=None):
    ws.append(vals)
    r = ws.max_row
    for i in range(1, len(vals) + 1):
        c = ws.cell(row=r, column=i)
        c.border, c.alignment = BORD, HAUT
        if gras: c.font = GRAS
        if fond: c.fill = TOTAL
        if fmt and i in fmt: c.number_format = fmt[i]

wb = Workbook()

# ---- 1. Synthese par exercice
ws = wb.active
ws.title = "1. Synthese par exercice"
ws["A1"] = "Consommation du personnel et du chef, par exercice (SARL LA DEMI LUNE)"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = ("Volumes non vendus, mesures poste par poste. Colonnes en euros : cout d'achat "
            "(prix facture au litre) et CA equivalent (prix de revente au litre releve en caisse).")
ws.append([])
F = {2: "0,00", 3: "0,00", 4: "0,00", 5: "0,00", 6: "#,##0.00 €", 7: "#,##0.00 €"}
entete(ws, ["Poste", "2022-2023 (L)", "2023-2024 (L)", "2024-2025 (L)", "Total 3 exercices (L)",
            "Cout d'achat (3 exercices)", "CA equivalent (3 exercices)"],
       [42, 15, 15, 15, 20, 22, 22])
tot_l = {e: 0.0 for e in EXOS}
tot_alc = {e: 0.0 for e in EXOS}
for p in POSTES:
    t = sum(p["l"].values())
    ca = t * p["pv"] if p["pv"] else None
    ligne(ws, [p["poste"]] + [round(p["l"][e], 2) for e in EXOS] + [round(t, 2),
          round(t * p["pa"], 2), round(ca, 2) if ca else "non determinable"], fmt=F)
    for e in EXOS:
        tot_l[e] += p["l"][e]
        if p["cat"] == "Alcool": tot_alc[e] += p["l"][e]
ligne(ws, ["Sous-total ALCOOL"] + [round(tot_alc[e], 2) for e in EXOS] + [round(sum(tot_alc.values()), 2), "", ""],
      gras=True, fond=True, fmt=F)
ligne(ws, ["TOTAL toutes boissons"] + [round(tot_l[e], 2) for e in EXOS] + [round(sum(tot_l.values()), 2), "", ""],
      gras=True, fond=True, fmt=F)
ws.append([])
ligne(ws, ["Jours d'ouverture (caisse, 1 ticket Z par jour)"] + [JOURS[e] for e in EXOS] + [J, "", ""])
ws.append([])
ws.append(["Le Picon n'a aucun prix de revente propre en caisse : il n'a jamais ete vendu seul "
           "sur les trois exercices, son seul usage enregistre est le Picon biere."])

# ---- 2. Comparaison avec le forfait du service
ws2 = wb.create_sheet("2. Forfait 5 % du service")
ws2["A1"] = "Le forfait de 5 %, tel que le service le calcule p. 70 et 71, compare a notre decompte"
ws2["A1"].font = Font(bold=True, size=13)
ws2.append([])
entete(ws2, ["Grandeur", "Perimetre", "1 exercice (L)", "3 exercices (L)", "Source"],
       [46, 30, 16, 16, 40])
F2 = {3: "0,00", 4: "0,00"}
ligne(ws2, ["5 % du vin disponible", "Alcool seul, exercice 1", round(forfait_vin, 2),
            round(forfait_vin * 3, 2), "Reponse 04/09/2026, p. 70 et 71 (1 676,83 L)"], fmt=F2)
ligne(ws2, ["5 % des autres alcools listes", "Alcool seul, exercice 1", round(forfait_autres, 2),
            round(forfait_autres * 3, 2), "Reponse 04/09/2026, p. 71 (1 229,23 L, liste partielle)"], fmt=F2)
ligne(ws2, ["Forfait 5 % du service, cumul", "Alcool seul", round(forfait_total, 2),
            round(forfait_total * 3, 2), "Reponse 04/09/2026, p. 71 (83,84 + 61,46)"],
      gras=True, fond=True, fmt=F2)
ligne(ws2, ["Notre decompte : alcool (Picon + Macvin)", "Alcool seul",
            round(sum(tot_alc.values()) / 3, 2), round(sum(tot_alc.values()), 2),
            "Onglet 1 ; memoire du 10/07/2026 (143 L sur 3 exercices)"], fmt=F2)
ligne(ws2, ["Notre decompte : toutes boissons", "Alcool + softs",
            round(sum(tot_l.values()) / 3, 2), round(sum(tot_l.values()), 2),
            "Onglet 1 (le forfait du service vise 'toutes les boissons', p. 70)"], fmt=F2)
ws2.append([])
ws2.append(["Le service compare son forfait d'UN exercice (145,30 L) au volume revendique sur TROIS "
            "exercices (143 L). Rebases sur la meme periode, les deux grandeurs ne coincident pas."])

# ---- 3. Le changement d'assiette
ws3 = wb.create_sheet("3. Changement d'assiette")
ws3["A1"] = "Le meme taux de 5 %, applique a deux assiettes differentes (p. 70 et 71)"
ws3["A1"].font = Font(bold=True, size=13)
ws3.append([])
entete(ws3, ["Grandeur", "Calcul", "Resultat", "Source"], [50, 34, 20, 40])
def L3(a, b, c, d, f="#,##0.00"):
    ligne(ws3, [a, b, c, d], fmt={3: f})
L3("Volume d'alcool disponible, exercice 1", "1 676,83 + 1 229,23", round(VIN_DISPO + AUTRES_DISPO, 2), "p. 70 et 71", "0,00 \\L")
L3("Forfait 5 % exprime en volume", "5 % de 2 906,06 L", round(forfait_total, 2), "p. 71", "0,00 \\L")
L3("Forfait 5 % exprime en euros", "chiffre du courrier", round(MINORATION_EUR, 2), "p. 71", "#,##0.00 €")
L3("Prix au litre implique par les deux chiffres", "35 133,26 / 145,30", round(prix_implicite, 2),
   "calcul", "#,##0.00 €")
L3("CA TTC reconstitue implique par la minoration", "35 133,26 / 5 %", round(assiette_ca, 2),
   "p. 71", "#,##0.00 €")
L3("CA TTC declare implique", "35 133,26 / 8,70 %", round(ca_declare, 2), "p. 71", "#,##0.00 €")
L3("Ecart entre CA reconstitue et CA declare", "702 665,20 / 403 830,57 - 1",
   round(100 * (assiette_ca / ca_declare - 1), 1), "calcul", "0,0 %%")
ws3.append([])
ws3.append(["Aucune boisson de la carte n'approche 241,80 euros le litre : le montant de 35 133,26 euros "
            "n'est donc pas la contrepartie des 145,30 litres, il est 5 % d'un chiffre d'affaires global "
            "dont la part cuisine est obtenue par un coefficient."])

# ---- 4. Methode et sources
ws4 = wb.create_sheet("4. Methode et sources")
ws4["A1"] = "Methode de calcul et sources, poste par poste"
ws4["A1"].font = Font(bold=True, size=13)
ws4.append([])
entete(ws4, ["Poste", "Methode de calcul", "Source"], [36, 74, 52])
for p in POSTES:
    ligne(ws4, [p["poste"], p["methode"], p["source"]])
ligne(ws4, ["Repartition par exercice",
            "Chaque volume triennal est reparti au prorata des jours d'ouverture de l'exercice "
            "(222, 221 et 219 jours, soit 662 jours).",
            "Caisse : un ticket Z par jour d'ouverture (annexe H)"])
ligne(ws4, ["Valorisation en euros",
            "Cout d'achat = litres x prix d'achat au litre issu des factures. CA equivalent = "
            "litres x prix de revente au litre releve en caisse.",
            "base_disparu_ajuste.json (prix_achat_l, prix_revente_l)"])
ws4.append([])
ws4.append(["Piece etablie par le script scripts/reponse1-conso-personnel.py (reproductible)."])

os.makedirs(os.path.dirname(SORTIE), exist_ok=True)
wb.save(SORTIE)

print(f"Jours d'ouverture : {JOURS} (total {J})")
print(f"Picon chef      : {picon_chef} L sur 3 exercices")
print(f"Macvin chef     : {macvin_chef} L (fourchette {macvin_bas} - {macvin_haut} L)")
print(f"Coca personnel  : {coca_staff} L")
print(f"ALCOOL 3 exos   : {sum(tot_alc.values()):.2f} L | TOUTES BOISSONS : {sum(tot_l.values()):.2f} L")
for e in EXOS:
    print(f"  {e} : alcool {tot_alc[e]:.2f} L, total {tot_l[e]:.2f} L")
print(f"Forfait service : {forfait_total:.2f} L / exercice, soit {forfait_total*3:.2f} L sur 3 exercices")
print(f"Prix implicite  : {prix_implicite:.2f} EUR/L")
print(f"CA reconstitue  : {assiette_ca:.2f} EUR | CA declare : {ca_declare:.2f} EUR "
      f"| ecart +{100*(assiette_ca/ca_declare-1):.1f} %")
print(f"Part vin + biere dans la base du service : {100*part_vin_biere:.1f} %")
print(f"Ecrit : {SORTIE}")
