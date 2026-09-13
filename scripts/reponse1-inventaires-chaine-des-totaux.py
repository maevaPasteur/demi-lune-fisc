#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REPONSE 1 - Piece "R1-inventaires-chaine-des-totaux.xlsx"
=========================================================
Objet : etablir, piece par piece et au centime, la chaine qui relie

    feuilles de comptage  ->  page recapitulative datee  ->  compte de stock

pour les trois clotures 31/03/2023, 31/03/2024 et 31/03/2025.

Le service (reponse du 04/09/2026, p. 33) n'objecte plus le contenu des
inventaires mais leur DATE : "les corrections etant intervenues
post-proposition de rectifications". La chaine ci-dessous se lit entierement
sur les PDF d'origine, remis au verificateur avant sa proposition.

SOURCES
-------
1. Les trois PDF scannes, joints au memoire :
   public/documents/inventaires/Inventaire_Demi_Lune_{2023,2024,2025}-03-31.pdf
   Les totaux manuscrits ou imprimes releves page par page sont recopies
   ci-dessous dans les constantes RELEVE_*, chacune portant la page du PDF ou
   elle se verifie. Aucun de ces nombres n'est calcule : ils sont LUS.
2. La transcription publiee des trois inventaires :
   public/documents/inventaires/inventaire_{2023,2024,2025}-03-31.csv
   (boissons seules ; les articles d'epicerie et d'entretien portes sur les
   memes feuilles en sont exclus).
3. La variation de stock boissons retenue par le service. Elle figure en pied
   des annexes n. 7-1, 7-2 et 7-3 de la proposition de rectification sous le
   libelle "312000 STOCK BOISSONS DEMI LUNE", et au tableau consolide des
   pages 34 et 35 de la proposition a la ligne "Variation stocks Boissons".
   La reponse du 04/09/2026 la rappelle p. 74 sous le numero "310200".

Tout le reste est calcule par ce script et controle par des assertions.

Sortie : public/documents/pieces-reponse-1/R1-inventaires-chaine-des-totaux.xlsx
"""

import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INV_DIR = os.path.join(ROOT, "public/documents/inventaires")
OUT = os.path.join(
    ROOT, "public/documents/pieces-reponse-1/R1-inventaires-chaine-des-totaux.xlsx"
)

# --------------------------------------------------------------------------- #
# 1. Ce qui est LU sur les PDF d'origine
# --------------------------------------------------------------------------- #

# --- Cloture du 31/03/2023 : etat dactylographie, deux tableaux totalises ---- #
RELEVE_2023 = {
    "feuilles": [
        # (page du PDF, intitule porte sur la feuille, famille, total imprime)
        (10, "Tableau « SANS ALCOOL », colonnes Produit sans alcool / Quantite / "
             "Prix unitaire HT / Total, 41 lignes, TOTAL imprime en pied",
         "sans_alcool", 727.70),
        (11, "Tableau « ALCOOL », colonnes Produit alcoolise / Quantite / "
             "Prix unitaire HT / Total, 45 lignes, sans total de page",
         "alcool", None),
        (12, "Suite du tableau « ALCOOL », 18 lignes, TOTAL imprime en pied "
             "de tableau", "alcool", 3050.29),
    ],
    "recapitulatif": {
        "page": 1,
        "titre": "INVENTAIRE H.T. DEMI LUNE 31/03/2023",
        "lignes": [
            ("BOISSONS SANS ALCOOL", 727.70),
            ("ALCOOLS ET VINS", 3050.29),
            ("ALIMENTAIRE (Frais-Surgeles-Epicerie)", 1734.15),
            ("CREMERIE", 158.85),
            ("VIANDE-CHARCUTERIE", 255.63),
            ("PRODUITS D'ENTRETIEN ET DIVERS (Gel fondue, ...)", 1103.99),
        ],
        "total_imprime": 7030.61,
    },
    # Lignes portees sur la page 10 mais hors du perimetre "boissons" du CSV.
    "hors_boissons": [
        ("sucre morceaux", 0, 11.06, 0.00),
        ("petites galettes st michel", 2, 35.06, 70.12),
        ("petites madeleines st michel", 2, 18.38, 36.76),
        ("mix crackers", 5, 4.63, 23.15),
    ],
}

# --- Cloture du 31/03/2024 : comptage porte sur la fiche client FCBS -------- #
# En-tete imprime sur chaque page : "FRANCHE-COMTE BOISSONS SERVICE / FICHE
# CLIENT / Le 7/03/24 a 13:55:33 / Semaine 10 Jour 67 / Page : n".
# Le 67e jour de 2024 (annee bissextile : 31 + 29 + 7) est le 7 mars 2024.
RELEVE_2024 = {
    "edition_fiche": "Le 7/03/24 a 13:55:33 - Semaine 10 Jour 67 (en-tete imprime, p. 2 a 7)",
    "feuilles": [
        # (page du PDF, sous-total "sans alcool", sous-total "avec alcool")
        (2, 198.79, 210.70),
        (3, 585.13, None),
        (4, 34.75, 561.32),
        (5, None, 900.99),
        (6, None, 1704.49),
        (7, 51.92, 559.57),
    ],
    "report_manuscrit": {
        "page": 7,
        "sans_alcool": 870.59,
        "alcool": 3937.07,
        "total": 4807.66,
    },
    "recapitulatif": {
        "page": 1,
        "titre": "INVENTAIRE H.T. AU 31 MARS 2024",
        "lignes": [
            ("Boissons sans alcool", 870.59),
            ("Vins & Alcools", 3937.07),
            ("Produits d'entretien et Divers (Serviettes, Essuie-mains...)", 1000.30),
            ("Alimentation (Frais, surgeles, viandes, cremerie...)", 4065.17),
        ],
        # La page imprime "Total : 4 807.60 EUR" la ou ses deux lignes font
        # 4 807,66 EUR. Le report manuscrit de la page 7 porte 4 807,66, et
        # c'est ce montant que reprend la comptabilite.
        "total_boissons_imprime": 4807.60,
        "total_general_imprime": 9873.07,
    },
}

# --- Cloture du 31/03/2025 : cinq feuilles manuscrites numerotees 1/5 a 5/5 -- #
RELEVE_2025 = {
    "feuilles": [
        # (page du PDF, numero porte sur la feuille, famille, total manuscrit)
        (8, "numero non lu", "alcool", 557.99),
        (9, "2/5", "alcool", 1524.04),
        (10, "3/5", "alcool", 1686.48),
        (11, "4/5", "sans_alcool", 473.41),
        (12, "5/5", "sans_alcool", 287.33),
    ],
    "recapitulatif": {
        "page": 1,
        "titre": "INVENTAIRE HT 31/03/2025",
        "lignes": [
            ("BOISSONS sans alcool", 765.74),
            ("BOISSONS vins et alcools", 3768.51),
            ("ALIMENTATION (frais, surgeles, epicerie...)", 3654.92),
            ("DIVERS NON COMESTIBLES (Produits d'entretien, serviettes...)", 1585.00),
        ],
        "total_boissons_imprime": 4534.25,
        "total_general_imprime": 9774.17,
    },
}

# --- Variation de stock boissons retenue par le service ---------------------- #
# Annexes 7-1 (p. 176), 7-2 (p. 190) et 7-3 (Proposition_4, p. 204) : ligne
# "312000 STOCK BOISSONS DEMI LUNE". Tableau consolide, proposition p. 34-35.
VARIATION_SERVICE = {
    "2022-2023": -800.83,
    "2023-2024": -1029.67,
    "2024-2025": -273.41,
}

# --------------------------------------------------------------------------- #
# 2. Lecture de la transcription publiee
# --------------------------------------------------------------------------- #


def lire_csv(date):
    chemin = os.path.join(INV_DIR, f"inventaire_{date}.csv")
    with open(chemin, encoding="utf-8") as f:
        return [
            {
                "produit": r["produit"].strip(),
                "categorie": r["categorie"].strip(),
                "quantite": (r["quantite"] or "").strip(),
                "pu": (r["prix_unitaire_ht"] or "").strip(),
                "valeur": float((r["valeur_ht"] or "0").replace(",", ".")),
                "fiabilite": (r.get("fiabilite") or "").strip(),
                "page": (r.get("page") or "").strip(),
            }
            for r in csv.DictReader(f, delimiter=";")
        ]


CSV = {d: lire_csv(d) for d in ("2023-03-31", "2024-03-31", "2025-03-31")}


def somme(lignes, categorie=None, page=None):
    s = sum(
        l["valeur"]
        for l in lignes
        if (categorie is None or l["categorie"] == categorie)
        and (page is None or l["page"] == str(page))
    )
    return round(s, 2)


# --------------------------------------------------------------------------- #
# 3. Controles
# --------------------------------------------------------------------------- #
controles = []


def verifier(libelle, calcule, attendu, tolerance=0.005):
    ecart = round(calcule - attendu, 2)
    ok = abs(ecart) < tolerance
    controles.append((libelle, calcule, attendu, ecart, ok))
    return ok


# --- 2023 ------------------------------------------------------------------- #
c23_sa = somme(CSV["2023-03-31"], "boisson_sans_alcool")
c23_al = somme(CSV["2023-03-31"], "alcool")
hors_b_2023 = round(sum(x[3] for x in RELEVE_2023["hors_boissons"]), 2)

verifier(
    "31/03/2023 - transcription « sans alcool » augmentee des 4 lignes d'epicerie "
    "de la page 10, contre le total imprime en pied de cette page",
    round(c23_sa + hors_b_2023, 2),
    727.70,
)
verifier(
    "31/03/2023 - transcription « alcools », contre le total imprime en pied "
    "du tableau ALCOOL (page 12)",
    c23_al,
    3050.29,
)
recap23 = dict(RELEVE_2023["recapitulatif"]["lignes"])
verifier(
    "31/03/2023 - somme des six familles de la page recapitulative datee, "
    "contre le total general imprime sur cette page",
    round(sum(v for _, v in RELEVE_2023["recapitulatif"]["lignes"]), 2),
    RELEVE_2023["recapitulatif"]["total_imprime"],
)
stock_2023 = round(recap23["BOISSONS SANS ALCOOL"] + recap23["ALCOOLS ET VINS"], 2)

# --- 2024 ------------------------------------------------------------------- #
sa24 = round(sum(f[1] for f in RELEVE_2024["feuilles"] if f[1] is not None), 2)
al24 = round(sum(f[2] for f in RELEVE_2024["feuilles"] if f[2] is not None), 2)
verifier(
    "31/03/2024 - somme des sous-totaux « sans alcool » des six feuilles de "
    "comptage, contre le report manuscrit de la page 7",
    sa24,
    RELEVE_2024["report_manuscrit"]["sans_alcool"],
)
verifier(
    "31/03/2024 - somme des sous-totaux « avec alcool » des six feuilles de "
    "comptage, contre le report manuscrit de la page 7",
    al24,
    RELEVE_2024["report_manuscrit"]["alcool"],
)
verifier(
    "31/03/2024 - report manuscrit de la page 7, contre la somme de ses deux "
    "composantes",
    round(sa24 + al24, 2),
    RELEVE_2024["report_manuscrit"]["total"],
)
recap24 = dict(RELEVE_2024["recapitulatif"]["lignes"])
verifier(
    "31/03/2024 - page recapitulative datee, lignes « Boissons sans alcool » et "
    "« Vins & Alcools », contre le report manuscrit de la page 7",
    round(recap24["Boissons sans alcool"] + recap24["Vins & Alcools"], 2),
    RELEVE_2024["report_manuscrit"]["total"],
)
stock_2024 = round(recap24["Boissons sans alcool"] + recap24["Vins & Alcools"], 2)

# --- 2025 ------------------------------------------------------------------- #
al25 = round(sum(f[3] for f in RELEVE_2025["feuilles"] if f[2] == "alcool"), 2)
sa25 = round(sum(f[3] for f in RELEVE_2025["feuilles"] if f[2] == "sans_alcool"), 2)
recap25 = dict(RELEVE_2025["recapitulatif"]["lignes"])
verifier(
    "31/03/2025 - somme des trois feuilles de comptage d'alcools (pages 8, 9 et 10), "
    "contre la ligne « vins et alcools » de la page recapitulative datee",
    al25,
    recap25["BOISSONS vins et alcools"],
)
verifier(
    "31/03/2025 - somme des deux feuilles de comptage de boissons sans alcool "
    "(pages 11 et 12), contre la transcription publiee",
    sa25,
    somme(CSV["2025-03-31"], "boisson_sans_alcool"),
)
stock_2025 = round(
    recap25["BOISSONS sans alcool"] + recap25["BOISSONS vins et alcools"], 2
)
ecart_sa_2025 = round(recap25["BOISSONS sans alcool"] - sa25, 2)

# --- Rapprochement avec le compte 310200 ------------------------------------ #
# Convention du plan comptable : variation de stock = stock initial - stock final.
var_2324 = round(stock_2023 - stock_2024, 2)
var_2425 = round(stock_2024 - stock_2025, 2)
verifier(
    "Exercice 2023-2024 - stock initial moins stock final, d'apres les deux pages "
    "recapitulatives datees, contre la variation portee au compte 310200",
    var_2324,
    VARIATION_SERVICE["2023-2024"],
)
verifier(
    "Exercice 2024-2025 - valeur absolue de l'ecart entre les deux pages "
    "recapitulatives datees, contre la valeur absolue du compte 310200",
    abs(var_2425),
    abs(VARIATION_SERVICE["2024-2025"]),
)

# --------------------------------------------------------------------------- #
# 4. Ecriture du classeur
# --------------------------------------------------------------------------- #
ENTETE = PatternFill("solid", fgColor="1F3864")
GRIS = PatternFill("solid", fgColor="F2F2F2")
VERT = PatternFill("solid", fgColor="E2EFDA")
ORANGE = PatternFill("solid", fgColor="FCE4D6")
BLANC = Font(color="FFFFFF", bold=True)
BORD = Border(*[Side(style="thin", color="BFBFBF")] * 4)
EUR = "# ##0.00"


def entete(ws, labels):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row=1, column=i, value=lab)
        c.fill = ENTETE
        c.font = BLANC
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORD
    ws.freeze_panes = ws.cell(row=2, column=1)


def largeurs(ws, valeurs):
    for i, w in enumerate(valeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def ligne(ws, r, valeurs, cols_eur=(), gras=False, fill=None):
    for i, v in enumerate(valeurs, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.border = BORD
        c.alignment = Alignment(vertical="top", wrap_text=(i <= 2))
        if gras:
            c.font = Font(bold=True)
        if fill:
            c.fill = fill
        if i in cols_eur:
            c.number_format = EUR
            c.alignment = Alignment(horizontal="right")
    return r + 1


wb = Workbook()

# --- Notice ---------------------------------------------------------------- #
ws = wb.active
ws.title = "Notice"
notice = [
    ["R1 - Inventaires de stocks : la chaine des totaux, des feuilles de comptage "
     "au compte 310200"],
    [""],
    ["Objet",
     "Le service ne conteste plus le contenu des inventaires produits, mais leur "
     "date : « les corrections etant intervenues post-proposition de "
     "rectifications » (reponse du 04/09/2026, p. 33). Cette piece etablit que "
     "les totaux repris par la comptabilite, puis par le service lui-meme, se "
     "lisent entierement sur les etats d'origine que le verificateur detenait "
     "avant sa proposition."],
    ["Ce qui est lu",
     "Les totaux de page et les pages recapitulatives datees des trois PDF "
     "d'inventaire, joints au memoire. Chaque nombre lu porte, dans les "
     "feuilles qui suivent, la page du PDF ou il se verifie."],
    ["Ce qui est calcule",
     "Les sommes, les ecarts et le rapprochement avec le compte 310200. Le "
     "script controle chacun d'eux par une assertion ; la feuille « Controles » "
     "les recapitule."],
    ["Resultat",
     f"Sur {len(controles)} controles, {sum(1 for c in controles if c[4])} sont "
     f"nuls au centime."],
    ["Perimetre de la transcription",
     "Les fichiers inventaire_AAAA-MM-JJ.csv retiennent les seules boissons. Les "
     "feuilles de comptage portent aussi du cafe, des thes, du sucre, de la "
     "biscuiterie, des pailles et des produits d'entretien, qui entrent dans les "
     "totaux imprimes mais pas dans la transcription. L'ecart est chiffre ligne "
     "a ligne pour la cloture 2023."],
    ["Unites", "Euros hors taxes. Quantites en bouteilles, futs ou unites."],
    ["Script", "scripts/reponse1-inventaires-chaine-des-totaux.py (reproductible)."],
]
for r, l in enumerate(notice, start=1):
    for c, v in enumerate(l, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if c == 1:
            cell.font = Font(bold=True)
ws.cell(row=1, column=1).font = Font(bold=True, size=14)
largeurs(ws, [26, 110])

# --- Cloture 2023 ---------------------------------------------------------- #
ws = wb.create_sheet("31-03-2023")
entete(ws, ["Piece", "Ce qu'elle porte", "Famille", "Montant lu ou calcule",
            "Nature du montant"])
r = 2
for page, desc, fam, total in RELEVE_2023["feuilles"]:
    r = ligne(ws, r, [f"PDF 2023, page {page}", desc, fam,
                      total if total is not None else "",
                      "total imprime en pied de tableau" if total else
                      "pas de total de page, le tableau continue page 12"],
              cols_eur=(4,))
r = ligne(ws, r, ["", "", "", "", ""])
for lib, val in RELEVE_2023["recapitulatif"]["lignes"]:
    r = ligne(ws, r, [f"PDF 2023, page 1 (datee)",
                      f"« {RELEVE_2023['recapitulatif']['titre']} » : {lib}",
                      "", val, "montant imprime"], cols_eur=(4,))
r = ligne(ws, r, ["PDF 2023, page 1 (datee)", "TOTAL imprime", "",
                  RELEVE_2023["recapitulatif"]["total_imprime"], "montant imprime"],
          cols_eur=(4,), gras=True, fill=GRIS)
r = ligne(ws, r, ["", "", "", "", ""])
r = ligne(ws, r, ["Transcription CSV", "boissons sans alcool, 37 lignes",
                  "sans_alcool", c23_sa, "calcule"], cols_eur=(4,))
for lib, q, pu, val in RELEVE_2023["hors_boissons"]:
    r = ligne(ws, r, ["PDF 2023, page 10",
                      f"{lib} : {q} x {pu:.2f} EUR (porte sur la page des boissons, "
                      f"hors perimetre de la transcription)", "hors boissons", val,
                      "montant imprime"], cols_eur=(4,))
r = ligne(ws, r, ["Controle", "transcription « sans alcool » + les 4 lignes "
                  "d'epicerie de la page 10", "sans_alcool",
                  round(c23_sa + hors_b_2023, 2),
                  "egal au total imprime de la page 10 : 727,70 EUR"],
          cols_eur=(4,), gras=True, fill=VERT)
r = ligne(ws, r, ["Controle", "transcription « alcools », 63 lignes", "alcool",
                  c23_al, "egal au total imprime en pied de la page 12 : "
                  "3 050,29 EUR"], cols_eur=(4,), gras=True, fill=VERT)
r = ligne(ws, r, ["Stock boissons au 31/03/2023", "727,70 + 3 050,29", "",
                  stock_2023, "somme des deux lignes de la page datee"],
          cols_eur=(4,), gras=True, fill=GRIS)
largeurs(ws, [28, 62, 16, 18, 46])

# --- Cloture 2024 ---------------------------------------------------------- #
ws = wb.create_sheet("31-03-2024")
entete(ws, ["Piece", "Ce qu'elle porte", "Sous-total sans alcool",
            "Sous-total avec alcool", "Nature du montant"])
r = 2
r = ligne(ws, r, ["PDF 2024, pages 2 a 7",
                  "Fiche client Franche-Comte Boissons Service, en-tete imprime : "
                  + RELEVE_2024["edition_fiche"] +
                  ". La trame de produits du fournisseur sert de liste de comptage ; "
                  "les quantites et les valeurs sont portees au stylo.",
                  "", "", "en-tete imprime"], gras=True, fill=GRIS)
for page, sa, al in RELEVE_2024["feuilles"]:
    r = ligne(ws, r, [f"PDF 2024, page {page}",
                      "sous-totaux manuscrits en pied de page",
                      sa if sa is not None else "",
                      al if al is not None else "",
                      "montants manuscrits"], cols_eur=(3, 4))
r = ligne(ws, r, ["Somme des six feuilles", "calculee par ce script", sa24, al24,
                  "calcule"], cols_eur=(3, 4), gras=True, fill=VERT)
r = ligne(ws, r, [f"PDF 2024, page {RELEVE_2024['report_manuscrit']['page']}",
                  "report manuscrit : « Total Sans Alcool = ... / Avec Alcool = ... »",
                  RELEVE_2024["report_manuscrit"]["sans_alcool"],
                  RELEVE_2024["report_manuscrit"]["alcool"],
                  "montants manuscrits, identiques a la somme calculee"],
          cols_eur=(3, 4), gras=True, fill=VERT)
r = ligne(ws, r, ["", "", "", "", ""])
for lib, val in RELEVE_2024["recapitulatif"]["lignes"]:
    r = ligne(ws, r, ["PDF 2024, page 1 (datee)",
                      f"« {RELEVE_2024['recapitulatif']['titre']} » : {lib}",
                      val if "sans alcool" in lib else "",
                      val if "Vins" in lib else "",
                      "montant imprime"], cols_eur=(3, 4))
r = ligne(ws, r, ["PDF 2024, page 1 (datee)",
                  "« Total : 4 807.60 EUR » imprime, la ou les deux lignes "
                  "ci-dessus font 4 807,66 EUR. Le report manuscrit de la page 7 "
                  "porte 4 807,66, et c'est ce montant que reprend la comptabilite.",
                  "", RELEVE_2024["recapitulatif"]["total_boissons_imprime"],
                  "coquille d'impression signalee par la societe"],
          cols_eur=(4,), fill=ORANGE)
r = ligne(ws, r, ["Stock boissons au 31/03/2024", "870,59 + 3 937,07", "",
                  stock_2024, "somme des deux lignes de la page datee"],
          cols_eur=(4,), gras=True, fill=GRIS)
largeurs(ws, [28, 72, 18, 18, 44])

# --- Cloture 2025 ---------------------------------------------------------- #
ws = wb.create_sheet("31-03-2025")
entete(ws, ["Piece", "Feuille", "Famille", "Total lu", "Nature du montant"])
r = 2
for page, num, fam, total in RELEVE_2025["feuilles"]:
    r = ligne(ws, r, [f"PDF 2025, page {page}", num, fam, total,
                      "total manuscrit en pied de feuille"], cols_eur=(4,))
r = ligne(ws, r, ["Somme des feuilles des pages 8, 9 et 10", "", "alcool", al25,
                  "calcule"], cols_eur=(4,), gras=True, fill=VERT)
r = ligne(ws, r, ["Somme des feuilles des pages 11 et 12", "", "sans_alcool", sa25,
                  "calcule"], cols_eur=(4,), gras=True, fill=VERT)
r = ligne(ws, r, ["", "", "", "", ""])
for lib, val in RELEVE_2025["recapitulatif"]["lignes"]:
    r = ligne(ws, r, ["PDF 2025, page 1 (datee)",
                      f"« {RELEVE_2025['recapitulatif']['titre']} »", lib, val,
                      "montant imprime"], cols_eur=(4,))
r = ligne(ws, r, ["Controle", "feuilles d'alcools contre la ligne « vins et "
                  "alcools » de la page datee", "alcool", al25,
                  "egal au centime : 3 768,51 EUR"], cols_eur=(4,), gras=True,
          fill=VERT)
r = ligne(ws, r, ["Reserve", "feuilles de boissons sans alcool contre la ligne "
                  "« sans alcool » de la page datee", "sans_alcool", ecart_sa_2025,
                  "ecart de 5,00 EUR, signale par la societe"], cols_eur=(4,),
          fill=ORANGE)
r = ligne(ws, r, ["Stock boissons au 31/03/2025", "765,74 + 3 768,51", "",
                  stock_2025, "somme des deux lignes de la page datee"],
          cols_eur=(4,), gras=True, fill=GRIS)
largeurs(ws, [34, 12, 16, 18, 46])

# --- Rapprochement 310200 --------------------------------------------------- #
ws = wb.create_sheet("Compte 310200")
entete(ws, ["Exercice", "Stock initial (page datee)", "Stock final (page datee)",
            "Stock initial moins stock final", "Variation retenue par le service",
            "Ecart"])
r = 2
r = ligne(ws, r, ["2022-2023", "stock au 31/03/2022 non repris dans nos etats",
                  stock_2023, "non calculable", VARIATION_SERVICE["2022-2023"],
                  "non calculable"], cols_eur=(3, 5))
r = ligne(ws, r, ["2023-2024", stock_2023, stock_2024, var_2324,
                  VARIATION_SERVICE["2023-2024"],
                  round(var_2324 - VARIATION_SERVICE["2023-2024"], 2)],
          cols_eur=(2, 3, 4, 5, 6), gras=True, fill=VERT)
r = ligne(ws, r, ["2024-2025", stock_2024, stock_2025, var_2425,
                  VARIATION_SERVICE["2024-2025"],
                  "montants egaux en valeur absolue ; le signe porte au compte "
                  "est discute a la page L, point 9"],
          cols_eur=(2, 3, 4, 5))
largeurs(ws, [14, 30, 30, 26, 26, 40])

# --- Controles -------------------------------------------------------------- #
ws = wb.create_sheet("Controles")
entete(ws, ["Controle", "Calcule", "Attendu", "Ecart", "Resultat"])
r = 2
for lib, calc, att, ec, ok in controles:
    r = ligne(ws, r, [lib, calc, att, ec, "nul" if ok else "A VERIFIER"],
              cols_eur=(2, 3, 4), fill=VERT if ok else ORANGE)
largeurs(ws, [96, 14, 14, 12, 14])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)

# --------------------------------------------------------------------------- #
# 5. Compte rendu console
# --------------------------------------------------------------------------- #
print(f"Ecrit : {OUT}\n")
for lib, calc, att, ec, ok in controles:
    print(f"  [{'OK ' if ok else 'KO '}] {lib}\n        calcule {calc} / attendu {att} / ecart {ec:+.2f}")
print()
print(f"Stock boissons : 31/03/2023 {stock_2023} | 31/03/2024 {stock_2024} | "
      f"31/03/2025 {stock_2025}")
print(f"Variation 2023-2024 : {var_2324:+.2f} | service {VARIATION_SERVICE['2023-2024']:+.2f}")
print(f"Variation 2024-2025 : {var_2425:+.2f} | service {VARIATION_SERVICE['2024-2025']:+.2f}")
print(f"Ecart « sans alcool » au 31/03/2025 entre feuilles et recapitulatif : {ecart_sa_2025:+.2f} EUR")
print(f"Lignes d'epicerie de la page 10 de 2023, hors perimetre boissons : {hors_b_2023:.2f} EUR")
if not all(c[4] for c in controles):
    raise SystemExit("Au moins un controle n'est pas nul.")
