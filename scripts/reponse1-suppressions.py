#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
reponse1-suppressions.py

Pièce jointe de la réfutation « A - LES SUPPRESSIONS DE NOTES (lignes DEL) »
de la réponse DDFiP 39 du 04/09/2026 (p. 3 à 17).

Tout est recalculé depuis les annexes de caisse remises au titre de
l'article L. 47 A, II du LPF :
  ANNEXE-C{1,2,3} : ticket ligne (détail des tickets)
  ANNEXE-E{1,2,3} : tpvevenement (journal des événements, lignes « DEL »)
  ANNEXE-F{1,2,3} : reglement (notes encaissées, mode de règlement)
  ANNEXE-G{1,2,3} : journal TVA (porte le grand total perpétuel « tot_sum »)

Sorties :
  public/documents/pieces-reponse-1/R1-suppressions-rapprochements.xlsx

Onglets (huit, comme l'annonce la page « Les suppressions de notes ») :
  1. Synthese                  : les chiffres clés de la réfutation, dont le
                                 rapprochement des 21 302 DEL avec le fichier
                                 des tickets (journée + ticket Z, puis montant
                                 au centime) et le catalogue des prix pratiqués
  2. 09-09-2022 rafales        : chaque rafale de DEL face aux notes encaissées
  3. 09-09-2022 DEL detail     : les 67 DEL de la journée, article par article
  4. 09-09-2022 notes          : les 25 notes encaissées de la journée
  5. 21-03-2024 DEL detail     : les DEL de la journée
  6. 21-03-2024 partage        : les notes n°13 et n°14 et le partage des quantités
  7. Chainage grand total      : contrôle du totalisateur perpétuel NF525
  8. DEL vs prix carte         : les 21 302 DEL confrontées aux prix de la carte,
                                 avec la ligne TOTAL

Le rapprochement DEL / tickets est délibérément publié en SYNTHÈSE, par
exercice, et non ligne à ligne : un détail ligne à ligne reviendrait à remettre
au service le relevé nominatif des 5 311 lignes que ce seul critère n'apparie
pas.
"""
import os
import collections
import xlrd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ICI = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(ICI, ".."))
CAISSE = os.path.join(ROOT, "public/documents/caisse-enregistreuse")
SORTIE = os.path.join(ROOT, "public/documents/pieces-reponse-1")
EXOS = ["2022-2023", "2023-2024", "2024-2025"]

JOUR_A = "2022-09-09"   # contre-exemple « conversion au forfait » (p. 7 à 15)
JOUR_B = "2024-03-21"   # contre-exemple « paiement séparé » (p. 16 et 17)


def annexe(lettre, i, nom, ex):
    return os.path.join(CAISSE, f"ANNEXE-{lettre}{i}_{nom}_{ex}.xls")


def rows(path):
    sh = xlrd.open_workbook(path).sheet_by_index(0)
    for r in range(1, sh.nrows):
        yield [sh.cell_value(r, c) for c in range(sh.ncols)]


def c(x):
    """euros -> centimes entiers (comparaisons exactes)."""
    return int(round(float(x) * 100))


def e(x):
    return round(float(x), 2)


# ---------------------------------------------------------------- chargement
# Catalogue des prix unitaires pratiqués, et lignes de tickets, par exercice.
prix_ex = {}          # exercice -> {centimes: {libellés}}
lignes_ex = {}        # exercice -> {date: [ (no, heure, ttc, lib, qte, pu) ]}
notes_ex = {}         # exercice -> {date: {no: {heure, ttc, id, regl:[(mode,montant)]}}}
del_ex = {}           # exercice -> [ (date, heure, montant, id, z) ]
zindex_ex = {}        # exercice -> {(date, no_z): {montants des lignes en centimes}}
tickets_c = set()     # (date, n° de note) présentes dans l'annexe C
ttc_c = {}            # (date, n° de note) -> total TTC de la note (annexe C)

for i, ex in enumerate(EXOS, 1):
    prix = collections.defaultdict(set)
    lignes = collections.defaultdict(list)
    for r in rows(annexe("C", i, "detail-tickets", ex)):
        d = str(r[0])[:10]
        if not d.startswith("20"):
            continue
        try:
            q = float(r[11]); pu = float(r[13]); ttc = float(r[5])
        except (ValueError, TypeError):
            continue
        lib = str(r[10]).strip()
        if pu > 0:
            prix[c(pu)].add(lib)
        lignes[d].append((int(r[2]), str(r[1])[:5], e(ttc), lib, q, e(pu)))
    prix_ex[ex] = prix
    lignes_ex[ex] = lignes

    # Index (journée, n° de ticket Z) -> montants TTC des lignes de note.
    # Sert au rapprochement du fichier des événements avec celui des tickets.
    idx = collections.defaultdict(set)
    for r in rows(annexe("C", i, "detail-tickets", ex)):
        d = str(r[0])[:10]
        if not d.startswith("20"):
            continue
        try:
            idx[(d, int(float(r[4])))].add(c(r[16] or 0))
            tickets_c.add((d, int(float(r[2]))))
            ttc_c[(d, int(float(r[2])))] = e(r[5])
        except (ValueError, TypeError):
            continue
    zindex_ex[ex] = idx

    notes = collections.defaultdict(dict)
    for r in rows(annexe("F", i, "reglements", ex)):
        d = str(r[0])[:10]
        if not d.startswith("20"):
            continue
        no = int(r[2])
        n = notes[d].setdefault(no, {"heure": str(r[1])[:5], "ttc": e(r[5]),
                                     "id": int(r[8]), "regl": []})
        n["regl"].append((str(r[9]), e(r[10])))
    notes_ex[ex] = notes

    dels = []
    for r in rows(annexe("E", i, "tpvevenement", ex)):
        d = str(r[1])[:10]
        if not d.startswith("20") or str(r[3]) != "DEL":
            continue
        dels.append((d, str(r[2])[:5], e(r[8]), int(r[9]), int(r[6])))
    del_ex[ex] = dels

EX_A = EXOS[0]   # 09/09/2022 -> exercice 2022-2023
EX_B = EXOS[1]   # 21/03/2024 -> exercice 2023-2024 (clos au 31/03/2024)


# ------------------------------------------------- 1. DEL contre prix de la carte
TOTAL_DEL = sum(len(v) for v in del_ex.values())
classe = collections.Counter()
valeurs = collections.Counter()
detail_classe = []
for ex in EXOS:
    prix = prix_ex[ex]
    n_exact = n_mult = n_autre = 0
    for d, h, m, _id, _z in del_ex[ex]:
        v = c(m)
        valeurs[v] += 1
        if v in prix:
            n_exact += 1
        else:
            k = next((k for k in range(2, 25) if v % k == 0 and v // k in prix), None)
            if k:
                n_mult += 1
            else:
                n_autre += 1
    classe[ex] = (n_exact, n_mult, n_autre)
    detail_classe.append((ex, n_exact, n_mult, n_autre, n_exact + n_mult + n_autre))

TOT_EXACT = sum(v[0] for v in classe.values())
TOT_MULT = sum(v[1] for v in classe.values())
TOT_AUTRE = sum(v[2] for v in classe.values())
NB_VALEURS = len(valeurs)
TOT_CARTE = TOT_EXACT + TOT_MULT
PCT_CARTE = 100 * TOT_CARTE / TOTAL_DEL

# Catalogue des prix unitaires effectivement pratiqués : par exercice, puis
# nombre de prix distincts sur les trois exercices réunis.
PRIX_PAR_EX = [(ex, len(prix_ex[ex])) for ex in EXOS]
PRIX_UNION = len(set().union(*[set(prix_ex[ex]) for ex in EXOS]))

# Montant et médiane des suppressions (le service retient « 9 EUR » p. 3).
MONTANTS_DEL = sorted(m for ex in EXOS for _d, _h, m, _i, _z in del_ex[ex])
TOTAL_DEL_EUR = round(sum(MONTANTS_DEL), 2)
MEDIANE_DEL = (MONTANTS_DEL[len(MONTANTS_DEL) // 2] if len(MONTANTS_DEL) % 2
               else round((MONTANTS_DEL[len(MONTANTS_DEL) // 2 - 1]
                           + MONTANTS_DEL[len(MONTANTS_DEL) // 2]) / 2, 2))


# ------------------ 1 bis. Rapprochement des DEL avec le fichier des tickets
# Deux critères successifs, appliqués aux 21 302 lignes :
#   a) la ligne « DEL » se rattache-t-elle à une journée et à un numéro de
#      ticket Z présents dans le fichier des tickets (annexe C) ?
#   b) son montant se retrouve-t-il au centime sur une ligne de note du même
#      ticket Z ?
# Le second critère n'apparie pas tout : ce qui n'est pas apparié par lui est
# porté ici en nombre, exercice par exercice, sans relevé nominatif.
rappro = []
for ex in EXOS:
    idx = zindex_ex[ex]
    n_jz = n_m = 0
    for d, _h, m, _id, z in del_ex[ex]:
        k = (d, z)
        if k in idx:
            n_jz += 1
            if c(m) in idx[k]:
                n_m += 1
    n = len(del_ex[ex])
    rappro.append((ex, n, n_jz, round(100 * n_jz / n, 1), n_m,
                   round(100 * n_m / n, 1), n - n_m))

RAP_JZ = sum(r[2] for r in rappro)
RAP_M = sum(r[4] for r in rappro)
RAP_RESTE = TOTAL_DEL - RAP_M
PCT_RAP_JZ = 100 * RAP_JZ / TOTAL_DEL
PCT_RAP_M = 100 * RAP_M / TOTAL_DEL


# ---------------------------------- 1 ter. Canal d'encaissement des règlements
# (le rapport « X fois les espèces » cité par la page se lit ici)
regl_modes = collections.Counter()
for ex in EXOS:
    for d, ns in notes_ex[ex].items():
        for no, nt in ns.items():
            for md, mt in nt["regl"]:
                regl_modes[md] += mt
REGL_TOTAL = round(sum(regl_modes.values()), 2)
ESPECES = round(regl_modes.get("ESP", 0), 2)
PCT_ESPECES = 100 * ESPECES / REGL_TOTAL
PCT_BANCARISE = 100 - PCT_ESPECES
RAPPORT_DEL_ESP = TOTAL_DEL_EUR / ESPECES

# Totaux de notes distincts (pour comparaison : population « note » vs « article »)
totaux_notes = collections.Counter()
nb_notes = 0
for ex in EXOS:
    for d, ns in notes_ex[ex].items():
        for no, n in ns.items():
            totaux_notes[c(n["ttc"])] += 1
            nb_notes += 1


# ------------------------------------------- 2. Chaînage du grand total perpétuel
# Le journal TVA porte « tot_sum », totalisateur perpétuel non remis à zéro.
# S'il progresse exactement du TTC de chaque ticket, aucune recette n'a disparu.
chainage = []
for i, ex in enumerate(EXOS, 1):
    vus = {}
    ordre = []
    for r in rows(annexe("G", i, "journal-tva", ex)):
        d = str(r[0])[:10]
        if not d.startswith("20"):
            continue
        try:
            ts = e(r[9]); ttc = e(r[7])
        except (ValueError, TypeError):
            continue
        k = (d, str(r[1])[:5], r[2])
        if k in vus:
            continue
        vus[k] = True
        ordre.append((ts, ttc, d))
    ordre.sort()
    ruptures = 0
    prev = None
    for ts, ttc, _d in ordre:
        if prev is not None and abs(round(ts - prev, 2) - ttc) > 0.005:
            ruptures += 1
        prev = ts
    somme = round(sum(t for _, t, _ in ordre), 2)
    delta = round(ordre[-1][0] - ordre[0][0] + ordre[0][1], 2)
    chainage.append((ex, len(ordre), somme, delta, round(delta - somme, 2), ruptures,
                     len(del_ex[ex])))


# ------------------------------------------------------- 3. Journée du 09/09/2022
dels_a = [d for d in del_ex[EX_A] if d[0] == JOUR_A]
notes_a = notes_ex[EX_A][JOUR_A]
lignes_a = lignes_ex[EX_A][JOUR_A]

# composition de chaque note du jour
compo_a = collections.defaultdict(list)
for no, h, ttc, lib, q, pu in lignes_a:
    compo_a[no].append((lib, q, pu))

# index des prix pratiqués CE JOUR-LA (pu et qté x pu)
pu_jour_a = collections.defaultdict(set)
qpu_jour_a = collections.defaultdict(set)
for no, h, ttc, lib, q, pu in lignes_a:
    pu_jour_a[c(pu)].add(lib)
    qpu_jour_a[c(q * pu)].add(f"{lib} x{q:g}")

rafales_a = collections.OrderedDict()
for d, h, m, _id, z in dels_a:
    rafales_a.setdefault(h, []).append((m, _id, z))

par_total_a = collections.defaultdict(list)
for no, n in notes_a.items():
    par_total_a[c(n["ttc"])].append(no)

# Aucune cellule n'est laissee vide sous une colonne interrogative : la reponse
# est portee en toutes lettres (OUI / NON), et les deux colonnes qui decrivent
# la note rapprochee portent « sans objet » lorsqu'il n'y en a pas.
lignes_rafales = []
for h in sorted(rafales_a):
    lot = rafales_a[h]
    somme = round(sum(m for m, _, _ in lot), 2)
    match = par_total_a.get(c(somme), [])
    if match:
        no = match[0]
        n = notes_a[no]
        libelle = " + ".join(f"{q:g}x {lib} a {pu:.2f} EUR" for lib, q, pu in compo_a[no])
        cible = (f"note n°{no} encaissee a {n['heure']} pour {n['ttc']:.2f} EUR "
                 f"(id reglement {n['id']}, {'/'.join(md for md, _ in n['regl'])})")
    else:
        libelle = "sans objet"
        cible = "sans objet"
    montants = " + ".join(f"{m:.2f}" for m, _, _ in sorted(lot))
    au_prix = sum(1 for m, _, _ in lot if c(m) in pu_jour_a)
    lignes_rafales.append((h, len(lot), montants, somme,
                           "OUI" if match else "NON", cible, libelle,
                           f"{au_prix} sur {len(lot)}"))

NB_RAFALES_A = len(rafales_a)
NB_RAFALES_A_OK = sum(1 for l in lignes_rafales if l[4] == "OUI")

detail_a = []
for d, h, m, _id, z in dels_a:
    v = c(m)
    pu_match = sorted(pu_jour_a.get(v, set()))
    qpu_match = sorted(qpu_jour_a.get(v, set()))
    prix = prix_ex[EX_A]
    if v in prix:
        nature = "prix unitaire de la carte"
    else:
        k = next((k for k in range(2, 25) if v % k == 0 and v // k in prix), None)
        nature = f"quantite {k} x un prix de la carte" if k else "autre"
    s = round(sum(x[0] for x in rafales_a[h]), 2)
    match = par_total_a.get(c(s), [])
    detail_a.append((h, m, _id, z, nature,
                     ", ".join(pu_match[:4]),
                     ", ".join(qpu_match[:4]),
                     f"note n°{match[0]}" if match else ""))

notes_a_rows = []
for no in sorted(notes_a):
    n = notes_a[no]
    notes_a_rows.append((no, n["heure"], n["ttc"], n["id"],
                         " / ".join(f"{md} {mt:.2f}" for md, mt in n["regl"]),
                         " + ".join(f"{q:g}x {lib} ({pu:.2f})" for lib, q, pu in compo_a[no])))
CA_JOUR_A = round(sum(n["ttc"] for n in notes_a.values()), 2)
DEL_JOUR_A = round(sum(m for _d, _h, m, _i, _z in dels_a), 2)
ESP_JOUR_A = round(sum(mt for n in notes_a.values() for md, mt in n["regl"]
                       if md == "ESP"), 2)


# ------------------------------------------------------- 4. Journée du 21/03/2024
dels_b_jour = [d for d in del_ex[EX_B] if d[0] == JOUR_B]
notes_b = notes_ex[EX_B][JOUR_B]
lignes_b = lignes_ex[EX_B][JOUR_B]
compo_b = collections.defaultdict(list)
for no, h, ttc, lib, q, pu in lignes_b:
    compo_b[no].append((lib, q, pu))

# prix des articles des notes n°13 et n°14 (la table partagee citee p. 17)
pu_1314 = collections.defaultdict(set)
for no in (13, 14):
    for lib, q, pu in compo_b[no]:
        pu_1314[c(pu)].add(lib)

# prix pratiques ailleurs le meme soir
pu_soir = collections.defaultdict(set)
for no, h, ttc, lib, q, pu in lignes_b:
    if no in (13, 14):
        continue
    pu_soir[c(pu)].add(f"{lib} (note n°{no})")

detail_b = []
for d, h, m, _id, z in dels_b_jour:
    v = c(m)
    prix = prix_ex[EX_B]
    if v in prix:
        nature = "prix unitaire de la carte"
    else:
        k = next((k for k in range(2, 25) if v % k == 0 and v // k in prix), None)
        nature = f"quantite {k} x un prix de la carte" if k else "hors carte"
    detail_b.append((h, m, _id, z, nature,
                     ", ".join(sorted(pu_1314.get(v, set()))),
                     ", ".join(sorted(prix.get(v, set()))[:5])))

# partage des quantites entre les notes 13 et 14
q13 = {lib: q for lib, q, pu in compo_b[13]}
q14 = {lib: q for lib, q, pu in compo_b[14]}
pux = {lib: pu for lib, q, pu in compo_b[13]}
partage = []
for lib in q13:
    a, b = q13[lib], q14.get(lib, 0)
    tot = round(a + b, 2)
    quart = round(tot / 4, 4)
    conforme = abs(a - quart) < 1e-6
    dels_meme_prix = sum(1 for _d, _h, m, _i, _z in dels_b_jour
                         if _h == "20:16" and c(m) == c(pux[lib]))
    partage.append((lib, pux[lib], a, b, tot, round(quart, 2),
                    "1/4 - 3/4 exact" if conforme else "part ajustee a la main",
                    dels_meme_prix))

CA_JOUR_B = round(sum(n["ttc"] for n in notes_b.values()), 2)
DEL_JOUR_B = round(sum(m for _d, _h, m, _i, _z in dels_b_jour), 2)
ESP_JOUR_B = round(sum(mt for n in notes_b.values() for md, mt in n["regl"]
                       if md == "ESP"), 2)
DEL_2016 = [x for x in dels_b_jour if x[1] == "20:16"]


# --------------------------------------------------------- fréquence des montants
# Le service ecrit p. 17 que la truite (18,60), la salade verte (2,90) et le
# menu galette (19,50) « n'existent pas en DEL dans le fichier tpvenement ».
CONTROLE_P17 = [("TRUITE JURASSIENNE", 18.60), ("Salade Verte", 2.90),
                ("Menu 'Galette'", 19.50)]
freq_p17 = [(lib, m, valeurs.get(c(m), 0)) for lib, m in CONTROLE_P17]


# ---------------------------------- 5. Ecart entre l'annexe C et l'annexe F
# La page retient 16 605 notes (annexe F, les notes encaissees). L'annexe C en
# porte deux de plus. Elles sont identifiees ici, avec leur total TTC.
notes_f = {(d, no) for ex in EXOS for d, ns in notes_ex[ex].items() for no in ns}
ecart_cf = sorted(tickets_c - notes_f)
ECART_CF = (", ".join(f"{d} note n°{no} cloturee a {ttc_c[(d, no)]:.2f} EUR"
                      for d, no in ecart_cf)
            or "aucun")
ECART_FC = sorted(notes_f - tickets_c)


# ------------------------------------------------------------------- ecriture
os.makedirs(SORTIE, exist_ok=True)
wb = openpyxl.Workbook()

TITRE = Font(bold=True, color="FFFFFF", size=11)
FOND = PatternFill("solid", fgColor="1F4E5F")
GRAS = Font(bold=True)


def feuille(nom, colonnes, data, largeurs=None):
    ws = wb.create_sheet(nom[:31])
    ws.append(colonnes)
    for cell in ws[1]:
        cell.font = TITRE
        cell.fill = FOND
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in data:
        ws.append(list(row))
    for j, _ in enumerate(colonnes, 1):
        L = largeurs[j - 1] if largeurs and j - 1 < len(largeurs) else 18
        ws.column_dimensions[get_column_letter(j)].width = L
    ws.freeze_panes = "A2"
    return ws


ws = wb.active
ws.title = "Synthese"
synth = [
    ("Pièce", "R1-suppressions-rapprochements.xlsx"),
    ("Objet", "Réfutation de la partie A (lignes « DEL ») de la réponse DDFiP 39 du 04/09/2026"),
    ("Source", "ANNEXE-C/E/F/G 1-2-3, fichiers remis au titre de l'article L. 47 A, II du LPF"),
    ("Script", "scripts/reponse1-suppressions.py (100 % recalculé, aucune saisie manuelle)"),
    ("", ""),
    ("Lignes « DEL » sur les trois exercices", TOTAL_DEL),
    ("Montant total de ces lignes (EUR)", TOTAL_DEL_EUR),
    ("Suppression médiane (EUR)", MEDIANE_DEL),
    ("Valeurs distinctes prises par ces lignes", NB_VALEURS),
    ("dont montant = un prix unitaire de la carte", TOT_EXACT),
    ("dont montant = quantité entière x un prix de la carte", TOT_MULT),
    ("dont ni l'un ni l'autre", TOT_AUTRE),
    ("Lignes égales à un prix de la carte ou à un multiple entier", TOT_CARTE),
    ("Part expliquée par un prix de la carte", f"{PCT_CARTE:.1f} %"),
    ("", ""),
    ("Prix unitaires distincts pratiqués, trois exercices réunis", PRIX_UNION),
    ("   dont exercice " + PRIX_PAR_EX[0][0], PRIX_PAR_EX[0][1]),
    ("   dont exercice " + PRIX_PAR_EX[1][0], PRIX_PAR_EX[1][1]),
    ("   dont exercice " + PRIX_PAR_EX[2][0], PRIX_PAR_EX[2][1]),
    ("Lecture du catalogue", f"Les {PRIX_UNION} prix unitaires distincts sont le catalogue "
                             f"des trois exercices réunis ; le classement ci-dessus est fait "
                             f"exercice par exercice, chaque ligne « DEL » étant confrontée au "
                             f"seul catalogue de son propre exercice "
                             f"({PRIX_PAR_EX[0][1]}, {PRIX_PAR_EX[1][1]} puis "
                             f"{PRIX_PAR_EX[2][1]} prix)."),
    ("", ""),
    ("RAPPROCHEMENT AVEC LE FICHIER DES TICKETS (annexe C)", ""),
    ("Lignes « DEL » rattachées à une journée et à un n° de ticket Z du fichier des tickets",
     f"{RAP_JZ} sur {TOTAL_DEL}, soit {PCT_RAP_JZ:.1f} %"),
    ("dont montant retrouvé au centime sur une ligne de note du même ticket Z",
     f"{RAP_M} sur {TOTAL_DEL}, soit {PCT_RAP_M:.1f} %"),
    ("Lignes non appariées par ce seul critère du montant",
     f"{RAP_RESTE} sur {TOTAL_DEL}, soit {100 - PCT_RAP_M:.1f} %"),
    ("   exercice " + rappro[0][0], f"{rappro[0][4]} appariées sur {rappro[0][1]} ({rappro[0][5]:.1f} %), "
                                    f"{rappro[0][6]} non appariées"),
    ("   exercice " + rappro[1][0], f"{rappro[1][4]} appariées sur {rappro[1][1]} ({rappro[1][5]:.1f} %), "
                                    f"{rappro[1][6]} non appariées"),
    ("   exercice " + rappro[2][0], f"{rappro[2][4]} appariées sur {rappro[2][1]} ({rappro[2][5]:.1f} %), "
                                    f"{rappro[2][6]} non appariées"),
    ("Portée de ce rapprochement", "Le rapprochement est publié en synthèse, exercice par "
                                   "exercice, et non ligne à ligne. Le critère du montant au "
                                   "centime sur le même ticket Z est un critère unique et "
                                   "volontairement strict : il n'apparie ni les lignes déplacées "
                                   "vers une autre table, ni celles refacturées au forfait, ni "
                                   "celles dont la quantité a été corrigée. La démonstration ne "
                                   "repose donc pas sur lui seul, mais aussi sur la forme des "
                                   "montants, sur le totalisateur perpétuel et sur le canal "
                                   "d'encaissement."),
    ("", ""),
    ("Notes encaissées sur les trois exercices (annexe F)", nb_notes),
    ("Notes présentes dans le fichier des tickets (annexe C)", len(tickets_c)),
    ("Écart entre les deux annexes", f"{len(tickets_c) - nb_notes} note(s) : {ECART_CF}"),
    ("Totaux de notes distincts", len(totaux_notes)),
    ("Lecture", f"Les DEL forment une population d'ARTICLES ({NB_VALEURS} valeurs, celles de "
                f"la carte), pas une population de NOTES ({len(totaux_notes)} totaux distincts)."),
    ("", ""),
    ("CANAL D'ENCAISSEMENT (annexe F)", ""),
    ("Total des règlements encaissés, trois exercices (EUR)", REGL_TOTAL),
    ("dont espèces (EUR)", ESPECES),
    ("Part des espèces", f"{PCT_ESPECES:.2f} %"),
    ("Part tracée par un tiers (banque, ANCV, émetteur de titres)", f"{PCT_BANCARISE:.2f} %"),
    ("Rapport entre le montant des lignes « DEL » et les espèces encaissées",
     f"{TOTAL_DEL_EUR:.2f} / {ESPECES:.2f} = {RAPPORT_DEL_ESP:.2f} fois"),
    ("", ""),
    ("Tickets contrôlés au grand total perpétuel", sum(x[1] for x in chainage)),
    ("Ruptures du totalisateur perpétuel", sum(x[5] for x in chainage)),
    ("", ""),
    (f"{JOUR_A} : lignes DEL", len(dels_a)),
    (f"{JOUR_A} : rafales de DEL", NB_RAFALES_A),
    (f"{JOUR_A} : rafales dont la somme = le total exact d'une note encaissée", NB_RAFALES_A_OK),
    (f"{JOUR_A} : CA encaissé de la journée", CA_JOUR_A),
    (f"{JOUR_A} : total des lignes DEL de la journée", DEL_JOUR_A),
    (f"{JOUR_A} : espèces encaissées ce jour-là", ESP_JOUR_A),
    ("", ""),
    (f"{JOUR_B} : lignes DEL", len(dels_b_jour)),
    (f"{JOUR_B} : lignes DEL à 20h16", len(DEL_2016)),
    (f"{JOUR_B} : CA encaissé de la journée", CA_JOUR_B),
    (f"{JOUR_B} : total des lignes DEL de la journée", DEL_JOUR_B),
    (f"{JOUR_B} : espèces encaissées ce jour-là", ESP_JOUR_B),
    (f"{JOUR_B} : total des 8 DEL de 20h16", round(sum(x[2] for x in DEL_2016), 2)),
    (f"{JOUR_B} : total de la table partagée (notes n°13 + n°14)",
     round(notes_b[13]["ttc"] + notes_b[14]["ttc"], 2)),
    ("", ""),
    ("Contrôle de l'affirmation p. 17 (montants « qui n'existent pas en DEL »)", ""),
]
for lib, m, n in freq_p17:
    synth.append((f"   {lib} à {m:.2f} EUR : occurrences en DEL", n))
for k, v in synth:
    ws.append([k, v])
ws.column_dimensions["A"].width = 62
ws.column_dimensions["B"].width = 96
for r in range(1, ws.max_row + 1):
    ws.cell(r, 1).font = GRAS
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")

feuille("09-09-2022 rafales",
        ["Heure", "Nb DEL", "Montants des DEL de la rafale (EUR)",
         "Somme des DEL (EUR)",
         "Somme = total exact d'une note encaissée ? (OUI / NON)",
         "Note encaissée correspondante", "Composition de cette note",
         "DEL de la rafale égales à un prix pratiqué ce jour-là"],
        lignes_rafales, [10, 9, 46, 20, 30, 52, 70, 26])

feuille("09-09-2022 DEL detail",
        ["Heure", "Montant (EUR)", "id tpvenement", "n° Z", "Nature du montant",
         "Article(s) de la carte à ce prix, ce jour-là",
         "Article(s) x quantité à ce montant, ce jour-là",
         "Note reconstituée par la rafale"],
        detail_a, [10, 14, 14, 9, 30, 42, 42, 24])

feuille("09-09-2022 notes",
        ["Note", "Heure", "Total TTC (EUR)", "id règlement", "Règlement(s)",
         "Composition"],
        notes_a_rows, [8, 10, 16, 14, 26, 96])

feuille("21-03-2024 DEL detail",
        ["Heure", "Montant (EUR)", "id tpvenement", "n° Z", "Nature du montant",
         "Article des notes n°13/14 à ce prix",
         "Article(s) de la carte à ce prix (exercice)"],
        detail_b, [10, 14, 14, 9, 32, 30, 52])

feuille("21-03-2024 partage",
        ["Article", "Prix unitaire (EUR)", "Qté note n°13", "Qté note n°14",
         "Total table", "Quart théorique", "Lecture du partage",
         "Nb DEL à ce prix à 20h16"],
        partage, [26, 16, 14, 14, 12, 16, 26, 22])

feuille("Chainage grand total",
        ["Exercice", "Tickets", "Somme des TTC (EUR)",
         "Progression du totalisateur perpétuel (EUR)", "Écart (EUR)",
         "Ruptures de chaînage", "Lignes DEL de l'exercice"],
        chainage, [14, 10, 22, 40, 14, 20, 22])

ws_carte = feuille("DEL vs prix carte",
        ["Exercice", "DEL = prix unitaire exact", "DEL = quantité x prix",
         "DEL ni l'un ni l'autre", "Total DEL"],
        detail_classe + [("TOTAL", TOT_EXACT, TOT_MULT, TOT_AUTRE, TOTAL_DEL)],
        [14, 26, 24, 22, 12])
for col in range(1, 6):
    ws_carte.cell(ws_carte.max_row, col).font = GRAS
ws_carte.append([f"{TOT_CARTE} lignes sur {TOTAL_DEL} valent un prix de la carte "
                 f"ou un multiple entier de ce prix, soit {PCT_CARTE:.1f} % "
                 f"({100 * TOT_EXACT / TOTAL_DEL:.1f} % un prix unitaire exact, "
                 f"{100 * TOT_MULT / TOTAL_DEL:.1f} % une quantité entière x un prix). "
                 f"Les {TOT_AUTRE} lignes restantes font {100 * TOT_AUTRE / TOTAL_DEL:.1f} %."])
ws_carte.cell(ws_carte.max_row, 1).font = GRAS

top = sorted(valeurs.items(), key=lambda kv: -kv[1])[:60]
ws2 = wb["DEL vs prix carte"]
ws2.append([])
ws2.append(["Les 60 montants les plus fréquents (montant EUR, occurrences, "
            "article(s) de la carte à ce prix)"])
ws2.cell(ws2.max_row, 1).font = GRAS
for v, n in top:
    libs = set()
    for ex in EXOS:
        libs |= prix_ex[ex].get(v, set())
    if libs:
        lecture = ", ".join(sorted(libs)[:6])
    else:
        # Aucun prix unitaire a ce montant : on porte la decomposition en
        # quantite entiere x prix de la carte, plutot qu'une cellule vide.
        lecture = "aucun prix unitaire a ce montant"
        for k in range(2, 25):
            if v % k:
                continue
            sous = set()
            for ex in EXOS:
                sous |= prix_ex[ex].get(v // k, set())
            if sous:
                lecture = (f"{k} x {v // k / 100:.2f} EUR : "
                           + ", ".join(sorted(sous)[:5]))
                break
    ws2.append([v / 100, n, lecture])

chemin = os.path.join(SORTIE, "R1-suppressions-rapprochements.xlsx")
wb.save(chemin)

print(f"OK -> {chemin}")
print(f"DEL total : {TOTAL_DEL} | valeurs distinctes : {NB_VALEURS}")
print(f"  prix exact {TOT_EXACT} ({100*TOT_EXACT/TOTAL_DEL:.1f} %)"
      f" | multiple {TOT_MULT} ({100*TOT_MULT/TOTAL_DEL:.1f} %)"
      f" | autre {TOT_AUTRE} ({100*TOT_AUTRE/TOTAL_DEL:.1f} %)")
print(f"Rapprochement tickets : jour+Z {RAP_JZ}/{TOTAL_DEL} ({PCT_RAP_JZ:.1f} %)"
      f" | montant au centime {RAP_M}/{TOTAL_DEL} ({PCT_RAP_M:.1f} %)"
      f" | non appariees {RAP_RESTE}")
print(f"Prix unitaires distincts : {PRIX_UNION} (union) | par exercice "
      f"{[n for _, n in PRIX_PAR_EX]}")
print(f"Montant DEL : {TOTAL_DEL_EUR} EUR | mediane {MEDIANE_DEL} EUR")
print(f"Reglements : {REGL_TOTAL} EUR dont especes {ESPECES} EUR "
      f"({PCT_ESPECES:.2f} %), bancarise {PCT_BANCARISE:.2f} % | "
      f"rapport DEL/especes {RAPPORT_DEL_ESP:.3f}")
print(f"Notes : {nb_notes} | totaux distincts : {len(totaux_notes)}")
print(f"Chainage : {sum(x[5] for x in chainage)} rupture(s) sur "
      f"{sum(x[1] for x in chainage)} tickets")
print(f"{JOUR_A} : {len(dels_a)} DEL ({DEL_JOUR_A} EUR), {NB_RAFALES_A} rafales dont "
      f"{NB_RAFALES_A_OK} = total exact d'une note ; CA du jour {CA_JOUR_A} EUR ; "
      f"especes du jour {ESP_JOUR_A} EUR")
print(f"{JOUR_B} : {len(dels_b_jour)} DEL ({DEL_JOUR_B} EUR) dont {len(DEL_2016)} a 20h16 "
      f"({round(sum(x[2] for x in DEL_2016), 2)} EUR) ; CA du jour {CA_JOUR_B} EUR ; "
      f"especes du jour {ESP_JOUR_B} EUR ; table partagee "
      f"{round(notes_b[13]['ttc'] + notes_b[14]['ttc'], 2)} EUR")
for lib, m, n in freq_p17:
    print(f"  p.17 « {lib} {m:.2f} EUR n'existe pas en DEL » -> {n} occurrences")
