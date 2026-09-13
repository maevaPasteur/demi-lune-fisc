#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REPONSE 1 - Sur-versement au verre : separation des assiettes + FOURCHETTE bornee
par le bilan matiere, boisson par boisson et exercice par exercice.

Le service (reponse du 04/09/2026, p. 80) calcule un taux moyen unique de 17,60 %
(720 L / 4 091 L), l'applique indistinctement au verre, au pichet de 50 cl et a la
bouteille de 75 cl, en deduit des contenants de 17,64 cl, 58,80 cl et 88,20 cl,
juge les deux derniers impossibles, et reporte alors la totalite du sur-versement
sur le seul verre pour conclure a une impossibilite physique.

Ce script fait deux choses :

 1. SEPARER LES ASSIETTES. A partir des ventes de caisse (ANNEXE-D, deja
    normalisees dans calculsBoissons/boissonsHorsCocktail.json), il ventile le
    volume de vin vendu entre le VERRE (verse a la main), le PICHET (rempli au
    contenant, pre-mesure) et la BOUTEILLE (contenant scelle ou verse a table).
    Le sur-versement n'est applique qu'a ce qui est verse a la main.

 2. BORNER LE SUR-VERSEMENT PAR LE STOCK. Pour chaque boisson et chaque exercice,
    identite de comptabilite matiere :
        disponible = achats + stock d'ouverture - stock de cloture
        disponible au service = disponible - usage cuisine (plats et desserts)
        ecart = disponible au service - volume vendu en caisse aux doses de la carte
    Le sur-versement ne peut pas depasser cet ecart. La fourchette compatible avec
    le stock est donc [0 % ; ecart / base servie a la main].

Sources (lecture seule, deja produites par le pipeline calculsBoissons) :
  src/data/calculsBoissons/consoTotaleParBoisson.json   (achats et stocks par exercice,
        conso par composante : boissons seches, cocktails, plats/desserts, menus)
  src/data/calculsBoissons/boissonsHorsCocktail.json    (ventes de caisse par boisson
        ET par format de service : Verre / Pichet 50-75 cl / Bouteille)
Elles derivent des annexes de caisse ANNEXE-B et ANNEXE-D, des factures fournisseur
et des inventaires de cloture au 31/03.

Taux de reference, verifies en ligne le 10/09/2026 :
  - Kerr WC, Patterson D, Koenen MA, Greenfield TK, « Alcohol Content Variation of
    Bar and Restaurant Drinks in Northern California », Alcoholism: Clinical and
    Experimental Research, 2008, 32(9):1623-1629. Verre de vin servi : 6,18 oz en
    moyenne pour une mesure de reference de 5 oz, soit + 23,6 % de volume.
    Cocktails : 0,85 oz d'ethanol pour 0,60 oz de reference, soit + 42 %.
    https://pmc.ncbi.nlm.nih.gov/articles/PMC2574782/
  - Wansink B, van Ittersum K, « Shape of glass and amount of alcohol poured:
    comparative study of effect of practice and concentration », BMJ, 2005,
    331(7531):1512-1514, DOI 10.1136/bmj.331.7531.1512. 86 barmen de Philadelphie,
    6,3 ans d'experience en moyenne, mesure demandee 44,3 ml : TABLE 2, ligne
    « All drinks », colonne « Average » : 54,6 ml verses dans le verre bas et large
    (+ 23,2 %) et 46,4 ml dans le verre haut et etroit (+ 4,7 %). Le taux de + 20 %
    retenu pour les spiritueux au verre est INFERIEUR au depassement mesure pour le
    verre bas et large. Article NI retracte NI corrige (verifie sur PubMed et
    Crossref le 10/09/2026).
    https://pmc.ncbi.nlm.nih.gov/articles/PMC1322248/
  - Note de prudence : Kerr et coll. mesurent les SHOTS de spiritueux servis a la
    mesure de reference (1,51 oz pour 1,50 oz, n = 61). C'est pourquoi les taux du
    vin et du cocktail ne sont PAS appliques aux spiritueux au verre, et pourquoi
    la variante « plafonne_sans_spiritueux_l » est publiee : le poste ne pese que
    12,25 L sur 475,36 L.
  - Le taux de + 23,6 % du vin est un calcul (6,18 oz / 5 oz) a partir de la Table 2
    de Kerr et coll. et de la definition du verre de reference americain qu'elle pose
    (« 5 ounces (148 ml) of 12 %ABV wine »). L'article publie de son cote + 43 % en
    ethanol, chiffre PLUS eleve, que nous ne retenons pas.

Sorties :
  src/data/reponse1Calculs/sur-versement-fourchette.json
  public/documents/pieces-reponse-1/R1-sur-versement-fourchette.xlsx
"""
import os, json, collections

ICI = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(ICI, ".."))
CALC = os.path.join(ROOT, "src/data/calculsBoissons")
DATA = os.path.join(ROOT, "src/data/reponse1Calculs")
PIECES = os.path.join(ROOT, "public/documents/pieces-reponse-1")
EXOS = ["2022-2023", "2023-2024", "2024-2025"]

# Categories concernees : vins (macvin inclus) et alcools forts au verre.
# Biere, cremant et petillants sont traites sur leurs pages dediees.
VIN = {"vin_blanc", "vin_rouge", "vin", "vin_rose", "vin_de_liqueur"}
SPIRIT = {"aperitif", "liqueur", "eau_de_vie", "spiritueux", "digestif"}
INSCOPE = VIN | SPIRIT

TAUX_VIN, TAUX_SPIRIT, TAUX_COCKTAIL = 0.236, 0.20, 0.42

# Regimes de service, d'apres le format de vente enregistre en caisse.
MAIN = "verre"        # verse a la main : le sur-versement s'y applique
PREMESURE = "pichet"  # rempli au contenant calibre : sur-versement nul par construction
SCELLE = "bouteille"  # bouteille servie telle quelle : sur-versement nul par construction


def regime(format_service):
    f = (format_service or "").lower()
    if "pichet" in f:
        return PREMESURE
    if "bouteille" in f:
        return SCELLE
    return MAIN


def charger():
    conso = json.load(open(os.path.join(CALC, "consoTotaleParBoisson.json"),
                          encoding="utf-8"))["boissons"]
    hors = json.load(open(os.path.join(CALC, "boissonsHorsCocktail.json"),
                          encoding="utf-8"))["boissons"]
    return conso, hors


def ventilation(hors):
    """{nom: {exo: {verre/pichet/bouteille: litres}}} a partir de la caisse,
    plus le NOMBRE de contenants vendus par regime et par exercice."""
    v = collections.defaultdict(lambda: {e: collections.defaultdict(float) for e in EXOS})
    n = {e: collections.defaultdict(float) for e in EXOS}
    for b in hors:
        if b.get("categorie") not in INSCOPE:
            continue
        r = regime(b.get("format_service"))
        for e in EXOS:
            pp = (b.get("par_periode", {}).get(e, {}) or {})
            v[b["nom_canonique"]][e][r] += pp.get("volume_l", 0.0) or 0.0
            n[e][r] += pp.get("quantite", 0.0) or 0.0
    return v, n


def main():
    conso, hors = charger()
    vent, nb = ventilation(hors)

    boissons = []
    for b in conso:
        if b.get("categorie") not in INSCOPE:
            continue
        nom = b["nom_canonique"]
        taille = b.get("taille_achat_cl") or 0
        inv = b.get("inventaire_fin_contenants_par_periode") or {}
        ach = b.get("achats_litres_par_periode") or {}
        stock_fin = {e: (inv.get(e) or 0) * taille / 100.0 for e in EXOS}

        lignes = []
        for i, e in enumerate(EXOS):
            pp = (b.get("par_periode", {}) or {}).get(e, {}) or {}
            de = pp.get("detail_exact_l", {}) or {}
            seches = de.get("boissons_seches", 0.0) or 0.0
            cocktails = de.get("ingredients_cocktails", 0.0) or 0.0
            cuisine = de.get("plats_desserts", 0.0) or 0.0
            menus = ((pp.get("menu_estime_l") or {}).get("moyen", 0.0)) or 0.0

            achats = ach.get(e) or 0.0
            # Pas d'inventaire au 31/03/2022 : pour le premier exercice, la
            # variation de stock est prise a zero (convention neutre, annoncee).
            stock_ini = stock_fin[EXOS[i - 1]] if i > 0 else stock_fin[e]
            dispo = achats + stock_ini - stock_fin[e]
            dispo_service = dispo - cuisine

            vv = vent.get(nom, {}).get(e, {})
            v_verre = vv.get(MAIN, 0.0)
            v_pichet = vv.get(PREMESURE, 0.0)
            v_bouteille = vv.get(SCELLE, 0.0)

            # L'alcool des menus est consomme (il sort du stock) mais la caisse
            # n'en detaille pas le format : il entre dans le volume vendu et reste
            # HORS de la base de sur-versement, comme dans le memoire du 10/07/2026.
            vendu_nominal = seches + cocktails + menus
            base_main = v_verre + cocktails          # ce qui est verse a la main
            taux_sec = TAUX_VIN if b["categorie"] in VIN else TAUX_SPIRIT
            retenu = v_verre * taux_sec + cocktails * TAUX_COCKTAIL

            ecart = dispo_service - vendu_nominal
            borne = ecart > 0 and base_main > 0
            taux_max = (ecart / base_main) if borne else 0.0
            # PLAFONNEMENT : le sur-versement ne peut jamais depasser l'ecart de
            # bilan matiere de la ligne. Quand l'ecart est negatif (la conso depasse
            # les achats FACTURES connus), le bilan matiere ne borne rien : la ligne
            # est signalee, pas plafonnee.
            plafonne = min(retenu, ecart) if borne else retenu

            lignes.append({
                "exercice": e,
                "achats_l": round(achats, 2),
                "stock_ouverture_l": round(stock_ini, 2),
                "stock_cloture_l": round(stock_fin[e], 2),
                "cuisine_l": round(cuisine, 2),
                "disponible_l": round(dispo, 2),
                "disponible_service_l": round(dispo_service, 2),
                "vendu_nominal_l": round(vendu_nominal, 2),
                "verre_l": round(v_verre, 2),
                "pichet_l": round(v_pichet, 2),
                "bouteille_l": round(v_bouteille, 2),
                "cocktails_l": round(cocktails, 2),
                "base_main_l": round(base_main, 2),
                "ecart_l": round(ecart, 2),
                "taux_min": 0.0,
                "taux_max": round(taux_max, 4),
                "taux_retenu": round(retenu / base_main, 4) if base_main > 0 else 0.0,
                "surversement_retenu_l": round(retenu, 2),
                "surversement_plafonne_l": round(plafonne, 2),
                "compatible": (retenu <= ecart + 1e-9) if borne else None,
                "borne": borne,
                "stock_ouverture_estime": i == 0,
            })
        if any(l["vendu_nominal_l"] > 0 or l["disponible_l"] > 0 for l in lignes):
            cumul = {
                "achats_l": round(sum(l["achats_l"] for l in lignes), 2),
                "disponible_service_l": round(sum(l["disponible_service_l"] for l in lignes), 2),
                "vendu_nominal_l": round(sum(l["vendu_nominal_l"] for l in lignes), 2),
                "ecart_l": round(sum(l["ecart_l"] for l in lignes), 2),
                "base_main_l": round(sum(l["base_main_l"] for l in lignes), 2),
                "verre_l": round(sum(l["verre_l"] for l in lignes), 2),
                "pichet_l": round(sum(l["pichet_l"] for l in lignes), 2),
                "bouteille_l": round(sum(l["bouteille_l"] for l in lignes), 2),
                "cocktails_l": round(sum(l["cocktails_l"] for l in lignes), 2),
                "retenu_l": round(sum(l["surversement_retenu_l"] for l in lignes), 2),
                "plafonne_l": round(sum(l["surversement_plafonne_l"] for l in lignes), 2),
            }
            bm, ec = cumul["base_main_l"], cumul["ecart_l"]
            cumul["taux_max"] = round(ec / bm, 4) if (bm > 0 and ec > 0) else 0.0
            cumul["taux_retenu"] = round(cumul["retenu_l"] / bm, 4) if bm > 0 else 0.0
            cumul["achats_connus"] = cumul["achats_l"] > 0
            cumul["compatible_cumul"] = (cumul["retenu_l"] <= ec + 1e-9) if bm > 0 else None
            boissons.append({"nom": nom, "categorie": b["categorie"],
                             "cumul": cumul, "lignes": lignes})

    # ---- Agregat par regime de service, exercice par exercice -------------
    regimes = {e: collections.defaultdict(float) for e in EXOS}
    for b in boissons:
        cat = b["categorie"]
        for l in b["lignes"]:
            r = regimes[l["exercice"]]
            r["disponible_service"] += l["disponible_service_l"]
            r["vendu_nominal"] += l["vendu_nominal_l"]
            r["ecart"] += l["ecart_l"]
            if cat in VIN:
                r["vin_verre"] += l["verre_l"]
                r["vin_pichet"] += l["pichet_l"]
                r["vin_bouteille"] += l["bouteille_l"]
            else:
                r["spiritueux_verre"] += l["verre_l"]
                r["spiritueux_bouteille"] += l["bouteille_l"]
            r["cocktails"] += l["cocktails_l"]

    total = collections.defaultdict(float)
    for e in EXOS:
        for k, v in regimes[e].items():
            total[k] += v

    # Variantes plafonnees ligne a ligne par le bilan matiere.
    v_retenu = v_plafonne = v_prudent = 0.0
    # Ventilation du sur-versement PLAFONNE par regime de service : le plafond de
    # la ligne est reparti au prorata entre sa composante « verre » et sa
    # composante « cocktail ».
    pl = {"vin_verre": 0.0, "spiritueux_verre": 0.0, "cocktails": 0.0}
    n_bornees = n_lignes = 0
    lignes_incompatibles = []
    achats_incomplets = []
    for b in boissons:
        c = b["cumul"]
        if c["retenu_l"] > 0 and not c["achats_connus"]:
            achats_incomplets.append({"boisson": b["nom"], "retenu_l": c["retenu_l"]})
        for l in b["lignes"]:
            v_retenu += l["surversement_retenu_l"]
            if l["surversement_retenu_l"] > 0:
                k = l["surversement_plafonne_l"] / l["surversement_retenu_l"]
                tx = TAUX_VIN if b["categorie"] in VIN else TAUX_SPIRIT
                cle = "vin_verre" if b["categorie"] in VIN else "spiritueux_verre"
                pl[cle] += l["verre_l"] * tx * k
                pl["cocktails"] += l["cocktails_l"] * TAUX_COCKTAIL * k
            if l["base_main_l"] > 0:
                n_lignes += 1
            if l["borne"]:
                n_bornees += 1
                v_plafonne += l["surversement_plafonne_l"]
                v_prudent += l["surversement_plafonne_l"]
                if not l["compatible"]:
                    lignes_incompatibles.append({
                        "boisson": b["nom"], "categorie": b["categorie"],
                        "exercice": l["exercice"],
                        "base_main_l": l["base_main_l"],
                        "verre_l": l["verre_l"], "cocktails_l": l["cocktails_l"],
                        "retenu_l": l["surversement_retenu_l"],
                        "ecart_l": l["ecart_l"],
                        "plafonne_l": l["surversement_plafonne_l"],
                        "retranche_l": round(l["surversement_retenu_l"] - l["ecart_l"], 2),
                        "taux_retenu": l["taux_retenu"], "taux_max": l["taux_max"],
                        "cumul_achats_l": c["achats_l"],
                        "cumul_ecart_l": c["ecart_l"],
                        "cumul_retenu_l": c["retenu_l"],
                        "compatible_sur_3_exercices": c["compatible_cumul"],
                    })
            else:
                # Ecart negatif : la consommation depasse les achats FACTURES connus,
                # donc le bilan matiere ne borne rien (achats incomplets, cf. Intermarche).
                v_plafonne += l["surversement_retenu_l"]

    def bloc(r):
        sv_vin = r["vin_verre"] * TAUX_VIN
        sv_spirit = r["spiritueux_verre"] * TAUX_SPIRIT
        sv_cock = r["cocktails"] * TAUX_COCKTAIL
        base_totale = (r["vin_verre"] + r["vin_pichet"] + r["vin_bouteille"]
                       + r["spiritueux_verre"] + r["spiritueux_bouteille"] + r["cocktails"])
        base_main = r["vin_verre"] + r["spiritueux_verre"] + r["cocktails"]
        sv = sv_vin + sv_spirit + sv_cock
        return {
            "disponible_service_l": round(r["disponible_service"], 2),
            "vendu_nominal_l": round(r["vendu_nominal"], 2),
            "ecart_l": round(r["ecart"], 2),
            "taux_max_compatible": round(r["ecart"] / base_main, 4) if base_main else 0.0,
            "vin_verre_l": round(r["vin_verre"], 2),
            "vin_pichet_l": round(r["vin_pichet"], 2),
            "vin_bouteille_l": round(r["vin_bouteille"], 2),
            "spiritueux_verre_l": round(r["spiritueux_verre"], 2),
            "spiritueux_bouteille_l": round(r["spiritueux_bouteille"], 2),
            "cocktails_l": round(r["cocktails"], 2),
            "base_totale_l": round(base_totale, 2),
            "base_main_l": round(base_main, 2),
            "sv_vin_l": round(sv_vin, 2),
            "sv_spiritueux_l": round(sv_spirit, 2),
            "sv_cocktails_l": round(sv_cock, 2),
            "sv_total_l": round(sv, 2),
            "taux_sur_base_totale": round(sv / base_totale, 4) if base_totale else 0.0,
            "taux_sur_base_main": round(sv / base_main, 4) if base_main else 0.0,
        }

    res = {
        "meta": {
            "objet": "Sur-versement au verre : separation des regimes de service et "
                     "fourchette bornee par le bilan matiere, par boisson et par exercice",
            "sources": [
                "src/data/calculsBoissons/consoTotaleParBoisson.json",
                "src/data/calculsBoissons/boissonsHorsCocktail.json",
                "public/documents/caisse-enregistreuse/ANNEXE-B et ANNEXE-D",
            ],
            "taux": {"vin_verre": TAUX_VIN, "spiritueux_verre": TAUX_SPIRIT,
                     "cocktails": TAUX_COCKTAIL, "pichet": 0.0, "bouteille": 0.0},
            "kerr_2008": "Kerr WC, Patterson D, Koenen MA, Greenfield TK, Alcoholism: "
                         "Clinical and Experimental Research, 2008, 32(9):1623-1629 - "
                         "verre de vin 6,18 oz pour 5 oz de reference (+23,6 %), "
                         "cocktails 0,85 oz d'ethanol pour 0,60 oz (+42 %)",
            "wansink_2005": "Wansink B, van Ittersum K, BMJ, 2005, 331(7531):1512-1514 - "
                            "86 barmen, mesure demandee 44,3 ml, verse 54,6 ml en verre "
                            "bas et large (+23,2 %) et 46,4 ml en verre haut et etroit (+4,7 %)",
            "convention_stock_ouverture": "Pas d'inventaire au 31/03/2022 : la variation de "
                                          "stock du premier exercice est prise a zero.",
        },
        "variantes": {
            "retenu_l": round(v_retenu, 2),
            "plafonne_stock_l": round(v_plafonne, 2),
            "plafonne_stock_prudent_l": round(v_prudent, 2),
            "plafonne_par_regime": {k: round(x, 2) for k, x in pl.items()},
            # Variante de controle : Kerr et coll. mesurent les shots de spiritueux
            # servis A la mesure de reference (1,51 oz pour 1,50 oz). On verifie
            # donc ce que devient le total si le poste « spiritueux au verre »
            # est ramene a zero.
            "plafonne_sans_spiritueux_l": round(v_plafonne - pl["spiritueux_verre"], 2),
            "verres_15cl_surcomptes": round(pl["vin_verre"] / 0.15),
            "doses_4cl_surcomptees": round(pl["spiritueux_verre"] / 0.04),
            "lignes_bornees": n_bornees,
            "lignes_base_main": n_lignes,
            "lignes_incompatibles": lignes_incompatibles,
            "boissons_achats_incomplets": achats_incomplets,
            "retenu_achats_incomplets_l": round(sum(x["retenu_l"] for x in achats_incomplets), 2),
        },
        "contenants_vendus": {e: {k: round(x, 1) for k, x in nb[e].items()} for e in EXOS},
        "contenants_vendus_total": {k: round(sum(nb[e][k] for e in EXOS), 1)
                                    for k in (MAIN, PREMESURE, SCELLE)},
        "regimes_par_exercice": {e: bloc(regimes[e]) for e in EXOS},
        "regimes_total": bloc(total),
        "boissons": boissons,
    }

    # --- Total du perimetre publie -----------------------------------------
    # Attention : « regimes_total » agrege les 57 boissons du fichier, y compris
    # celles qui ne sont jamais versees a la main (pichet ou bouteille seuls).
    # Le total publie dans la reponse ne porte que sur les boissons effectivement
    # versees a la main, seules bornables par un taux de sur-versement. C'est ce
    # perimetre, et lui seul, que la feuille « Fourchette 3 exercices » restitue.
    res["total_borne"] = total_borne(res["boissons"])
    res["regimes_total"]["perimetre"] = (
        "57 boissons, toutes assiettes confondues. Le total publie dans la reponse "
        "est celui de « total_borne » : les 46 boissons versees a la main."
    )

    os.makedirs(DATA, exist_ok=True)
    with open(os.path.join(DATA, "sur-versement-fourchette.json"), "w", encoding="utf-8") as f:
        json.dump(res, f, ensure_ascii=False, indent=1)

    ecrire_xlsx(res)

    t = res["regimes_total"]
    print("=== Regimes de service, total 3 exercices (litres) ===")
    for k in ("vin_verre_l", "vin_pichet_l", "vin_bouteille_l", "spiritueux_verre_l",
              "spiritueux_bouteille_l", "cocktails_l", "base_totale_l", "base_main_l",
              "sv_vin_l", "sv_spiritueux_l", "sv_cocktails_l", "sv_total_l"):
        print(f"  {k:26s} {t[k]:10.2f}")
    print(f"  taux sur base totale     {t['taux_sur_base_totale']*100:9.2f} %")
    print(f"  taux sur base main levee {t['taux_sur_base_main']*100:9.2f} %")
    print(f"  disponible au service    {t['disponible_service_l']:10.2f} L")
    print(f"  vendu aux doses carte    {t['vendu_nominal_l']:10.2f} L")
    print(f"  ecart                    {t['ecart_l']:10.2f} L")
    print(f"  taux max compatible      {t['taux_max_compatible']*100:9.2f} %")
    print("  contenants vendus        ", res["contenants_vendus_total"])
    print("  variantes                ", {k: v for k, v in res["variantes"].items()
                                          if k != "lignes_incompatibles"})
    print()
    print("=== Lignes ou le taux des sources depasse le plafond de stock ===")
    for x in res["variantes"]["lignes_incompatibles"]:
        print(f"  {x['boisson'][:26]:26s} {x['exercice']}  retenu {x['retenu_l']:6.2f} L  "
              f"plafond {x['ecart_l']:6.2f} L  retranche {x['retranche_l']:5.2f} L  "
              f"cumul 3 exos : achats {x['cumul_achats_l']:7.1f} L  ecart {x['cumul_ecart_l']:7.2f} L "
              f"retenu {x['cumul_retenu_l']:6.2f} L  "
              f"{'compatible' if x['compatible_sur_3_exercices'] else 'NON compatible'}")
    print(f"  {len(res['variantes']['lignes_incompatibles'])} lignes sur "
          f"{res['variantes']['lignes_bornees']} bornees "
          f"({res['variantes']['lignes_base_main']} lignes a base non nulle)")
    print("  boissons sans achat facture connu :",
          res["variantes"]["boissons_achats_incomplets"])
    print()
    print("=== Fourchette par boisson (total 3 exercices) ===")
    for b in sorted(res["boissons"], key=lambda x: -sum(l["base_main_l"] for l in x["lignes"]))[:20]:
        bm = sum(l["base_main_l"] for l in b["lignes"])
        ec = sum(l["ecart_l"] for l in b["lignes"])
        rt = sum(l["surversement_retenu_l"] for l in b["lignes"])
        if bm <= 0:
            continue
        print(f"  {b['nom'][:34]:34s} base main {bm:8.1f} L  ecart {ec:8.1f} L  "
              f"retenu {rt:7.1f} L  taux max {100*ec/bm if bm else 0:7.1f} %  "
              f"{'compatible' if rt <= ec else 'non borne (achats incomplets)'}")


def fr(x, dec=2):
    return f"{x:,.{dec}f}".replace(",", " ").replace(".", ",")


def total_borne(boissons):
    """Totaux du perimetre publie : les boissons reellement versees a la main.

    Une boisson servie uniquement au pichet ou a la bouteille n'a pas de base
    versee a la main : aucun taux de sur-versement ne peut lui etre applique,
    elle sort donc du perimetre borne. Additionner les colonnes de la feuille
    « Fourchette par boisson », qui les conserve, donne un autre total.
    """
    B = [b for b in boissons if b["cumul"]["base_main_l"] > 0]
    s = lambda cle: round(sum(b["cumul"][cle] for b in B), 2)
    ecart, base = s("ecart_l"), s("base_main_l")
    retenu, plafonne, vendu = s("retenu_l"), s("plafonne_l"), s("vendu_nominal_l")
    return {
        "boissons": len(B),
        "boissons_hors_perimetre": len(boissons) - len(B),
        "achats_l": s("achats_l"),
        "disponible_service_l": s("disponible_service_l"),
        "vendu_nominal_l": vendu,
        "verre_l": s("verre_l"),
        "pichet_l": s("pichet_l"),
        "bouteille_l": s("bouteille_l"),
        "cocktails_l": s("cocktails_l"),
        "ecart_l": ecart,
        "base_main_l": base,
        "retenu_l": retenu,
        "plafonne_l": plafonne,
        "taux_max_compatible": round(ecart / base, 4) if base else 0.0,
        "taux_retenu_sur_base_main": round(retenu / base, 4) if base else 0.0,
        "taux_effectif_sur_vendu": round(plafonne / vendu, 4) if vendu else 0.0,
        "part_de_l_ecart_couverte": round(retenu / ecart, 4) if ecart else 0.0,
        "rapport_ecart_sur_retenu": round(ecart / retenu, 2) if retenu else 0.0,
    }


def ecrire_xlsx(res):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    blanc = Font(bold=True, color="FFFFFF")
    tete = PatternFill("solid", fgColor="0F766E")
    surligne = PatternFill("solid", fgColor="CCE7E2")
    alerte = PatternFill("solid", fgColor="FBD5D5")
    trait = Side(style="thin", color="D8DEE4")
    bord = Border(left=trait, right=trait, top=trait, bottom=trait)
    droite = Alignment(horizontal="right")

    def entetes(ws, cols, ligne):
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=ligne, column=c)
            cell.font, cell.fill, cell.border = blanc, tete, bord
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def largeurs(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def ligne_bordee(ws, depuis=2):
        r = ws.max_row
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).border = bord
            if c >= depuis:
                ws.cell(row=r, column=c).alignment = droite
        return r

    # ---- Feuille 1 : les regimes de service ------------------------------
    ws = wb.active
    ws.title = "Regimes de service"
    ws.append(["Sur-versement par regime de service : le verre, le pichet et la bouteille "
               "ne sont pas la meme assiette"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Le sur-versement ne s'applique qu'a ce qui est verse a la main. Un pichet est "
               "rempli au contenant calibre et une bouteille est servie telle quelle : leur "
               "sur-versement est nul par construction. Volumes issus de la caisse (ANNEXE-D), "
               "ventiles par format de vente."])
    ws["A2"].font = Font(italic=True, size=9, color="64748B")
    ws.append([])
    cols = ["Exercice", "Regime de service", "Volume vendu (L)", "Taux applique",
            "Sur-versement (L)"]
    entetes(ws, cols, 4)
    LIB = [("vin_verre_l", "Vin au verre (verse a la main)", "vin"),
           ("vin_pichet_l", "Vin au pichet 50 / 75 cl (contenant calibre)", "zero"),
           ("vin_bouteille_l", "Vin a la bouteille ou demi-bouteille", "zero"),
           ("spiritueux_verre_l", "Spiritueux, aperitifs et digestifs au verre", "spirit"),
           ("spiritueux_bouteille_l", "Spiritueux vendus a la bouteille", "zero"),
           ("cocktails_l", "Alcool des cocktails (hors cremant)", "cocktail")]
    TX = {"vin": (TAUX_VIN, "sv_vin_l"), "spirit": (TAUX_SPIRIT, "sv_spiritueux_l"),
          "cocktail": (TAUX_COCKTAIL, "sv_cocktails_l"), "zero": (0.0, None)}
    for e in EXOS + ["Total 3 exercices"]:
        b = res["regimes_total"] if e.startswith("Total") else res["regimes_par_exercice"][e]
        for cle, lib, kind in LIB:
            taux, _ = TX[kind]
            vol = b[cle]
            ws.append([e, lib, round(vol, 2), f"{fr(taux * 100, 1)} %", round(vol * taux, 2)])
            ligne_bordee(ws, 3)
        ws.append([e, "TOTAL", b["base_totale_l"], "", b["sv_total_l"]])
        r = ligne_bordee(ws, 3)
        for c in range(1, len(cols) + 1):
            ws.cell(row=r, column=c).font = Font(bold=True)
            ws.cell(row=r, column=c).fill = surligne
        ws.append([e, "dont base reellement versee a la main", b["base_main_l"],
                   f"{fr(b['taux_sur_base_main'] * 100, 2)} %", b["sv_total_l"]])
        ligne_bordee(ws, 3)
        ws.append([e, "Taux rapporte a la base totale (verre + pichet + bouteille)",
                   b["base_totale_l"], f"{fr(b['taux_sur_base_totale'] * 100, 2)} %",
                   b["sv_total_l"]])
        ligne_bordee(ws, 3)
        ws.append([])
    ws.append([])
    ws.append(["Nombre de contenants vendus (ANNEXE-D), tous exercices"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    cv = res["contenants_vendus_total"]
    tot_cv = sum(cv.values())
    for cle, lib in (("verre", "Verres"), ("pichet", "Pichets 50 / 75 cl"),
                     ("bouteille", "Bouteilles et demi-bouteilles")):
        ws.append(["", lib, cv.get(cle, 0), f"{fr(100 * cv.get(cle, 0) / tot_cv, 1)} %", ""])
        ligne_bordee(ws, 3)
    ws.append([])
    ws.append(["Sur-versement apres plafonnement par le bilan matiere (voir feuille 2)"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    v = res["variantes"]
    for cle, lib in (("retenu_l", "Sur-versement calcule aux taux des sources"),
                     ("plafonne_stock_l", "Plafonne ligne a ligne par l ecart de stock"),
                     ("plafonne_stock_prudent_l",
                      "Plafonne, et ramene a zero quand les achats connus sont incomplets")):
        ws.append(["", lib, v[cle], "", ""])
        ligne_bordee(ws, 3)
    largeurs(ws, [16, 48, 18, 16, 18])
    ws.freeze_panes = "A5"

    # ---- Feuille 2 : la fourchette par boisson et par exercice ------------
    ws2 = wb.create_sheet("Fourchette par boisson")
    ws2.append(["Fourchette de sur-versement compatible avec le stock, boisson par boisson "
                "et exercice par exercice"])
    ws2["A1"].font = Font(bold=True, size=13)
    ws2.append(["Disponible = achats + stock d'ouverture - stock de cloture. Disponible au "
                "service = disponible - usage cuisine (plats et desserts). Ecart = disponible "
                "au service - volume vendu en caisse aux doses de la carte. Le sur-versement "
                "ne peut pas depasser cet ecart : le taux maximal compatible vaut donc "
                "ecart / base versee a la main. Le taux minimal est 0 %. Faute d'inventaire "
                "au 31/03/2022, la variation de stock du premier exercice est prise a zero."])
    ws2["A2"].font = Font(italic=True, size=9, color="64748B")
    ws2.append([])
    cols2 = ["Boisson", "Categorie", "Exercice", "Achats (L)", "Stock ouverture (L)",
             "Stock cloture (L)", "Disponible (L)", "Usage cuisine (L)",
             "Disponible au service (L)", "Vendu en caisse aux doses de la carte (L)",
             "dont verre (L)", "dont pichet (L)", "dont bouteille (L)",
             "dont cocktails (L)", "Base versee a la main (L)", "Ecart (L)",
             "Sur-versement minimal", "Sur-versement maximal compatible",
             "Taux retenu par la defense", "Sur-versement retenu (L)", "Compatible"]
    entetes(ws2, cols2, 4)
    for b in sorted(res["boissons"], key=lambda x: -sum(l["base_main_l"] for l in x["lignes"])):
        for l in b["lignes"]:
            cmp = ("non borne (achats incomplets)" if not l["borne"]
                   else ("oui" if l["compatible"] else "NON"))
            ws2.append([b["nom"], b["categorie"], l["exercice"], l["achats_l"],
                        l["stock_ouverture_l"], l["stock_cloture_l"], l["disponible_l"],
                        l["cuisine_l"], l["disponible_service_l"], l["vendu_nominal_l"],
                        l["verre_l"], l["pichet_l"], l["bouteille_l"], l["cocktails_l"],
                        l["base_main_l"], l["ecart_l"], "0,0 %",
                        (f"{fr(l['taux_max'] * 100, 1)} %" if l["borne"] else "non borne"),
                        f"{fr(l['taux_retenu'] * 100, 1)} %",
                        l["surversement_retenu_l"], cmp])
            r = ligne_bordee(ws2, 4)
            if l["borne"] and not l["compatible"]:
                for c in range(1, len(cols2) + 1):
                    ws2.cell(row=r, column=c).fill = alerte
        ws2.append([])
    # Total de la feuille, et mise en garde sur son perimetre : additionner ces
    # colonnes ne donne pas le total publie, parce que la feuille conserve les
    # boissons servies uniquement au pichet ou a la bouteille.
    lg = [l for b in res["boissons"] for l in b["lignes"]]
    som = lambda cle: round(sum(l[cle] for l in lg), 2)
    tb = res["total_borne"]
    ws2.append([f"TOTAL de la feuille : {len(lg)} lignes, {len(res['boissons'])} boissons",
                "", "", som("achats_l"), som("stock_ouverture_l"), som("stock_cloture_l"),
                som("disponible_l"), som("cuisine_l"), som("disponible_service_l"),
                som("vendu_nominal_l"), som("verre_l"), som("pichet_l"), som("bouteille_l"),
                som("cocktails_l"), som("base_main_l"), som("ecart_l"), "0,0 %", "", "",
                som("surversement_retenu_l"), ""])
    r = ligne_bordee(ws2, 4)
    for c in range(1, len(cols2) + 1):
        ws2.cell(row=r, column=c).font = Font(bold=True)
        ws2.cell(row=r, column=c).fill = surligne
    ws2.append([f"Ce total n'est pas celui que la reponse publie : il englobe "
                f"{tb['boissons_hors_perimetre']} boissons jamais versees a la main "
                f"(pichet ou bouteille seuls), auxquelles aucun taux de sur-versement "
                f"ne s'applique. Le total du perimetre publie est celui de la feuille "
                f"« Fourchette 3 exercices » : {fr(tb['disponible_service_l'], 1)} L "
                f"disponibles au service, {fr(tb['vendu_nominal_l'], 1)} L vendus aux doses "
                f"de la carte, soit {fr(tb['ecart_l'], 1)} L d'ecart."])
    ws2.cell(row=ws2.max_row, column=1).font = Font(italic=True, size=9, color="64748B")
    largeurs(ws2, [30, 16, 14] + [16] * 13 + [18, 24, 22, 20, 13])
    ws2.freeze_panes = "D5"

    # ---- Feuille 3 : la fourchette par boisson, 3 exercices cumules -------
    ws4 = wb.create_sheet("Fourchette 3 exercices")
    ws4.append(["Fourchette de sur-versement compatible avec le stock, par boisson, "
                "sur les trois exercices cumules"])
    ws4["A1"].font = Font(bold=True, size=13)
    ws4.append(["Le taux n'est pas postule : il est encadre. La borne basse est 0 %, la borne "
                "haute est ce que le bilan matiere autorise (ecart / base versee a la main). "
                "Le taux retenu est le PLUS PETIT des deux : celui de la mesure publiee, ou "
                "celui du plafond de stock. Une boisson sans achat facture connu n'est bornee "
                "par rien : elle est signalee comme telle."])
    ws4["A2"].font = Font(italic=True, size=9, color="64748B")
    ws4.append([])
    cols4 = ["Boisson", "Categorie", "Achats factures (L)", "Disponible au service (L)",
             "Vendu aux doses de la carte (L)", "dont verre (L)", "dont pichet (L)",
             "dont bouteille (L)", "dont cocktails (L)", "Ecart (L)",
             "Base versee a la main (L)",
             "Borne basse", "Borne haute compatible avec le stock",
             "Taux des sources", "Sur-versement aux taux des sources (L)",
             "Sur-versement plafonne par le stock (L)",
             "Taux effectif rapporte au volume vendu", "Statut"]
    entetes(ws4, cols4, 4)
    for b in sorted(res["boissons"], key=lambda x: -x["cumul"]["base_main_l"]):
        c = b["cumul"]
        if c["base_main_l"] <= 0:
            continue
        if not c["achats_connus"]:
            statut = "achats factures absents : non borne"
        elif c["ecart_l"] <= 0:
            statut = "conso superieure aux achats factures : non borne"
        elif c["compatible_cumul"]:
            statut = "compatible"
        else:
            statut = "plafonne"
        eff = (c["plafonne_l"] / c["vendu_nominal_l"]) if c["vendu_nominal_l"] > 0 else 0.0
        ws4.append([b["nom"], b["categorie"], c["achats_l"], c["disponible_service_l"],
                    c["vendu_nominal_l"], c["verre_l"], c["pichet_l"], c["bouteille_l"],
                    c["cocktails_l"], c["ecart_l"], c["base_main_l"], "0,0 %",
                    (f"{fr(c['taux_max'] * 100, 1)} %" if c["taux_max"] > 0 else "non borne"),
                    f"{fr(c['taux_retenu'] * 100, 1)} %", c["retenu_l"], c["plafonne_l"],
                    f"{fr(eff * 100, 1)} %", statut])
        r = ligne_bordee(ws4, 3)
        if statut != "compatible":
            for cc in range(1, len(cols4) + 1):
                ws4.cell(row=r, column=cc).fill = alerte
    # Total du perimetre publie : c'est cette ligne, et elle seule, que la
    # reponse reprend.
    ws4.append([f"TOTAL : {tb['boissons']} boissons versees a la main", "",
                tb["achats_l"], tb["disponible_service_l"], tb["vendu_nominal_l"],
                tb["verre_l"], tb["pichet_l"], tb["bouteille_l"], tb["cocktails_l"],
                tb["ecart_l"], tb["base_main_l"], "0,0 %",
                f"{fr(tb['taux_max_compatible'] * 100, 1)} %",
                f"{fr(tb['taux_retenu_sur_base_main'] * 100, 1)} %",
                tb["retenu_l"], tb["plafonne_l"],
                f"{fr(tb['taux_effectif_sur_vendu'] * 100, 1)} %",
                "perimetre publie dans la reponse"])
    r = ligne_bordee(ws4, 3)
    for cc in range(1, len(cols4) + 1):
        ws4.cell(row=r, column=cc).font = Font(bold=True)
        ws4.cell(row=r, column=cc).fill = surligne
    ws4.append([f"Lecture du total : {fr(tb['ecart_l'], 1)} L d'ecart pour "
                f"{fr(tb['base_main_l'], 1)} L verses a la main, soit une borne haute de "
                f"{fr(tb['taux_max_compatible'] * 100, 1)} %. Le sur-versement retenu, "
                f"{fr(tb['retenu_l'], 1)} L, n'en couvre que "
                f"{fr(tb['part_de_l_ecart_couverte'] * 100, 1)} %, soit un rapport de "
                f"{fr(tb['rapport_ecart_sur_retenu'], 2)} : les taux retenus n'absorbent "
                f"pas l'ecart, ils en expliquent une part. Les "
                f"{tb['boissons_hors_perimetre']} boissons servies uniquement au pichet ou "
                f"a la bouteille sont hors de ce total : elles portent elles aussi un ecart "
                f"de bilan matiere, mais aucun sur-versement ne leur est applique."])
    ws4.cell(row=ws4.max_row, column=1).font = Font(italic=True, size=9, color="64748B")
    largeurs(ws4, [32, 16, 18, 20, 22, 14, 14, 16, 16, 14, 20, 14, 26, 16, 24, 24, 22, 34])
    ws4.freeze_panes = "C5"

    # ---- Feuille 4 : les lignes plafonnees --------------------------------
    ws5 = wb.create_sheet("Lignes plafonnees")
    ws5.append(["Les lignes ou le taux de la mesure publiee depasse le plafond de stock, "
                "et le plafonnement applique"])
    ws5["A1"].font = Font(bold=True, size=13)
    ws5.append(["Sur les lignes ci-dessous, le sur-versement calcule au taux de la source "
                "depasse l'ecart de bilan matiere de l'exercice. Il est ramene a l'ecart. "
                "La derniere colonne indique si, sur les trois exercices cumules, le bilan "
                "matiere de la meme boisson reste compatible : le depassement provient alors "
                "du decoupage annuel du stock (inventaire arrete au 31/03) et non d'une "
                "impossibilite materielle."])
    ws5["A2"].font = Font(italic=True, size=9, color="64748B")
    ws5.append([])
    cols5 = ["Boisson", "Exercice", "Base versee a la main (L)", "dont verre (L)",
             "dont cocktails (L)", "Taux des sources", "Plafond de stock",
             "Sur-versement aux taux des sources (L)", "Plafond (ecart, L)",
             "Sur-versement retenu apres plafonnement (L)", "Retranche (L)",
             "Achats factures 3 exercices (L)", "Ecart 3 exercices (L)",
             "Sur-versement 3 exercices (L)", "Compatible sur 3 exercices"]
    entetes(ws5, cols5, 4)
    for x in res["variantes"]["lignes_incompatibles"]:
        ws5.append([x["boisson"], x["exercice"], x["base_main_l"], x["verre_l"],
                    x["cocktails_l"], f"{fr(x['taux_retenu'] * 100, 1)} %",
                    f"{fr(x['taux_max'] * 100, 1)} %", x["retenu_l"], x["ecart_l"],
                    x["plafonne_l"], x["retranche_l"], x["cumul_achats_l"],
                    x["cumul_ecart_l"], x["cumul_retenu_l"],
                    "oui" if x["compatible_sur_3_exercices"] else "NON"])
        ligne_bordee(ws5, 3)
    v = res["variantes"]
    ws5.append([])
    ws5.append([f"{len(v['lignes_incompatibles'])} lignes plafonnees sur "
                f"{v['lignes_bornees']} lignes bornees par le stock ; "
                f"total retranche : {fr(v['retenu_l'] - v['plafonne_stock_l'], 2)} L ; "
                f"sur-versement retenu apres plafonnement : "
                f"{fr(v['plafonne_stock_l'], 2)} L."])
    ws5.cell(row=ws5.max_row, column=1).font = Font(bold=True)
    ws5.append([])
    ws5.append(["Boissons sans achat facture connu sur les trois exercices "
                "(le bilan matiere ne les borne pas) :"])
    ws5.cell(row=ws5.max_row, column=1).font = Font(bold=True)
    for x in v["boissons_achats_incomplets"]:
        ws5.append([x["boisson"], "", "", "", "", "", "", x["retenu_l"]])
        ligne_bordee(ws5, 3)
    ws5.append([f"Variante prudente, ces boissons ramenees a zero : "
                f"{fr(v['plafonne_stock_prudent_l'], 2)} L."])
    ws5.cell(row=ws5.max_row, column=1).font = Font(bold=True)
    largeurs(ws5, [30, 14, 20, 16, 18, 16, 16, 26, 18, 28, 14, 22, 20, 22, 22])
    ws5.freeze_panes = "C5"

    # ---- Feuille 5 : sources et conventions -------------------------------
    ws3 = wb.create_sheet("Sources et conventions")
    ws3.append(["Sources des taux et conventions de calcul"])
    ws3["A1"].font = Font(bold=True, size=13)
    ws3.append([])
    for k, v in res["meta"].items():
        if isinstance(v, (str,)):
            ws3.append([k, v])
        elif isinstance(v, list):
            ws3.append([k, " ; ".join(v)])
        elif isinstance(v, dict):
            ws3.append([k, " ; ".join(f"{a} = {fr(x * 100, 1)} %" for a, x in v.items())])
        ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True)
        ws3.cell(row=ws3.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws3.append(["lien Kerr 2008", "https://pmc.ncbi.nlm.nih.gov/articles/PMC2574782/"])
    ws3.append(["lien Wansink 2005", "https://pmc.ncbi.nlm.nih.gov/articles/PMC1322248/"])
    largeurs(ws3, [32, 120])

    os.makedirs(PIECES, exist_ok=True)
    wb.save(os.path.join(PIECES, "R1-sur-versement-fourchette.xlsx"))


if __name__ == "__main__":
    main()
