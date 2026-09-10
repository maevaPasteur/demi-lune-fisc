#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REPONSE 1 - Partie G : reconciliation des annexes A, B, C, F, G (et H)
=====================================================================
Reponse DDFiP 39 du 04/09/2026, p. 38 a 45.

Le service oppose une « fluctuation » du chiffre d'affaires entre les annexes
B, C, F et G, et une « minoration » des quantites des annexes C par rapport
aux annexes B. Ce script refait le rapprochement a partir des fichiers de
caisse reellement remis et etablit, poste par poste, l'origine de chaque ecart.

READ-ONLY sur les sources. Sortie :
  public/documents/pieces-reponse-1/R1-reconciliation-annexes.xlsx

Lancer :  python3 scripts/reponse1-reconciliation-annexes.py
"""

import collections
import datetime
import os

import xlrd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CAISSE = os.path.join(ROOT, "public", "documents", "caisse-enregistreuse")
OUT_DIR = os.path.join(ROOT, "public", "documents", "pieces-reponse-1")
OUT = os.path.join(OUT_DIR, "R1-reconciliation-annexes.xlsx")

EXOS = [("2022-2023", "1"), ("2023-2024", "2"), ("2024-2025", "3")]


def xd(x):
    return datetime.datetime(*xlrd.xldate_as_tuple(x, 0)).strftime("%Y-%m-%d")


def sheet(nom):
    return xlrd.open_workbook(os.path.join(CAISSE, nom)).sheet_by_index(0)


# ---------------------------------------------------------------- lecture ---
def lire_A(i, lib):
    """Annexe A : synthese CA. Deux blocs : encaissements par mode, ventes par
    famille. Renvoie (total modes, total familles, nb tickets)."""
    sh = sheet("ANNEXE-A%s_synthese-CA_%s.xls" % (i, lib))
    modes, familles, nb = None, None, None
    bloc = None
    for r in range(sh.nrows):
        v = sh.row_values(r)
        if v[0] == "Encaissements Tickets Mode":
            bloc = "modes"
        elif v[0] == "Familles":
            bloc = "familles"
        elif v[0] == "TOTAL":
            if bloc == "modes":
                modes, nb = v[2], v[1]
            elif bloc == "familles":
                familles = v[2]
    return modes, familles, nb


def lire_B(i, lib):
    """Annexe B : prix de vente et quantites (extrait Excel du 13/02/2026).
    Derniere ligne = totaux : Qte, Total TTC, colonne A (prix x qte) ajoutee par
    le service, colonne B (ecart A - Total TTC) ajoutee par le service."""
    sh = sheet("ANNEXE-B%s_prix-vente-quantite_%s.xls" % (i, lib))
    t = sh.row_values(sh.nrows - 1)
    lignes = collections.Counter()
    montants = collections.Counter()
    n = 0
    for r in range(1, sh.nrows):
        v = sh.row_values(r)
        if v[0] == "":
            continue
        k = (xd(v[0]), str(v[2]), v[3], v[4])
        lignes[k] += 1
        montants[k] += v[9]
        n += 1
    return dict(qte=t[5], ca=t[9], ctrl_prix_qte=t[10], ctrl_ecart=t[11],
                nb_lignes=n, lignes=lignes, montants=montants)


def lire_C(i, lib):
    """Annexe C : detail des tickets (archive certifiee du 16/03/2026).
    Derniere ligne = totaux. Colonnes A a G ajoutees par le service."""
    sh = sheet("ANNEXE-C%s_detail-tickets_%s.xls" % (i, lib))
    t = sh.row_values(sh.nrows - 1)
    lignes = collections.Counter()
    montants = collections.Counter()
    tva_flag = collections.Counter()
    tva_corr = 0.0
    par_ticket = collections.defaultdict(float)
    n = 0
    for r in range(1, sh.nrows - 2):
        v = sh.row_values(r)
        if not v[0]:
            continue
        k = (v[0], str(int(v[2])), v[9], v[10])
        lignes[k] += 1
        montants[k] += v[16]
        par_ticket[int(v[8])] = v[5]
        n += 1
        if v[21]:
            tva_flag[v[21].replace("\xa0", " ")] += 1
            if isinstance(v[22], float):
                tva_corr += v[22]
    return dict(qte=t[11], ca=t[16], ctrl_prix_qte=t[17], ctrl_ecart=t[18],
                nb_lignes=n, lignes=lignes, montants=montants,
                tva_flag=tva_flag, tva_corr=tva_corr, par_ticket=dict(par_ticket))


def lire_F(i, lib):
    """Annexe F : reglements. Lignes TOTAL <mode> + TOTAL GENERAL."""
    sh = sheet("ANNEXE-F%s_reglements_%s.xls" % (i, lib))
    total, modes = None, {}
    par_ticket = collections.defaultdict(float)
    nb_lignes_regl = collections.Counter()
    n = 0
    for r in range(1, sh.nrows):
        v = sh.row_values(r)
        if v[8] == "TOTAL":
            modes[v[9]] = v[10]
        elif v[8] == "TOTAL GENERAL":
            total = v[10]
        elif v[0]:
            par_ticket[int(v[8])] += v[10]
            nb_lignes_regl[int(v[8])] += 1
            n += 1
    return dict(ca=total, modes=modes, par_ticket=dict(par_ticket),
                nb_lignes=n, nb_regl=nb_lignes_regl)


def lire_G(i, lib):
    """Annexe G : journal de TVA. La ligne TOTAL porte, dans la colonne
    « tot_sum », la mention « Difference somme Entre 01-04 et 31-03 » : ce
    n'est pas une somme de ventes mais un ECART DE COMPTEUR."""
    sh = sheet("ANNEXE-G%s_journal-tva_%s.xls" % (i, lib))
    for r in range(sh.nrows - 1, 0, -1):
        v = sh.row_values(r)
        if v[0] == "TOTAL":
            return dict(ca=v[9], tva=v[10])
    return dict(ca=None, tva=None)


def lire_H(i, lib):
    """Annexe H : liste des tickets. Porte a la fois la SOMME des tickets
    (tot_ttc) et l'ECART DU COMPTEUR PERPETUEL (tot_sum). On recalcule le
    compteur : a-nouveau, maximum atteint, valeur de la derniere ligne."""
    sh = sheet("ANNEXE-H%s_liste-tickets_%s.xls" % (i, lib))
    ouverture = sh.cell_value(1, 8)
    vals = [sh.cell_value(r, 8) for r in range(2, sh.nrows)
            if isinstance(sh.cell_value(r, 8), float)
            and isinstance(sh.cell_value(r, 0), str)
            and sh.cell_value(r, 0)[:2] == "20"]
    derniere = vals[-1]
    maxi = max(vals)
    tot = None
    for r in range(sh.nrows - 1, 0, -1):
        if sh.cell_value(r, 0) == "TOTAL":
            tot = sh.row_values(r)
            break
    return dict(ca=tot[6], ouverture=ouverture, maxi=maxi, derniere=derniere,
                cpt_max=round(maxi - ouverture, 2),
                cpt_derniere=round(derniere - ouverture, 2))


# ---------------------------------------------------------------- calculs ---
D = {}
for lib, i in EXOS:
    a_modes, a_fam, a_nb = lire_A(i, lib)
    B, C, F, G, H = lire_B(i, lib), lire_C(i, lib), lire_F(i, lib), lire_G(i, lib), lire_H(i, lib)

    # lignes presentes dans C et absentes de B (et inversement)
    manquantes = C["lignes"] - B["lignes"]
    en_trop = B["lignes"] - C["lignes"]
    detail_manq = []
    for k, n in sorted(manquantes.items()):
        part = C["montants"][k] / max(C["lignes"][k], 1) * n
        detail_manq.append((k, n, round(part, 2)))
    montant_manq = round(sum(x[2] for x in detail_manq), 2)

    # tickets regles en plusieurs lignes de reglement
    multi = sum(1 for t, n in F["nb_regl"].items() if n > 1)

    D[lib] = dict(a_modes=a_modes, a_fam=a_fam, a_nb=a_nb, B=B, C=C, F=F, G=G, H=H,
                  detail_manq=detail_manq, montant_manq=montant_manq,
                  nb_manq=sum(manquantes.values()), nb_trop=sum(en_trop.values()),
                  multi=multi)

# --------------------------------------------------------------- ecriture ---
os.makedirs(OUT_DIR, exist_ok=True)
wb = Workbook()

TITRE = Font(bold=True, color="FFFFFF", size=11)
FOND = PatternFill("solid", fgColor="1F3864")
GRAS = Font(bold=True)
EURO = '#,##0.00 "€"'
NB = '#,##0.00'


def entete(ws, cols, largeurs):
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=ws.max_row, column=c)
        cell.font = TITRE
        cell.fill = FOND
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for c, w in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"


# ---- 1. Reconciliation du chiffre d'affaires -------------------------------
ws = wb.active
ws.title = "1 CA par annexe"
entete(ws,
       ["Exercice", "Annexe", "Montant", "Ce que la colonne mesure reellement",
        "Ecart avec l'annexe C", "Explication de l'ecart"],
       [14, 26, 16, 62, 16, 76])

for lib, i in EXOS:
    d = D[lib]
    ref = d["C"]["ca"]
    rows = [
        ("A (ventes par famille)", d["a_fam"],
         "Recapitulatif annuel edite par la caisse, ventes ventilees par famille de produits",
         "Etat de synthese produit par la caisse, remis le 13/01/2026 ; perimetre distinct de l'archive extraite le 16/03/2026"),
        ("A (encaissements par mode)", d["a_modes"],
         "Meme recapitulatif, bloc « Encaissements Tickets Mode » (avoirs inclus)",
         "Difference avec le bloc familles = tickets a montant negatif (remboursements) isoles dans un bloc distinct du meme etat"),
        ("B (prix de vente / quantite)", d["B"]["ca"],
         "Somme de la colonne « Total TTC » de l'extrait Excel du 13/02/2026",
         "Extrait de travail anterieur : toutes ses lignes figurent dans l'annexe C, mais l'annexe C en contient %d de plus (%.2f EUR)" % (d["nb_manq"], d["montant_manq"])),
        ("C (detail des tickets)", d["C"]["ca"],
         "Somme de la colonne « Tot_rem Ttc » de l'archive certifiee du 16/03/2026",
         "REFERENCE : identique au total de l'annexe H (liste des tickets), au centime"),
        ("F (reglements)", d["F"]["ca"],
         "Somme des lignes de reglement (un ticket peut etre regle en plusieurs lignes)",
         "Perimetre encaissements : %d tickets regles en plusieurs lignes ; les avoirs et les differences de date d'encaissement jouent" % d["multi"]),
        ("G (journal de TVA)", d["G"]["ca"],
         "ECART DE COMPTEUR : la cellule porte la mention « Difference somme Entre 01-04 et 31-03 »",
         "Ce n'est pas une somme de ventes : c'est le compteur cumulatif de la caisse arrete a la DERNIERE LIGNE du fichier (fichier trie par n° de ticket, pas par ordre chronologique)"),
        ("H (liste des tickets)", d["H"]["ca"],
         "Somme de la colonne « tot_ttc » des tickets",
         "Egal a l'annexe C au centime : les deux fichiers disent la meme chose"),
    ]
    for nom, val, mesure, expl in rows:
        ecart = None if val is None else round(val - ref, 2)
        ws.append([lib, nom, val, mesure, ecart, expl])
        ws.cell(row=ws.max_row, column=3).number_format = EURO
        ws.cell(row=ws.max_row, column=5).number_format = EURO
        for c in (4, 6):
            ws.cell(row=ws.max_row, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.append([])

# ---- 2. Le compteur perpetuel ---------------------------------------------
ws = wb.create_sheet("2 Compteur perpetuel")
entete(ws,
       ["Exercice", "Compteur a l'ouverture (ligne « A NOUVEAUX »)",
        "Maximum atteint dans l'exercice", "Valeur portee par la derniere ligne du fichier",
        "Maximum - ouverture", "Derniere ligne - ouverture",
        "Total de l'annexe B", "Total de l'annexe G", "Lecture"],
       [14, 22, 20, 22, 18, 20, 18, 18, 70])
for lib, i in EXOS:
    d = D[lib]
    H = d["H"]
    ws.append([lib, H["ouverture"], H["maxi"], H["derniere"], H["cpt_max"], H["cpt_derniere"],
               d["B"]["ca"], d["G"]["ca"],
               "« Maximum - ouverture » = total de l'annexe B ; « derniere ligne - ouverture » = total de l'annexe G. "
               "Les deux chiffres que le service oppose sortent du MEME compteur, arrete a deux endroits differents du meme fichier."])
    for c in range(2, 9):
        ws.cell(row=ws.max_row, column=c).number_format = EURO
    ws.cell(row=ws.max_row, column=9).alignment = Alignment(wrap_text=True, vertical="top")

# ---- 3. Quantites B vs C et colonne de controle du service -----------------
ws = wb.create_sheet("3 Quantites B vs C")
entete(ws,
       ["Exercice", "Annexe", "Somme de la colonne Qte", "Nombre de lignes",
        "CA de l'annexe", "Colonne du service : prix unitaire x Qte",
        "Colonne du service : ecart avec le CA de l'annexe", "Lecture"],
       [14, 30, 18, 16, 16, 22, 24, 76])
for lib, i in EXOS:
    d = D[lib]
    for nom, x in (("B (extrait Excel du 13/02/2026)", d["B"]), ("C (archive du 16/03/2026)", d["C"])):
        ws.append([lib, nom, x["qte"], x["nb_lignes"], x["ca"], x["ctrl_prix_qte"], x["ctrl_ecart"],
                   "Colonnes A et B du fichier, ajoutees par le service lui-meme."])
        for c in (5, 6, 7):
            ws.cell(row=ws.max_row, column=c).number_format = EURO
        ws.cell(row=ws.max_row, column=3).number_format = NB
    ws.append([lib, "Ecart B - C", round(d["B"]["qte"] - d["C"]["qte"], 2),
               d["B"]["nb_lignes"] - d["C"]["nb_lignes"],
               round(d["B"]["ca"] - d["C"]["ca"], 2), None, None,
               "L'ecart de quantite ne produit aucun ecart de recette proportionne : "
               "la colonne de controle du service montre que « prix x Qte » s'ecarte du CA de %.2f EUR dans l'annexe B "
               "et de %.2f EUR seulement dans l'annexe C." % (d["B"]["ctrl_ecart"], d["C"]["ctrl_ecart"])])
    ws.cell(row=ws.max_row, column=3).number_format = NB
    ws.cell(row=ws.max_row, column=5).number_format = EURO
    for c in range(1, 9):
        ws.cell(row=ws.max_row, column=c).font = GRAS
    ws.cell(row=ws.max_row, column=8).alignment = Alignment(wrap_text=True, vertical="top")
    ws.append([])

# ---- 4. Exemple : la journee du 14/04/2022, tickets 11 a 16 ----------------
ws = wb.create_sheet("4 Exemple 14-04-2022")
entete(ws,
       ["Ticket", "Libelle", "Prix unitaire", "Qte annexe B", "Qte annexe C",
        "Montant TTC annexe B", "Montant TTC annexe C"],
       [10, 30, 14, 14, 14, 18, 18])
shB = sheet("ANNEXE-B1_prix-vente-quantite_2022-2023.xls")
shC = sheet("ANNEXE-C1_detail-tickets_2022-2023.xls")
bq, bm = {}, {}
for r in range(1, shB.nrows):
    v = shB.row_values(r)
    if v[0] == "":
        continue
    if xd(v[0]) == "2022-04-14" and str(v[2]) in ("11", "12", "13", "14", "15", "16"):
        bq[(str(v[2]), v[4])] = v[5]
        bm[(str(v[2]), v[4])] = v[9]
cq, cm, cpu = {}, {}, {}
for r in range(1, shC.nrows - 2):
    v = shC.row_values(r)
    if v[0] == "2022-04-14" and int(v[2]) in (11, 12, 13, 14, 15, 16):
        cq[(str(int(v[2])), v[10])] = v[11]
        cm[(str(int(v[2])), v[10])] = v[16]
        cpu[(str(int(v[2])), v[10])] = v[13]
for k in sorted(set(list(bq) + list(cq)), key=lambda k: (int(k[0]), k[1])):
    ws.append([k[0], k[1], cpu.get(k), bq.get(k), cq.get(k), bm.get(k), cm.get(k)])
    for c in (3, 6, 7):
        ws.cell(row=ws.max_row, column=c).number_format = EURO
ws.append(["TOTAL", "", "", round(sum(bq.values()), 2), round(sum(cq.values()), 2),
           round(sum(bm.values()), 2), round(sum(cm.values()), 2)])
for c in range(1, 8):
    ws.cell(row=ws.max_row, column=c).font = GRAS
for c in (6, 7):
    ws.cell(row=ws.max_row, column=c).number_format = EURO
ws.append([])
ws.append(["Lecture : ce sont les extraits que le service reproduit lui-meme en pages 42 (annexe B-1) et 43 (annexe C-1). "
           "Trois additions partagees (11/12, 13/14, 15/16). Les montants sont RIGOUREUSEMENT IDENTIQUES dans les deux annexes. "
           "Seule la colonne « Qte » differe : l'annexe C porte la part reelle sur chacune des deux demi-notes (0,50 + 0,50 = 1 article vendu), "
           "l'annexe B porte 0,50 sur la premiere et 1,00 sur la seconde (1,50 pour un seul article vendu). "
           "C'est l'annexe B qui sur-compte, pas l'annexe C qui minore."])

# ---- 5. Lignes presentes dans C et absentes de B ---------------------------
ws = wb.create_sheet("5 Lignes C absentes de B")
entete(ws, ["Exercice", "Date", "N° ticket", "Reference", "Libelle", "Nb lignes", "Montant TTC"],
       [14, 14, 12, 12, 34, 12, 16])
for lib, i in EXOS:
    d = D[lib]
    for (date, tick, ref, libelle), n, montant in d["detail_manq"]:
        ws.append([lib, date, tick, ref, libelle, n, montant])
        ws.cell(row=ws.max_row, column=7).number_format = EURO
    ws.append([lib, "TOTAL", "", "", "", d["nb_manq"], d["montant_manq"]])
    for c in range(1, 8):
        ws.cell(row=ws.max_row, column=c).font = GRAS
    ws.cell(row=ws.max_row, column=7).number_format = EURO
    ws.append([lib, "Lignes presentes dans B et absentes de C", "", "", "", d["nb_trop"], 0])
    ws.append([])
ws.append(["Lecture : l'annexe B (13/02/2026) est un sous-ensemble strict de l'annexe C (16/03/2026). "
           "Aucune ligne n'a disparu entre les deux remises ; l'archive certifiee en contient davantage. "
           "Une donnee alteree perd de l'information, elle n'en gagne pas."])

# ---- 6. Colonnes TVA ajoutees par le service -------------------------------
ws = wb.create_sheet("6 Colonnes TVA du service")
entete(ws,
       ["Exercice", "Lignes signalees « 10 % / 20 % »", "Lignes signalees « 20 % / 10 % »",
        "Total des lignes signalees", "Total de la colonne « Correction TVA Sur Taux »",
        "Sens de la correction"],
       [14, 24, 24, 20, 30, 46])
for lib, i in EXOS:
    C = D[lib]["C"]
    a = C["tva_flag"].get("10 % / 20 %", 0)
    b = C["tva_flag"].get("20 % / 10 %", 0)
    ws.append([lib, a, b, a + b, round(C["tva_corr"], 2),
               "En faveur du Tresor" if C["tva_corr"] > 0 else "EN DEFAVEUR DU TRESOR : la TVA due DIMINUE"])
    ws.cell(row=ws.max_row, column=5).number_format = EURO
ws.append([])
ws.append(["Lecture : les colonnes E (« Erreur Taux TVA dans Base Envoyee Faux/Correct ») et F (« Correction TVA Sur Taux ») "
           "ne sont pas des colonnes du logiciel de caisse : ce sont des colonnes AJOUTEES par le service dans son propre tableur "
           "(le fichier numerote 1 a 17 les colonnes de la caisse, puis A a G celles que le service ajoute). "
           "Le total de la colonne F, calcule par le service lui-meme, est negatif sur les trois exercices : "
           "si l'on retenait les taux que le service estime corrects, la TVA due par la societe DIMINUERAIT."])

wb.save(OUT)
print("Ecrit :", OUT)
for lib, i in EXOS:
    d = D[lib]
    print("%s | A(fam)=%.2f A(modes)=%.2f B=%.2f C=%.2f F=%.2f G=%.2f H=%.2f" %
          (lib, d["a_fam"], d["a_modes"], d["B"]["ca"], d["C"]["ca"], d["F"]["ca"], d["G"]["ca"], d["H"]["ca"]))
    print("     compteur : max-ouv=%.2f (=B)  derniere-ouv=%.2f (=G)" % (d["H"]["cpt_max"], d["H"]["cpt_derniere"]))
    print("     qte B=%.2f C=%.2f ecart=%.2f | ctrl prix*qte-CA : B=%.2f C=%.2f" %
          (d["B"]["qte"], d["C"]["qte"], d["B"]["qte"] - d["C"]["qte"], d["B"]["ctrl_ecart"], d["C"]["ctrl_ecart"]))
    print("     lignes C absentes de B : %d (%.2f EUR) | lignes B absentes de C : %d" %
          (d["nb_manq"], d["montant_manq"], d["nb_trop"]))
    print("     TVA colonnes service : %d lignes, correction totale %.2f EUR" %
          (sum(d["C"]["tva_flag"].values()), d["C"]["tva_corr"]))
