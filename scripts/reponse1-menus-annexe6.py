#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
R1 - Point 6 : alcool cuit dans les plats des menus (reponse DDFiP du 04/09/2026, p. 68 a 70).

Produit public/documents/pieces-reponse-1/R1-menus-alcool-annexe-6.xlsx :
  1. Origine de chaque dose des menus : le courriel des dirigeants du 30/03/2026, reproduit
     par le service page 37 de sa proposition de rectifications, et la valeur que le service
     retient lui-meme dans ses propres tableaux.
  2. Les 228 L, exercice par exercice et alcool par alcool (le grief de globalisation).
  3. L'annexe n° 6 du service (proposition p. 29-30, 35-36, 41-42), relevee ligne a ligne sur
     les images, confrontee aux factures relues par code article.
  4. Ce que le service oppose p. 69 et 70, et d'ou chaque chiffre vient dans son annexe n° 6.
  5. Le supplement qui reste a retrancher, exercice par exercice, carte et menus separes.
  6. Le Bailey's : les six termes des 236 cl lus en caisse par le service, confrontes aux notres.
  7. Les factures, une a une, pour les quatre alcools discutes au point 6.

Sources (lecture seule) :
  src/data/calculsBoissons/menusCompositionAlcool.json  (dose x probabilite par option de menu)
  src/data/calculsBoissons/menusVentesParPrix.json      (menus vendus, prix catalogue ou non)
  src/data/calculsBoissons/consoAlcoolMenus.json        (alcool des menus par exercice)
  src/data/calculsBoissons/itemsCaisse.json             (quantites vendues, annexe D)
  src/data/factures-fournisseur.json                    (factures FCBS, par code article)
  public/documents/inventaires/inventaire_*.csv         (stocks de cloture)
  public/documents/vins-boissons/_pipeline_conso_alcool.py (mapping libelle caisse -> plat/dose)

Annexe n° 6 : valeurs relevees sur les images des pages 29-30 (exercice 1), 35-36 (exercice 2)
et 41-42 (exercice 3) de la proposition de rectifications du 18/05/2026. Colonnes du service :
Achats FCBS / Achats Inter-M / Stocks debut / Stocks fin / Disponibles.
"""
import ast, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__))) + "/"
JS = BASE + "src/data/calculsBoissons/"
OUT = BASE + "public/documents/pieces-reponse-1/"
EXOS = ["2022-2023", "2023-2024", "2024-2025"]
LAB = {"2022-2023": "Exercice 1", "2023-2024": "Exercice 2", "2024-2025": "Exercice 3"}


def J(n):
    return json.load(open(JS + n, encoding="utf-8"))


# --------------------------------------------------------------------------
# 0. factures relues par code article
# --------------------------------------------------------------------------
FACT = json.load(open(BASE + "src/data/factures-fournisseur.json", encoding="utf-8"))["factures"]


def _periode(df):
    j, m, a = df.split("/")
    iso = f"{a}-{m}-{j}"
    if iso < "2022-04-01":
        return None
    if iso <= "2023-03-31":
        return "2022-2023"
    if iso <= "2024-03-31":
        return "2023-2024"
    if iso <= "2025-03-31":
        return "2024-2025"
    return None


def lignes_code(code):
    """Lignes facturees pour un code article (les lignes non livrees sont signalees)."""
    out = []
    for f in FACT:
        p = _periode(f["dateFacture"])
        if not p:
            continue
        for ln in f.get("lignes", []):
            if str(ln.get("code") or "") != code:
                continue
            out.append({"exo": p, "date": f["dateFacture"], "num": f.get("numeroFacture") or f.get("numero") or "",
                        "designation": ln.get("designation") or "", "qte": ln.get("quantite") or 0,
                        "pu": ln.get("prixUnitaireHT"), "ht": ln.get("montantHT")})
    out.sort(key=lambda r: (r["exo"], r["date"]))
    return out


def unites(code):
    q = {e: 0.0 for e in EXOS}
    for l in lignes_code(code):
        if l["ht"] is not None:
            q[l["exo"]] += l["qte"]
    return q


CODES = {
    "Ravelin blanc 75 cl": ("600867", 75.0),
    "Ravelin blanc 99 cl": ("600933", 99.0),
    "Ravelin BIB 10 L": ("600731", 1000.0),
    "Bailey's 70 cl": ("510133", 70.0),
    "Marc de Bourgogne Jacoulot 45° 70 cl": ("550252", 70.0),
    "Marc du Jura Tissot 50° 70 cl": ("520096", 70.0),
    "Liqueur Golden Eight poire William 70 cl": ("591050", 70.0),
    "Eau-de-vie poire William Massenez 40° 70 cl": ("520131", 70.0),
    "Eau-de-vie poire William Grande Reserve 42° 70 cl": ("520127", 70.0),
}

# --------------------------------------------------------------------------
# 1. doses des menus et leur origine
# --------------------------------------------------------------------------
COURRIEL = {
    ("Ravelin", 10): "« Fondue Jurassienne : 9 ou10 cl de vin/Pers. »",
    ("Vin jaune", 1): "« Sauce Jurassienne : 10 cl de vin jaune pour 1 litre de sauce pour 10 parts = 1 cl de vin jaune/Pers. »",
    ("Porto", 1): "« Sauce Forestière et Sauce Morilles : 10 cl de Porto pour 1 litre de sauce pour 10 parts = 1 cl de Porto/Pers. »",
    ("Calvados", 4): "« Camembert, Assiette Franc-Comtoise et Assiette du Père Grégoire : Ce sont des plats flambés devant les clients, environ 4 à 5 cl d’alcool/Pers. (Macvin ou Calvados) »",
    ("Macvin", 4): "« Babas : 4 cl d’alcool (Quelque soit l’alcool) » et « […] environ 4 à 5 cl d’alcool/Pers. (Macvin ou Calvados) »",
    ("Marc de bourgogne", 4): "« Babas : 4 cl d’alcool (Quelque soit l’alcool) »",
    ("Liqueur de poire", 4): "« Babas : 4 cl d’alcool (Quelque soit l’alcool) »",
    ("Baileys", 4): "« Babas : 4 cl d’alcool (Quelque soit l’alcool) » et « Coupes de glaces alcoolisées : 4 cl d’alcool (Quelque soit l’alcool) »",
    ("Crème de cassis", 4): "« Coupes de glaces alcoolisées : 4 cl d’alcool (Quelque soit l’alcool) »",
}
SERVICE_DOSE = {
    "Ravelin": "10 cl : dose du repère E, en-tête du tableau de la p. 66 (« Fondue 10 cl »)",
    "Vin jaune": "1 cl : repère E, p. 66, et « conformément aux dires des dirigeants »",
    "Porto": "1 cl : repère G, p. 67, « la dose d’un centilitre fournie par les dirigeants lors des opérations sur place »",
    "Calvados": "4 cl : repère D, en-tête du tableau de la p. 65",
    "Macvin": "4 cl : repère Q, p. 67 (« 4 cl, pour la préparation des plats/desserts »), 5 cl dans la proposition p. 48",
    "Marc de bourgogne": "aucune dose propre : le service ne discute pas la dose, seulement le volume disponible (p. 70)",
    "Liqueur de poire": "aucune dose propre : le service ne discute pas la dose, seulement le volume disponible (p. 70)",
    "Baileys": "aucune dose propre : le service ne discute pas la dose, seulement le volume disponible (p. 69)",
    "Crème de cassis": "aucune dose propre : le service retranche 10,64 / 11,32 / 11,32 L au titre des desserts (p. 69)",
}

mca = J("menusCompositionAlcool.json")
doses_rows = []
for carte, cv in mca["cartes"].items():
    for menu, mv in cv["menus"].items():
        for service, options in mv["services"].items():
            for it in options:
                doses_rows.append({
                    "carte": carte, "dates": cv["dates"], "menu": menu, "service": service,
                    "option": it["option"], "alcool": it["alcool"], "dose": it["dose_cl"],
                    "obligatoire": "oui" if it.get("obligatoire") else "non",
                    "n_options": it["nb_options_service"],
                    "pb": it["proba_basse"], "pm": it["proba_moyenne"], "ph": it["proba_haute"],
                    "uniforme": round(1.0 / it["nb_options_service"], 4),
                    "courriel": COURRIEL.get((it["alcool"], it["dose_cl"]), ""),
                    "service": SERVICE_DOSE.get(it["alcool"], ""),
                })

# --------------------------------------------------------------------------
# 2. les 228 L, exercice par exercice
# --------------------------------------------------------------------------
ALIAS = {"Crème de cassis": "Crème de Cassis", "Marc de bourgogne": "Marc de Bourgogne",
         "Baileys": "Bailey's", "Vin jaune": "Arbois Vin Jaune", "Liqueur de poire": "Liqueur de Poire"}
menus = {}
for e in EXOS:
    for a in J("consoAlcoolMenus.json")["exercices"][e]["alcools"]:
        menus.setdefault(ALIAS.get(a["alcool"], a["alcool"]), {x: 0.0 for x in EXOS})[e] += a["cl_moyen"]

# alcool de cuisine a la carte (hors menus), pour separer les deux grandeurs
src = open(BASE + "public/documents/vins-boissons/_pipeline_conso_alcool.py", encoding="utf-8").read()
FOOD = None
for node in ast.parse(src).body:
    if isinstance(node, ast.Assign) and getattr(node.targets[0], "id", "") == "FOOD":
        FOOD = ast.literal_eval(node.value)
assert FOOD, "FOOD introuvable"
items = {}
for it in J("itemsCaisse.json")["items"]:
    items.setdefault(it["produit"], {e: 0.0 for e in EXOS})
    for e in EXOS:
        items[it["produit"]][e] += it["quantite"].get(e, 0.0)
carte = {}
for libelle, ds in FOOD.items():
    q = items.get(libelle)
    if not q:
        continue
    for (alc, dose, _est) in ds:
        parts = [(alc, 1.0)] if alc != "Grand Marnier ou Calvados" else [("Grand Marnier", 0.5), ("Calvados", 0.5)]
        for (a, coef) in parts:
            d = carte.setdefault(ALIAS.get(a, a), {e: 0.0 for e in EXOS})
            for e in EXOS:
                d[e] += q[e] * dose * coef

# menus vendus au prix catalogue (denominateur du calcul des menus)
mvp = J("menusVentesParPrix.json")
menus_vendus = {}
for _p, pv in mvp["periodes"].items():
    for menu, mv in pv["menus"].items():
        d = menus_vendus.setdefault(menu, {e: [0.0, 0.0] for e in EXOS})
        for l in mv["lignes"]:
            d[pv["exercice"]][0] += l["quantite"]
            if l["prix_modifie"] == "NON":
                d[pv["exercice"]][1] += l["quantite"]

# --------------------------------------------------------------------------
# 3. annexe n° 6 du service, relevee sur les images
# --------------------------------------------------------------------------
# libelle -> exercice -> (achats FCBS, stock debut, stock fin, disponibles) ; None = case vide
ANNEXE6 = {
    "Ravelin Blanc 75 cl": {"2022-2023": (30, 4, 9, 25), "2023-2024": (24, 9, None, 33), "2024-2025": (0, None, None, 0)},
    "Ravelin Blanc 99 cl": {"2022-2023": (0, None, None, 0), "2023-2024": (0, None, 12, -12), "2024-2025": (0, 12, None, 12)},
    "BIB 10 L Ravelin Blanc": {"2022-2023": (0, None, None, 0), "2023-2024": (10, None, 1, 9), "2024-2025": (8, 1, None, 9)},
    "Baileys 70 cl": {"2022-2023": (4, None, None, 4), "2023-2024": (16, None, 1, 15), "2024-2025": (12, 1, 1, 12)},
    "Marc Du Jura 70 cl": {"2022-2023": (9, None, 1, 8), "2023-2024": (3, 1, 3, 1), "2024-2025": (0, 3, None, 3)},
    "Marc De Bourg 70 cl": {"2022-2023": (0, None, None, 0), "2023-2024": (1, None, None, 1), "2024-2025": (1, None, 1, 0)},
    "Liqueur Golden 8 70 cl": {"2022-2023": (2, 1.75, None, 3.75), "2023-2024": (12, None, 2, 10), "2024-2025": (11, 2, None, 13)},
    "Poire 70 cl": {"2022-2023": (1, 1, 1, 1), "2023-2024": (0, 1, None, 1), "2024-2025": (1, None, 1, 0)},
}
PAGE_A6 = {"2022-2023": "p. 29 et 30", "2023-2024": "p. 35 et 36", "2024-2025": "p. 41 et 42"}
# code article des factures correspondant a chaque ligne de l'annexe
A6_CODE = {
    "Ravelin Blanc 75 cl": "600867", "Ravelin Blanc 99 cl": "600933", "BIB 10 L Ravelin Blanc": "600731",
    "Baileys 70 cl": "510133", "Marc Du Jura 70 cl": "520096", "Marc De Bourg 70 cl": "550252",
    "Liqueur Golden 8 70 cl": "591050", "Poire 70 cl": "520131+520127",
}

# --------------------------------------------------------------------------
# 4. ce que le service oppose p. 69 et 70
# --------------------------------------------------------------------------
OPPOSE = [
    {"alcool": "Bailey’s", "page": "p. 69", "oppose": "280 + 1 050 + 70 = 1 400 cl",
     "qualif": "« la totalité des achats disponibles »",
     "origine": "colonne « Disponibles » de l’annexe n° 6 pour les deux premiers exercices (4 et 15 bouteilles) ; le troisième terme, 70 cl, ne s’y retrouve pas : l’annexe porte 12 bouteilles disponibles",
     "exact_a6": 4 + 15 + 12, "unite": 70.0, "achats_code": "510133"},
    {"alcool": "Marc de Bourgogne", "page": "p. 70", "oppose": "70 cl",
     "qualif": "« près de 6 fois supérieur à celui disponible »",
     "origine": "colonne « Disponibles » de la ligne « Marc De Bourg 70 cl » de l’annexe n° 6 : 0 + 1 + 0 = 1 bouteille ; deux des quatre bouteilles facturées sont portées à la ligne « Marc Du Jura 70 cl »",
     "exact_a6": 1, "unite": 70.0, "achats_code": "550252"},
    {"alcool": "Liqueur de poire", "page": "p. 70", "oppose": "140 cl",
     "qualif": "« près de 3 fois supérieur à celui disponible »",
     "origine": "colonne « Disponibles » de la ligne « Poire 70 cl » de l’annexe n° 6 (1 + 1 + 0 = 2 bouteilles), dont les achats 1 / 0 / 1 sont ceux des eaux-de-vie de poire William ; la liqueur de poire William est la ligne « Liqueur Golden 8 70 cl », achats 2 / 12 / 11",
     "exact_a6": 3.75 + 10 + 13, "unite": 70.0, "achats_code": "591050"},
]

# --------------------------------------------------------------------------
# 5. supplement exercice par exercice
# --------------------------------------------------------------------------
# volumes que le service dit avoir deja retranches, par exercice (cl), p. 65 a 69
DEJA = {
    "Calvados": {"2022-2023": 3310, "2023-2024": 3002, "2024-2025": 3005},
    "Arbois Vin Jaune": {"2022-2023": 3717, "2023-2024": 3693, "2024-2025": 3748},
    "Porto": {"2022-2023": 398, "2023-2024": 583, "2024-2025": 556},
    "Macvin": {"2022-2023": 7774, "2023-2024": 8832, "2024-2025": 7406},
    "Crème de Cassis": {"2022-2023": 1064, "2023-2024": 1132, "2024-2025": 1132},
    "Ravelin": {"2022-2023": 0, "2023-2024": 16257, "2024-2025": 15554},
}
DEJA_NOTE = {
    "Calvados": "repère D, p. 65 et 69",
    "Arbois Vin Jaune": "repère E, p. 66 et 69",
    "Porto": "repère G, p. 67 et 69",
    "Macvin": "repère Q, p. 67 et 69",
    "Crème de Cassis": "part « desserts » du volume unique des crèmes, p. 69",
    "Ravelin": "part « cuisine » du volume unique des vins en BIB, p. 69 ; nulle sur l’exercice 1, où l’annexe n° 6 du service porte 0 BIB de ravelin acheté et 0 disponible",
    "Bailey’s": "aucun retranchement identifié",
    "Marc de Bourgogne": "aucun retranchement identifié",
    "Liqueur de Poire": "aucun retranchement identifié",
}
ACHATS_CL = {
    "Ravelin": {e: unites("600867")[e] * 75 + unites("600933")[e] * 99 + unites("600731")[e] * 1000 for e in EXOS},
    "Bailey's": {e: unites("510133")[e] * 70 for e in EXOS},
    "Marc de Bourgogne": {e: unites("550252")[e] * 70 for e in EXOS},
    "Liqueur de Poire": {e: unites("591050")[e] * 70 for e in EXOS},
}
# marc de bourgogne : plafonne sur trois exercices aux 210 cl sortis du stock
PLAFOND_3ANS = {"Marc de Bourgogne": 210.0}

alcools4 = ["Ravelin", "Bailey's", "Liqueur de Poire", "Marc de Bourgogne"]
supp = []
for a in alcools4:
    c = carte.get(a, {e: 0.0 for e in EXOS})
    m = menus.get(a, {e: 0.0 for e in EXOS})
    dem = {e: c[e] + m[e] for e in EXOS}
    if a in PLAFOND_3ANS:
        reste = PLAFOND_3ANS[a]
        plaf = {}
        for e in EXOS:
            plaf[e] = min(dem[e], reste)
            reste -= plaf[e]
    else:
        plaf = {e: min(dem[e], ACHATS_CL.get(a, dem)[e]) for e in EXOS}
    dj = DEJA.get(a, {e: 0.0 for e in EXOS})
    supp.append({"alcool": a, "carte": c, "menus": m, "plaf": plaf, "deja": dj,
                 "supp": {e: max(0.0, plaf[e] - dj[e]) for e in EXOS},
                 "note": DEJA_NOTE.get(a if a != "Bailey's" else "Bailey’s", "")})

# --------------------------------------------------------------------------
# 6. Bailey's : les six termes du service
# --------------------------------------------------------------------------
BAILEYS_SERVICE = [20, 16, 80, 32, 28, 60]  # p. 69, dans l'ordre du courrier
verres = {e: items.get("Baileys", {}).get(e, 0.0) + items.get("BAILEYS", {}).get(e, 0.0) for e in EXOS}
coupes = {e: items.get("IVRESSE", {}).get(e, 0.0) for e in EXOS}
BAILEYS_NOUS = [("Bailey’s au verre, " + LAB[e], verres[e], 4.0, verres[e] * 4) for e in EXOS] + \
               [("Coupe glacée IVRESSE à la carte, " + LAB[e], coupes[e], 4.0, coupes[e] * 4) for e in EXOS]

# --------------------------------------------------------------------------
# ecriture xlsx
# --------------------------------------------------------------------------
HEAD = PatternFill("solid", fgColor="1F4E78")
HF = Font(bold=True, color="FFFFFF", size=10)
BAD = PatternFill("solid", fgColor="FCE4E4")
OK = PatternFill("solid", fgColor="E4F3E8")
thin = Side("thin", color="BFBFBF")
B = Border(thin, thin, thin, thin)


def sheet(wb, title, titre, cols, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.cell(1, 1, titre).font = Font(bold=True, size=12)
    for j, h in enumerate(cols, 1):
        c = ws.cell(3, j, h)
        c.fill = HEAD
        c.font = HF
        c.border = B
        c.alignment = Alignment("center", "center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A4"
    return ws


def put(ws, r, vals, fmts=None, fill=None, wrap=()):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.border = B
        if fmts and fmts[j - 1]:
            c.number_format = fmts[j - 1]
        if fill and fill(j):
            c.fill = fill(j)
        if j in wrap:
            c.alignment = Alignment(wrap_text=True, vertical="top")


wb = Workbook()

# --- feuille 1 : origine des doses
ws = sheet(wb, "1-Origine des doses", "1. Chaque dose des plats de menu, et la ligne du courriel des dirigeants du 30/03/2026 dont elle vient (proposition de rectifications, p. 37)",
           ["Carte", "Période", "Menu", "Service", "Option", "Alcool", "Dose retenue (cl)",
            "Plat obligatoire du menu", "Options du service", "Probabilité 1/n", "Probabilité retenue (basse)",
            "Probabilité retenue (moyenne)", "Probabilité retenue (haute)",
            "Ligne du courriel du 30/03/2026, reproduite p. 37 de la proposition",
            "Dose que le service retient lui-même"],
           [7, 20, 15, 9, 28, 17, 9, 11, 10, 10, 11, 12, 11, 62, 50], first=True)
r = 4
for x in doses_rows:
    put(ws, r, [x["carte"], x["dates"], x["menu"], x["service"], x["option"], x["alcool"], x["dose"],
                x["obligatoire"], x["n_options"], x["uniforme"], x["pb"], x["pm"], x["ph"],
                x["courriel"], x["service"]],
        [None] * 6 + ["0"] + [None, "0", "0.000", "0.000", "0.000", "0.000"] + [None, None],
        wrap=(14, 15))
    r += 1
ws.cell(r + 1, 1, "Aucune dose de ce tableau n’est une dose de la société : chacune est celle du courriel des dirigeants du 30 mars 2026, que le service reproduit page 37 de sa proposition de rectifications et dont il écrit page 68 qu’il « n’a fait que suivre les recommandations des dirigeants ».")
ws.cell(r + 2, 1, "Les deux fourchettes du courriel (« 9 ou10 cl » pour la fondue, « environ 4 à 5 cl » pour les plats flambés) sont converties comme le service les convertit lui-même : 10 cl pour la fondue (en-tête de son tableau du repère E, p. 66), 4 cl pour les plats flambés (en-tête de son tableau du repère D, p. 65, et p. 67 pour le macvin).")
ws.cell(r + 3, 1, "La probabilité de choix n’est pas lue en caisse : un menu se vend au forfait et le plat choisi à l’intérieur du menu ne fait pas l’objet d’une ligne de ticket distincte. C’est une convention uniforme 1/n sur les options du service, encadrée par une borne basse et une borne haute, la valeur moyenne étant retenue. Elle ne joue pas sur le ravelin : la fondue jurassienne est le plat unique du menu Dahu, probabilité 1,00.")

# --- feuille 2 : les 228 L exercice par exercice
ws = sheet(wb, "2-228 L par exercice", "2. Les 228 L d’alcool cuit dans les plats des menus, exercice par exercice (le service reproche p. 68 de « regrouper les volumes sur trois exercices »)",
           ["Alcool"] + [f"Menus {LAB[e]} (L)" for e in EXOS] + ["Menus 3 ans (L)", "Arrondi du tableau reproduit p. 68"],
           [20, 15, 15, 15, 15, 20])
r = 4
ordre = sorted(menus, key=lambda a: -sum(menus[a].values()))
for a in ordre:
    put(ws, r, [a] + [round(menus[a][e] / 100, 2) for e in EXOS] +
        [round(sum(menus[a].values()) / 100, 2), round(sum(menus[a].values()) / 100)],
        [None] + ["#,##0.00"] * 4 + ["#,##0"])
    r += 1
put(ws, r, ["TOTAL"] + [round(sum(menus[a][e] for a in menus) / 100, 2) for e in EXOS] +
    [round(sum(sum(menus[a].values()) for a in menus) / 100, 2), 228],
    [None] + ["#,##0.00"] * 4 + ["#,##0"])
for j in range(1, 7):
    ws.cell(r, j).font = Font(bold=True)
ws.cell(r + 2, 1, "Le total exact est de 228,31 L. Les neuf valeurs de la colonne de droite sont des arrondis à l’unité : leur somme fait 227 L, le total imprimé 228 L étant celui des valeurs exactes.")
ws.cell(r + 3, 1, "Aucun terme n’est négatif : une somme de volumes positifs donne le même résultat qu’elle soit présentée agrégée ou exercice par exercice. La présentation sur trois ans ne permet de compenser aucun exercice par un autre.")
r += 6
ws.cell(r, 1, "Menus vendus, et menus retenus dans le calcul (ceux facturés au prix catalogue)").font = Font(bold=True)
r += 1
for j, h in enumerate(["Menu"] + [f"Vendus {LAB[e]}" for e in EXOS] + [f"Retenus {LAB[e]}" for e in EXOS], 1):
    c = ws.cell(r, j, h)
    c.fill = HEAD
    c.font = HF
    c.border = B
r += 1
for menu in ["Dahu", "Demi Lune", "Bourguignon", "Franc-Comtois", "Vegetarien"]:
    d = menus_vendus.get(menu, {e: [0, 0] for e in EXOS})
    put(ws, r, [menu] + [round(d[e][0], 1) for e in EXOS] + [round(d[e][1], 1) for e in EXOS],
        [None] + ["#,##0.0"] * 6)
    r += 1
ws.cell(r + 1, 1, "Le ravelin des menus est le seul poste sans convention : la fondue jurassienne est le plat unique du menu Dahu. 250, 398 puis 351 menus Dahu facturés au prix catalogue, à 10 cl de vin par personne, donnent 2 500, 3 980 et 3 510 cl. Les menus facturés à un prix qui ne correspond à aucun prix de carte sont écartés du calcul, soit 31 menus Dahu sur les trois exercices.")

# --- feuille 3 : annexe n° 6 confrontee aux factures
ws = sheet(wb, "3-Annexe 6 vs factures", "3. L’annexe n° 6 du service (proposition p. 29-30, 35-36, 41-42), relevée sur les images, confrontée aux factures relues par code article",
           ["Ligne de l’annexe n° 6", "Page", "Exercice", "Achats FCBS (annexe)", "Stock début (annexe)",
            "Stock fin (annexe)", "Disponibles (annexe)", "Code article des factures",
            "Achats facturés (unités)", "Écart achats (annexe - factures)"],
           [26, 14, 12, 13, 12, 12, 12, 16, 14, 15])
r = 4
for lib, parex in ANNEXE6.items():
    code = A6_CODE[lib]
    for e in EXOS:
        a, si, sf, di = parex[e]
        if "+" in code:
            fq = sum(unites(c)[e] for c in code.split("+"))
        else:
            fq = unites(code)[e]
        ec = (a or 0) - fq
        put(ws, r, [lib, PAGE_A6[e], LAB[e], a, si if si is not None else "", sf if sf is not None else "",
                    di, code, fq, ec],
            [None, None, None, "#,##0.00", "#,##0.00", "#,##0.00", "#,##0.00", None, "#,##0.00", "#,##0.00"],
            fill=(lambda j: BAD if j == 10 else None) if abs(ec) > 0.001 else (lambda j: OK if j == 10 else None))
        r += 1
ws.cell(r + 1, 1, "Lecture : la colonne « Disponibles » du service est la somme des achats et du stock de début, diminuée du stock de fin. Elle se recalcule sur ses propres colonnes pour toutes les lignes relevées.")
ws.cell(r + 2, 1, "Ravelin : 84 bouteilles de 99 cl facturées sur l’exercice 1 (code 600933) ne figurent pas dans la colonne « Achats FCBS » de l’annexe ; les 24 bouteilles de 99 cl de l’exercice 2 y sont portées à la ligne « Ravelin Blanc 75 cl », et le stock de clôture de 12 bouteilles reste à la ligne « 99 cl », qui affiche un volume disponible de moins 12 bouteilles.")
ws.cell(r + 3, 1, "Marc : le fournisseur facture 4 bouteilles de marc de Bourgogne Jacoulot 45° (code 550252) et 10 bouteilles de marc du Jura Tissot 50° (code 520096). L’annexe en porte 2 et 12 : deux bouteilles de Bourgogne sont rangées à la ligne « Marc Du Jura 70 cl ».")
ws.cell(r + 4, 1, "Poire : la ligne « Poire 70 cl » de l’annexe porte 1, 0 puis 1 bouteille achetée, ce qui est exactement le compte des deux eaux-de-vie de poire William facturées (codes 520131 et 520127). La liqueur de poire William, qui entre dans le baba, est la ligne « Liqueur Golden 8 70 cl » : 2, 12 puis 11 bouteilles, soit les 25 bouteilles du code 591050.")

# --- feuille 4 : ce que le service oppose
ws = sheet(wb, "4-Ce qu’il oppose", "4. Les volumes disponibles que le service oppose p. 69 et 70, et leur origine dans sa propre annexe n° 6",
           ["Alcool", "Page", "Volume opposé", "Qualification du service",
            "D’où vient ce chiffre dans l’annexe n° 6", "Colonne « Disponibles » de l’annexe n° 6, 3 exercices (cl)",
            "Achats facturés, 3 exercices (cl)", "Volume que nous chiffrons (cl)",
            "Part du disponible de l’annexe"],
           [18, 8, 16, 26, 64, 18, 16, 16, 14])
r = 4
for o in OPPOSE:
    a6 = o["exact_a6"] * o["unite"]
    fact = sum(unites(o["achats_code"]).values()) * o["unite"]
    nom = {"Bailey’s": "Bailey's", "Marc de Bourgogne": "Marc de Bourgogne", "Liqueur de poire": "Liqueur de Poire"}[o["alcool"]]
    s = next(x for x in supp if x["alcool"] == nom)
    nous = sum(s["plaf"].values())
    put(ws, r, [o["alcool"], o["page"], o["oppose"], o["qualif"], o["origine"], a6, fact, round(nous, 1),
                round(nous / a6, 3) if a6 else ""],
        [None] * 5 + ["#,##0", "#,##0", "#,##0.0", "0.0%"], wrap=(4, 5))
    r += 1
ws.cell(r + 1, 1, "Bailey’s : les deux premiers termes du service, 280 et 1 050 cl, sont exactement les 4 et 15 bouteilles de sa colonne « Disponibles ». Le troisième, 70 cl, ne s’y retrouve pas : l’annexe porte 12 bouteilles disponibles sur le troisième exercice, soit 840 cl.")
ws.cell(r + 2, 1, "Marc de Bourgogne : le chiffre du service est fidèle à sa colonne. Rétablies dans leur ligne, les 4 bouteilles facturées et la bouteille inventoriée au 31/03/2025 donnent 3 bouteilles sorties du stock, soit 210 cl : c’est le plafond auquel nous ramenons notre demande.")
ws.cell(r + 3, 1, "Liqueur de poire : le service oppose la ligne « Poire 70 cl », qui est celle des eaux-de-vie de poire William. La liqueur de poire du baba est la ligne « Liqueur Golden 8 70 cl » de la même annexe.")

# --- feuille 5 : supplement exercice par exercice
ws = sheet(wb, "5-Supplément", "5. Le supplément qui reste à retrancher, exercice par exercice, carte et menus séparés (litres)",
           ["Alcool", "Exercice", "Cuisine à la carte (L)", "Plats des menus (L)", "Total cuisine (L)",
            "Plafonné aux achats ou au stock (L)", "Déjà retranché par le service (L)",
            "Supplément (L)", "Dont menus (L)", "Origine du retranchement du service"],
           [20, 12, 13, 12, 12, 14, 14, 12, 12, 58])
r = 4
for s in supp:
    for e in EXOS:
        dem = s["carte"][e] + s["menus"][e]
        sp = s["supp"][e]
        part_menus = min(sp, s["menus"][e]) if sp else 0.0
        put(ws, r, [s["alcool"], LAB[e], round(s["carte"][e] / 100, 2), round(s["menus"][e] / 100, 2),
                    round(dem / 100, 2), round(s["plaf"][e] / 100, 2), round(s["deja"][e] / 100, 2),
                    round(sp / 100, 2), round(part_menus / 100, 2), s["note"] if e == EXOS[0] else ""],
            [None, None] + ["#,##0.00"] * 7 + [None], wrap=(10,))
        r += 1
tot = {k: sum(sum(s[k].values()) for s in supp) for k in ("carte", "menus", "plaf", "supp")}
put(ws, r, ["TOTAL", "3 exercices", round(tot["carte"] / 100, 2), round(tot["menus"] / 100, 2),
            round((tot["carte"] + tot["menus"]) / 100, 2), round(tot["plaf"] / 100, 2),
            round(sum(sum(s["deja"].values()) for s in supp) / 100, 2), round(tot["supp"] / 100, 2),
            round(sum(min(s["supp"][e], s["menus"][e]) for s in supp for e in EXOS) / 100, 2), ""],
    [None, None] + ["#,##0.00"] * 7 + [None])
for j in range(1, 11):
    ws.cell(r, j).font = Font(bold=True)
ws.cell(r + 2, 1, "Le supplément est calculé exercice par exercice et planché à zéro : un litre retranché de l’exercice 2 ne neutralise pas un litre cuit sur l’exercice 1, les bases étant établies exercice par exercice, et la reconstitution du service l’étant aussi.")
ws.cell(r + 3, 1, "Ces 98,2 L sont déjà compris dans les 169,9 L demandés au point 5 (page « Alcool de cuisine sous-déduit ») : ils ne s’y ajoutent pas.")

# --- feuille 6 : Bailey's, les six termes
ws = sheet(wb, "6-Bailey’s en caisse", "6. Les 236 cl que le service lit en caisse pour le Bailey’s (p. 69) et les six termes de notre propre chiffrage hors menus",
           ["Terme", "Articles vendus", "Dose (cl)", "Volume (cl)"], [40, 14, 10, 12])
r = 4
for lib, q, d, v in BAILEYS_NOUS:
    put(ws, r, [lib, q, d, v], [None, "#,##0.0", "0", "#,##0.0"])
    r += 1
put(ws, r, ["TOTAL, notre chiffrage hors menus", "", "", sum(v for _l, _q, _d, v in BAILEYS_NOUS)],
    [None, None, None, "#,##0.0"])
for j in range(1, 5):
    ws.cell(r, j).font = Font(bold=True)
r += 2
ws.cell(r, 1, "Les six termes additionnés par le service p. 69 : " + " + ".join(str(x) for x in BAILEYS_SERVICE) + " = " + str(sum(BAILEYS_SERVICE)) + " cl")
r += 1
ws.cell(r, 1, "Nos six termes, triés : " + ", ".join(str(int(v)) for _l, _q, _d, v in sorted(BAILEYS_NOUS, key=lambda x: x[3])) + " cl")
r += 1
ws.cell(r, 1, "Cinq des six termes sont identiques. Le sixième diffère de 12 cl, soit trois coupes.")

# --- feuille 7 : factures
ws = sheet(wb, "7-Factures", "7. Les factures du fournisseur Franche-Comté Boissons Services pour les alcools discutés au point 6, une à une",
           ["Produit", "Code article", "Exercice", "Date", "Facture", "Désignation facturée",
            "Quantité", "Prix unitaire HT", "Montant HT", "Ligne livrée"],
           [34, 11, 12, 11, 11, 48, 10, 12, 12, 11])
r = 4
for nom, (code, cont) in CODES.items():
    for l in lignes_code(code):
        put(ws, r, [nom, code, LAB[l["exo"]], l["date"], l["num"], l["designation"], l["qte"],
                    l["pu"] if l["pu"] is not None else "", l["ht"] if l["ht"] is not None else "",
                    "oui" if l["ht"] is not None else "non (non livrée)"],
            [None] * 6 + ["#,##0.0", "#,##0.00", "#,##0.00", None])
        r += 1
    q = unites(code)
    put(ws, r, [nom + " : total facturé", code, "3 exercices", "", "", f"{cont:.0f} cl l’unité",
                sum(q.values()), "", "", f"{sum(q.values()) * cont:.0f} cl"],
        [None] * 6 + ["#,##0.0", None, None, None])
    for j in range(1, 11):
        ws.cell(r, j).font = Font(bold=True)
    r += 2

# --- feuille 8 : ce que la valorisation au prix du verre represente
# Le service ecrit p. 38 de sa proposition que sa reconstitution croise, pour chaque famille,
# « le prix de vente moyen ainsi que les quantites d'articles enregistrees dans la caisse ».
# On applique donc au volume cuit dans les menus le prix moyen de vente lu dans la caisse,
# exercice par exercice, pour le format sous lequel l'alcool est vendu au comptoir.
IT = {i["produit"]: i for i in J("itemsCaisse.json")["items"]}
FORMAT_VENTE = {
    "Ravelin": (["Verre ALIGOTE", "Verre de vin"], 15.0, "verre de vin blanc maison, 15 cl"),
    "Macvin": (["Macvin du Jura"], 6.0, "verre de macvin, 6 cl"),
    "Arbois Vin Jaune": (["Vin jaune"], 12.0, "verre de vin jaune, 12 cl"),
    "Calvados": (["Calvados (4cl)"], 4.0, "verre de Calvados, 4 cl"),
    "Porto": (["Porto"], 6.0, "verre de porto, 6 cl"),
    "Bailey's": (["Baileys", "BAILEYS"], 4.0, "verre de Bailey’s, 4 cl"),
    "Marc de Bourgogne": (["Vieux Marc Bourgogne"], 4.0, "verre de vieux marc de Bourgogne, 4 cl"),
    "Liqueur de Poire": (["Liqueur de Poire"], 4.0, "verre de liqueur de poire, 4 cl"),
}


def prix_moyen(noms, e):
    q = sum(IT[n]["quantite"].get(e, 0.0) for n in noms if n in IT)
    c = sum((IT[n].get("ca_ttc") or {}).get(e, 0.0) for n in noms if n in IT)
    if q and c:
        return c / q
    ps = []
    for e2 in EXOS:
        q2 = sum(IT[n]["quantite"].get(e2, 0.0) for n in noms if n in IT)
        c2 = sum((IT[n].get("ca_ttc") or {}).get(e2, 0.0) for n in noms if n in IT)
        if q2 and c2:
            ps.append(c2 / q2)
    return sum(ps) / len(ps) if ps else 0.0


ws = sheet(wb, "8-Valorisation", "8. Ce que représente, en chiffre d’affaires liquides, le fait de compter l’alcool cuit dans un plat de menu comme du volume vendable au verre",
           ["Alcool", "Exercice", "Volume cuit dans les menus (L)", "Format de vente au comptoir",
            "Prix moyen lu en caisse (TTC)", "Prix au centilitre (TTC)",
            "Chiffre d’affaires attribué au volume cuit (TTC)"],
           [20, 12, 15, 30, 14, 13, 18])
r = 4
tot_val = 0.0
for a, (noms, vol, lib) in FORMAT_VENTE.items():
    for e in EXOS:
        p = prix_moyen(noms, e)
        cl = menus.get(a, {e: 0.0})[e]
        v = cl * p / vol
        tot_val += v
        put(ws, r, [a, LAB[e], round(cl / 100, 2), lib, round(p, 3), round(p / vol, 4), round(v, 2)],
            [None, None, "#,##0.00", None, "#,##0.000", "#,##0.0000", "#,##0.00"])
        r += 1
put(ws, r, ["TOTAL", "3 exercices",
            round(sum(sum(menus[a].values()) for a in FORMAT_VENTE) / 100, 2), "", "", "", round(tot_val, 2)],
    [None, None, "#,##0.00", None, None, None, "#,##0.00"])
for j in range(1, 8):
    ws.cell(r, j).font = Font(bold=True)
ws.cell(r + 2, 1, "La crème de cassis est exclue de ce tableau : elle n’a aucun prix de vente propre, ni sur la carte ni en caisse, où elle n’apparaît qu’à l’intérieur des cocktails. Les 15,87 L de crème de cassis des menus ne sont donc pas valorisés ici, et le total porte sur 212,4 L des 228,3 L.")
ws.cell(r + 3, 1, "Prix moyen de vente = chiffre d’affaires TTC de l’article divisé par la quantité vendue, exercice par exercice (src/data/calculsBoissons/itemsCaisse.json). C’est la définition que le service donne lui-même page 38 de sa proposition : « le prix de vente moyen ainsi que les quantités d’articles enregistrées dans la caisse ». Le champ « prix_unitaire » du même fichier n’est pas utilisé : il porte des valeurs aberrantes sur trois lignes.")
ws.cell(r + 4, 1, "Ce tableau chiffre le seul effet de la valorisation au niveau du chiffre d’affaires liquides. Il ne préjuge pas du prix unitaire que le service a réellement appliqué à chacun de ces alcools : ses annexes 7-1, 7-2 et 7-3 ne sont pas reproduites dans la réponse du 04/09/2026.")

os.makedirs(OUT, exist_ok=True)
wb.save(OUT + "R1-menus-alcool-annexe-6.xlsx")
print("-> public/documents/pieces-reponse-1/R1-menus-alcool-annexe-6.xlsx")

# --------------------------------------------------------------------------
# recapitulatif console
# --------------------------------------------------------------------------
print("\n== 228 L par exercice ==")
for e in EXOS:
    print(f"   {LAB[e]} : {sum(menus[a][e] for a in menus)/100:.2f} L")
print(f"   total : {sum(sum(menus[a].values()) for a in menus)/100:.2f} L")

print("\n== supplement par alcool ==")
for s in supp:
    print(f"   {s['alcool']:20} carte {sum(s['carte'].values())/100:6.2f}  menus {sum(s['menus'].values())/100:6.2f}"
          f"  plafonne {sum(s['plaf'].values())/100:6.2f}  supplement {sum(s['supp'].values())/100:6.2f}")
print(f"   TOTAL supplement : {sum(sum(s['supp'].values()) for s in supp)/100:.2f} L")
print(f"   dont menus : {sum(min(s['supp'][e], s['menus'][e]) for s in supp for e in EXOS)/100:.2f} L")
print(f"   dont carte : {sum(max(0.0, s['supp'][e]-s['menus'][e]) for s in supp for e in EXOS)/100:.2f} L")

print("\n== annexe 6 vs factures, disponibles cumules ==")
for lib, parex in ANNEXE6.items():
    di = sum(parex[e][3] for e in EXOS)
    code = A6_CODE[lib]
    fq = sum(sum(unites(c)[e] for c in code.split("+")) for e in EXOS)
    print(f"   {lib:26} disponibles annexe {di:7.2f}  achats annexe {sum(parex[e][0] for e in EXOS):5}"
          f"  achats factures {fq:5.0f}")

print(f"\n== valorisation CA liquides (hors creme de cassis) : {tot_val:.2f} EUR TTC ==")
print("\n== Bailey's ==")
print("   service :", BAILEYS_SERVICE, "=", sum(BAILEYS_SERVICE), "cl")
print("   nous    :", sorted(int(v) for _l, _q, _d, v in BAILEYS_NOUS), "=",
      sum(v for _l, _q, _d, v in BAILEYS_NOUS), "cl")
