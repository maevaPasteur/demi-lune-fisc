#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REPONSE 1 - Cremant : le tableau de la page 62 relu jour par jour.

OBJET. Le service (reponse du 04/09/2026, p. 62-63, repris p. 84-85) conduit
son test au niveau ANNUEL : « la lecture des donnees de caisse permet
d'affirmer que 118,08 litres ont ete utilises pour la confection des articles
vendus » (p. 63), puis il compare ce total au volume disponible de l'exercice.
Une lecture annuelle suppose implicitement qu'une bouteille ouverte peut etre
reportee d'un service au suivant sans limite : c'est le seul cas ou 118,08 L
peuvent sortir de 158 bouteilles.

Ce script ne discute aucune des donnees du service. Il reprend EXACTEMENT son
perimetre - les seuls libelles de son tableau « DONNEES ISSUES LOGICIEL DE
CAISSE » de la page 62, avec SES doses (cocktails 6 cl, verre 12 cl) et SES
quantites disponibles (275, 325 puis 322 bouteilles de 75 cl) - et se borne a
le lire au niveau ou la contrainte physique existe : la JOURNEE de service.

Il ne retient AUCUN sur-versement : le serveur verse la dose exacte du service.
La seule perte modelisee est arithmetique : si le besoin d'une journee n'est
pas un multiple exact de 75 cl, la derniere bouteille ouverte n'est pas vide.

Trois grandeurs sont produites, toutes dans le perimetre du service :
  1. lecture annuelle   : ceil(besoin annuel / 75 cl) bouteilles ;
  2. lecture journaliere: somme des ceil(besoin du jour / 75 cl) ;
  3. lecture par libelle: somme des ceil(besoin du jour et du libelle / 75 cl),
     c'est-a-dire le cas ou aucune bouteille ne serait partagee entre deux
     articles. L'ecart entre 3 et 2 mesure ce que vaut, en bouteilles, la
     premiere objection du service (« le cremant est utilise pour de nombreux
     articles »), objection deja integree au calcul.

Sources (lecture seule, reproductible) :
  public/documents/caisse-enregistreuse/ANNEXE-C{1,2,3}_detail-tickets_{exo}.xls
Sorties :
  public/documents/pieces-reponse-1/R1-cremant-perimetre-du-service.xlsx
  public/documents/pieces-reponse-1/R1-cremant-perimetre-du-service.csv
"""
import os, math, csv, collections
import xlrd

ICI = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(ICI, ".."))
CAISSE = os.path.join(ROOT, "public/documents/caisse-enregistreuse/")
PIECES = os.path.join(ROOT, "public/documents/pieces-reponse-1/")
EXOS = ["2022-2023", "2023-2024", "2024-2025"]
C = {e: CAISSE + f"ANNEXE-C{i}_detail-tickets_{e}.xls" for i, e in enumerate(EXOS, 1)}
BOUTEILLE = 75.0

# Les seuls libelles servis au verre que le tableau de la page 62 retient, avec
# les doses de sa colonne « Dose (en centilitres) ». Rien n'est ajoute.
PERIMETRE_P62 = {
    "La Vouivre": 6,        # ref. 0357, « La Vouivre ( H ) », dose 6
    "Père Gregoire": 6,     # ref. 0333, « Père Gregoire ( N ) », dose 6
    "KITTYKIR": 6,          # ref. 0001, « KITTYKIR ( O ) », dose 6
    "Crémant": 12,          # ref. 0348, « Crémant (verre) », dose 12
}
# Les deux dernieres lignes du meme tableau : des bouteilles de 75 cl vendues
# scellees (ref. 0362 « Crémant du Jura (Bouteille) » et 1897 « Crément du
# Jura », dose 75). Elles sortent du stock mais ne sont jamais ouvertes au
# service. On s'en tient STRICTEMENT a ces deux libelles, pour que le total
# recompte dans l'annexe C soit exactement celui de son tableau (6 + 4 = 10
# pour le premier exercice). Le libelle « Crémant Rosé », present une fois dans
# la caisse du premier exercice, est donc laisse de cote : l'inclure
# diminuerait le disponible, donc jouerait en notre faveur.
BOUTEILLES_SCELLEES = ["Crémant du Jura", "Crément du Jura"]

# Quantites disponibles retenues PAR LE SERVICE (p. 62 et p. 84).
DISPO_BOUTEILLES = {"2022-2023": 275, "2023-2024": 325, "2024-2025": 322}

# Quantites vendues lues par le service dans son tableau de la page 62,
# premier exercice, pour controle de la correspondance avec l'annexe C.
CONTROLE_P62 = {"La Vouivre": 622, "Père Gregoire": 503,
                "KITTYKIR": 234, "Crémant": 242}


def lignes(fn):
    sh = xlrd.open_workbook(fn).sheet_by_index(0)
    for r in range(1, sh.nrows):
        lib = str(sh.cell_value(r, 10)).strip()
        try:
            q = float(sh.cell_value(r, 11))
        except ValueError:
            q = 0.0
        yield str(sh.cell_value(r, 0))[:10], lib, q


def collecte(ex):
    """Besoin en cremant, par jour et par libelle, dans le perimetre p. 62."""
    jour = collections.defaultdict(float)
    jour_lib = collections.defaultdict(lambda: collections.defaultdict(float))
    qte_lib = collections.defaultdict(float)
    scellees = 0.0
    for date, lib, q in lignes(C[ex]):
        if q <= 0:
            continue
        if lib in PERIMETRE_P62:
            cl = q * PERIMETRE_P62[lib]
            jour[date] += cl
            jour_lib[date][lib] += cl
            qte_lib[lib] += q
        elif lib in BOUTEILLES_SCELLEES:
            scellees += q
    return dict(jour), jour_lib, dict(qte_lib), scellees


REPORTS = [1, 2, 3, 5, 7]


def report(jour, k):
    """Bouteilles a ouvrir si une bouteille ouverte sert k journees de service
    consecutives : les journees sont groupees par blocs de k, et chaque bloc
    n'ouvre que les bouteilles necessaires a son besoin total. k = 1 est le cas
    retenu (aucun report d'un service a l'autre)."""
    days = sorted(jour)
    bt = 0
    for i in range(0, len(days), k):
        bloc = sum(jour[d] for d in days[i:i + k])
        bt += math.ceil(bloc / BOUTEILLE - 1e-9)
    return bt


def calcule(ex):
    jour, jour_lib, qte_lib, scellees = collecte(ex)
    besoin = sum(jour.values())
    bt_annuel = math.ceil(besoin / BOUTEILLE - 1e-9)
    bt_jour = sum(math.ceil(v / BOUTEILLE - 1e-9) for v in jour.values())
    bt_libelle = sum(math.ceil(v / BOUTEILLE - 1e-9)
                     for d in jour_lib.values() for v in d.values())
    dispo = DISPO_BOUTEILLES[ex] * BOUTEILLE
    dispo_service = dispo - scellees * BOUTEILLE
    sorti = bt_jour * BOUTEILLE
    reports = {}
    for k in REPORTS:
        b = report(jour, k)
        reports[k] = {"bouteilles": b, "sorti_cl": b * BOUTEILLE,
                      "jete_cl": b * BOUTEILLE - besoin,
                      "reste_cl": dispo_service - b * BOUTEILLE}
    return {
        "reports": reports,
        "exercice": ex, "jours": len(jour), "besoin_cl": besoin,
        "qte_lib": qte_lib, "scellees": scellees,
        "bt_annuel": bt_annuel, "bt_jour": bt_jour, "bt_libelle": bt_libelle,
        "dispo_cl": dispo, "dispo_service_cl": dispo_service,
        "sorti_cl": sorti, "jete_cl": sorti - besoin,
        "reste_cl": dispo_service - sorti,
        "jour": jour,
    }


def fr(x, dec=2):
    return f"{x:,.{dec}f}".replace(",", " ").replace(".", ",")


def ecrire_xlsx(res):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    blanc = Font(bold=True, color="FFFFFF")
    gras = Font(bold=True)
    tete = PatternFill("solid", fgColor="0F766E")
    surligne = PatternFill("solid", fgColor="CCE7E2")
    trait = Side(style="thin", color="D8DEE4")
    bord = Border(left=trait, right=trait, top=trait, bottom=trait)
    droite = Alignment(horizontal="right")

    def entetes(ws, cols):
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=ws.max_row, column=c)
            cell.font, cell.fill, cell.border = blanc, tete, bord
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)

    def largeurs(ws, l):
        for i, w in enumerate(l, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # --- 1. Le perimetre, et sa correspondance avec le tableau p. 62 --------
    ws = wb.active
    ws.title = "Perimetre p. 62"
    ws.append(["Le perimetre du service : les quatre libelles servis au verre "
               "du tableau « DONNEES ISSUES LOGICIEL DE CAISSE », page 62"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Doses de sa colonne « Dose (en centilitres) ». Quantites recomptees "
               "ligne a ligne dans l'annexe C (detail des tickets)."])
    ws.append([])
    entetes(ws, ["Exercice", "Libelle de caisse", "Dose retenue par le service (cl)",
                 "Quantite lue dans l'annexe C", "Quantite du tableau p. 62 (exercice 1)",
                 "Cremant (cl)"])
    for r in res:
        for lib, dose in PERIMETRE_P62.items():
            q = r["qte_lib"].get(lib, 0.0)
            ws.append([r["exercice"], lib, dose, round(q, 1),
                       CONTROLE_P62[lib] if r["exercice"] == "2022-2023" else "",
                       round(q * dose, 1)])
        ws.append([r["exercice"], "Bouteilles de 75 cl vendues scellees", 75,
                   round(r["scellees"], 1), 10 if r["exercice"] == "2022-2023" else "",
                   round(r["scellees"] * 75, 1)])
        ws.append([r["exercice"], "TOTAL servi au verre", "", "", "",
                   round(r["besoin_cl"], 1)])
        for c in range(1, 7):
            ws.cell(row=ws.max_row, column=c).font = gras
    largeurs(ws, [14, 40, 18, 18, 20, 14])

    # --- 2. Lecture annuelle contre lecture journaliere ---------------------
    ws2 = wb.create_sheet("Annuel contre journalier")
    ws2.append(["Le meme besoin, lu au niveau de l'annee puis au niveau de la journee"])
    ws2["A1"].font = Font(bold=True, size=12)
    ws2.append(["Aucun sur-versement : le serveur verse la dose exacte du service. "
                "La seule perte est le solde de la derniere bouteille ouverte."])
    ws2.append([])
    entetes(ws2, ["Exercice", "Jours ou l'un de ces articles est vendu", "Cremant servi (L)",
                  "Bouteilles au niveau annuel", "Bouteilles jour par jour",
                  "Bouteilles en plus", "Volume correspondant (L)"])
    t = [0, 0.0, 0, 0]
    for r in res:
        ws2.append([r["exercice"], r["jours"], round(r["besoin_cl"] / 100, 2),
                    r["bt_annuel"], r["bt_jour"], r["bt_jour"] - r["bt_annuel"],
                    round((r["bt_jour"] - r["bt_annuel"]) * 0.75, 2)])
        t[0] += r["jours"]; t[1] += r["besoin_cl"]
        t[2] += r["bt_annuel"]; t[3] += r["bt_jour"]
    ws2.append(["Total 3 exercices", t[0], round(t[1] / 100, 2), t[2], t[3],
                t[3] - t[2], round((t[3] - t[2]) * 0.75, 2)])
    for c in range(1, 8):
        ws2.cell(row=ws2.max_row, column=c).font = gras
        ws2.cell(row=ws2.max_row, column=c).fill = surligne
    largeurs(ws2, [18, 24, 18, 22, 22, 18, 22])

    # --- 3. Ce que vaut la mutualisation entre articles ---------------------
    ws3 = wb.create_sheet("Mutualisation")
    ws3.append(["Premiere objection du service (p. 62) : « le cremant est utilise "
                "pour de nombreux articles »"])
    ws3["A1"].font = Font(bold=True, size=12)
    ws3.append(["Elle est integree au calcul : une meme bouteille sert, dans la "
                "journee, aux quatre libelles. Colonne de gauche : ce que couterait "
                "une bouteille distincte par article et par jour."])
    ws3.append([])
    entetes(ws3, ["Exercice", "Une bouteille par article et par jour",
                  "Bouteille partagee dans la journee (retenu)",
                  "Bouteilles economisees", "Volume economise (L)"])
    u = [0, 0]
    for r in res:
        ws3.append([r["exercice"], r["bt_libelle"], r["bt_jour"],
                    r["bt_libelle"] - r["bt_jour"],
                    round((r["bt_libelle"] - r["bt_jour"]) * 0.75, 2)])
        u[0] += r["bt_libelle"]; u[1] += r["bt_jour"]
    ws3.append(["Total 3 exercices", u[0], u[1], u[0] - u[1],
                round((u[0] - u[1]) * 0.75, 2)])
    for c in range(1, 6):
        ws3.cell(row=ws3.max_row, column=c).font = gras
        ws3.cell(row=ws3.max_row, column=c).fill = surligne
    largeurs(ws3, [18, 34, 34, 22, 22])

    # --- 4. Le bilan matiere, dans le seul perimetre du service -------------
    ws4 = wb.create_sheet("Bilan dans son perimetre")
    ws4.append(["Bilan matiere du cremant, dans le seul perimetre de la page 62, "
                "sans aucun sur-versement"])
    ws4["A1"].font = Font(bold=True, size=12)
    ws4.append([])
    entetes(ws4, ["Exercice", "Disponible retenu par le service (L)",
                  "Bouteilles vendues scellees", "Disponible pour le service (L)",
                  "Bouteilles ouvertes", "Sorti du stock (L)", "Servi au verre (L)",
                  "Solde de bouteille perdu (L)", "Reste (L)"])
    s = [0.0] * 6
    for r in res:
        ws4.append([r["exercice"], round(r["dispo_cl"] / 100, 2), r["scellees"],
                    round(r["dispo_service_cl"] / 100, 2), r["bt_jour"],
                    round(r["sorti_cl"] / 100, 2), round(r["besoin_cl"] / 100, 2),
                    round(r["jete_cl"] / 100, 2), round(r["reste_cl"] / 100, 2)])
        s[0] += r["dispo_cl"]; s[1] += r["scellees"]; s[2] += r["dispo_service_cl"]
        s[3] += r["sorti_cl"]; s[4] += r["besoin_cl"]; s[5] += r["jete_cl"]
    ws4.append(["Total 3 exercices", round(s[0] / 100, 2), s[1], round(s[2] / 100, 2),
                sum(r["bt_jour"] for r in res), round(s[3] / 100, 2),
                round(s[4] / 100, 2), round(s[5] / 100, 2),
                round((s[2] - s[3]) / 100, 2)])
    for c in range(1, 10):
        ws4.cell(row=ws4.max_row, column=c).font = gras
        ws4.cell(row=ws4.max_row, column=c).fill = surligne
    largeurs(ws4, [18, 24, 20, 24, 18, 20, 20, 24, 14])

    # --- 5. Ce que devient la perte si la bouteille est conservee ----------
    ws5 = wb.create_sheet("Report d un jour a l autre")
    ws5.append(["Seconde objection du service (p. 62) : les dirigeants « sont susceptibles "
                "de savoir parfaitement conserver cette boisson d'un jour a l'autre »"])
    ws5["A1"].font = Font(bold=True, size=12)
    ws5.append(["Meme perimetre, memes doses, aucun sur-versement. Les journees de service "
                "sont groupees par blocs de k journees consecutives ; chaque bloc n'ouvre "
                "que les bouteilles necessaires a son besoin total. k = 1 : aucun report "
                "(cas retenu). k = 2 : la bouteille sert d'un jour a l'autre."])
    ws5.append([])
    entetes(ws5, ["Hypothese de conservation", "Exercice", "Bouteilles ouvertes",
                  "Sorti du stock (L)", "Servi au verre (L)",
                  "Solde de bouteille perdu (L)", "Reste (L)"])
    for k in REPORTS:
        lab = {1: "Aucun report d'un service a l'autre (retenu)",
               2: "Report d'un jour a l'autre"}.get(k, f"Report sur {k} journees de service")
        tj = tr = 0.0
        for r in res:
            x = r["reports"][k]
            ws5.append([lab, r["exercice"], x["bouteilles"],
                        round(x["sorti_cl"] / 100, 2), round(r["besoin_cl"] / 100, 2),
                        round(x["jete_cl"] / 100, 2), round(x["reste_cl"] / 100, 2)])
            tj += x["jete_cl"]; tr += x["reste_cl"]
        ws5.append([lab, "Total 3 exercices",
                    sum(r["reports"][k]["bouteilles"] for r in res),
                    round(sum(r["reports"][k]["sorti_cl"] for r in res) / 100, 2),
                    round(sum(r["besoin_cl"] for r in res) / 100, 2),
                    round(tj / 100, 2), round(tr / 100, 2)])
        for c in range(1, 8):
            ws5.cell(row=ws5.max_row, column=c).font = gras
            ws5.cell(row=ws5.max_row, column=c).fill = surligne
    largeurs(ws5, [42, 18, 20, 20, 20, 26, 14])

    for w in (ws, ws2, ws3, ws4, ws5):
        for row in w.iter_rows(min_row=4):
            for cell in row:
                cell.border = bord
                if isinstance(cell.value, (int, float)):
                    cell.alignment = droite

    os.makedirs(PIECES, exist_ok=True)
    wb.save(PIECES + "R1-cremant-perimetre-du-service.xlsx")


def ecrire_csv(res):
    with open(PIECES + "R1-cremant-perimetre-du-service.csv", "w",
              newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Cremant : le perimetre de la page 62 lu jour par jour"])
        w.writerow(["Libelles et doses du tableau « DONNEES ISSUES LOGICIEL DE CAISSE » "
                    "(reponse du 04/09/2026, p. 62) : La Vouivre 6 cl, Pere Gregoire 6 cl, "
                    "KITTYKIR 6 cl, Cremant au verre 12 cl. Bouteille = 75 cl. "
                    "Aucun sur-versement."])
        w.writerow(["Source : ANNEXE-C, detail des tickets (lecture seule)."])
        w.writerow([])
        w.writerow(["Exercice", "Date", "Cremant servi (cl)",
                    "Bouteilles a ouvrir ce jour", "Solde de bouteille perdu (cl)"])
        for r in res:
            for d in sorted(r["jour"]):
                cl = r["jour"][d]
                b = math.ceil(cl / BOUTEILLE - 1e-9)
                w.writerow([r["exercice"], d, f"{cl:.1f}".replace(".", ","), b,
                            f"{b * BOUTEILLE - cl:.1f}".replace(".", ",")])
            w.writerow([r["exercice"], "TOTAL", f"{r['besoin_cl']:.1f}".replace(".", ","),
                        r["bt_jour"], f"{r['jete_cl']:.1f}".replace(".", ",")])


def main():
    res = [calcule(ex) for ex in EXOS]
    ecrire_xlsx(res)
    ecrire_csv(res)
    tot = collections.Counter()
    for r in res:
        print(f"\n=== {r['exercice']} : {r['jours']} jours ou l'un de ces quatre articles est vendu")
        print(f"  servi au verre (perimetre p. 62)  : {fr(r['besoin_cl']/100)} L")
        print(f"  lecture annuelle                  : {r['bt_annuel']} bouteilles")
        print(f"  lecture jour par jour             : {r['bt_jour']} bouteilles "
              f"(+{r['bt_jour']-r['bt_annuel']} = {fr((r['bt_jour']-r['bt_annuel'])*0.75)} L)")
        print(f"  une bouteille par article et jour : {r['bt_libelle']} bouteilles "
              f"(la mutualisation en economise {r['bt_libelle']-r['bt_jour']})")
        print(f"  sorti du stock                    : {fr(r['sorti_cl']/100)} L")
        print(f"  solde de bouteille perdu          : {fr(r['jete_cl']/100)} L "
              f"({fr(100*r['jete_cl']/r['dispo_cl'],1)} % du disponible retenu)")
        print(f"  disponible pour le service        : {fr(r['dispo_service_cl']/100)} L")
        print(f"  reste                             : {fr(r['reste_cl']/100)} L")
        for k in ("besoin_cl", "sorti_cl", "jete_cl", "reste_cl", "dispo_service_cl"):
            tot[k] += r[k]
        for k in ("bt_annuel", "bt_jour", "bt_libelle", "jours"):
            tot[k] += r[k]
    print(f"\n=== TOTAL 3 exercices ({tot['jours']} jours concernes)")
    print(f"  servi {fr(tot['besoin_cl']/100)} L | annuel {tot['bt_annuel']} bt | "
          f"jour par jour {tot['bt_jour']} bt | ecart {tot['bt_jour']-tot['bt_annuel']} bt "
          f"= {fr((tot['bt_jour']-tot['bt_annuel'])*0.75)} L")
    print(f"  mutualisation : {tot['bt_libelle']-tot['bt_jour']} bouteilles economisees "
          f"= {fr((tot['bt_libelle']-tot['bt_jour'])*0.75)} L")
    print(f"  perdu en solde de bouteille : {fr(tot['jete_cl']/100)} L | "
          f"reste {fr(tot['reste_cl']/100)} L")
    print("\n  si une bouteille ouverte sert plusieurs journees de service :")
    for k in REPORTS:
        j = sum(r["reports"][k]["jete_cl"] for r in res)
        rs = sum(r["reports"][k]["reste_cl"] for r in res)
        det = " / ".join(fr(r["reports"][k]["jete_cl"] / 100) for r in res)
        print(f"     k={k} : jete {fr(j/100):>8} L  ({det})  reste {fr(rs/100):>8} L")
    print("\nPiece :", PIECES + "R1-cremant-perimetre-du-service.xlsx")


if __name__ == "__main__":
    main()
