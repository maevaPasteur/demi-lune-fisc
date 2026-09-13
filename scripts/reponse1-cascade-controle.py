#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REPONSE 1 - Controle de coherence de la cascade des 10 622 L.

Verifie, en lecture seule, que la cascade produite par
scripts/reponse1-cascade-valeurs.py boucle exactement et que les pieces
publiees portent les memes valeurs. A relancer apres toute modification d'un
poste : si un seul controle echoue, le dossier se contredit.

Chaine de production a relancer dans cet ordre :
  1. python3 scripts/reponse1-cascade-valeurs.py     (producteur unique)
  2. python3 scripts/reponse1-abattements.py         (R1-abattements-double-emploi.xlsx)
  3. python3 scripts/reponse1-coefficients.py        (R1-coefficients-et-extrapolation.xlsx)
  4. python3 scripts/reponse1-biere-freinte.py       (R1-biere-freinte.xlsx)
  5. python3 scripts/reponse1-cascade-controle.py    (ce script)
Le patch des PAGES (scripts/reponse1-cascade-coherence.py) est a lancer a part :
il ecrit dans src/data/reponse1/.

Execution : python3 scripts/reponse1-cascade-controle.py
Sortie : aucun fichier. Code de retour 1 si un controle echoue.
"""
import json
import os
import sys

ICI = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(ICI, ".."))
CASC = os.path.join(ROOT, "src/data/reponse1Calculs/cascade-10622.json")
PIECES = os.path.join(ROOT, "public/documents/pieces-reponse-1")

RESULTATS = []


def chk(libelle, condition):
    RESULTATS.append((libelle, bool(condition)))
    print(("  OK   " if condition else "  ECHEC ") + libelle)


def main():
    if not os.path.exists(CASC):
        raise SystemExit("Lancer d'abord : python3 scripts/reponse1-cascade-valeurs.py")
    d = json.load(open(CASC, encoding="utf-8"))
    P = {p["cle"]: p for p in d["postes"]}
    A = {a["cle"]: a for a in d["assiettes"]}
    B = d["doubles_comptages"]["biere"]
    O = d["doubles_comptages"]["offerts"]

    print("1. La cascade boucle")
    chk("somme des postes exacts + résidu = 10 622,00 L",
        abs(d["attribue_l"] + d["residu_l"] - 10622.0) < 0.005)
    chk("somme des postes arrondis + résidu arrondi = 10 622 L",
        sum(p["litres_arrondi"] for p in d["postes"]) + d["residu_arrondi_l"] == 10622)
    chk("aucun poste négatif", all(p["litres"] >= 0 for p in d["postes"]) and d["residu_l"] >= 0)

    print("2. Les assiettes s'additionnent")
    chk("A + cuisine et menus = B",
        A["A"]["litres_arrondi"] + P["cuisine"]["litres_arrondi"] == A["B"]["litres_arrondi"])
    chk("A + résidu = C", A["A"]["litres_arrondi"] + d["residu_arrondi_l"] == A["C"]["litres_arrondi"])
    chk("B + résidu = D", A["B"]["litres_arrondi"] + d["residu_arrondi_l"] == A["D"]["litres_arrondi"])
    chk("assiette A = somme de l'itémisation, en litres",
        abs(A["A"]["litres"] - sum(x["litres"] for x in d["itemisation"])) < 0.005)
    chk("assiette A = somme de l'itémisation, en euros",
        abs(A["A"]["euros"] - sum(x["euros"] for x in d["itemisation"])) < 0.05)
    chk("l'itémisation ne reprend que des postes de la cascade, une seule fois",
        abs(sum(x["litres"] for x in d["itemisation"])
            - sum(P[k]["litres"] for k in ("cremant", "surversement", "degustation",
                                           "biere", "chef", "offerts"))) < 0.005)
    chk("itémisation par exercice cohérente avec les totaux",
        all(abs(sum(x["litres_par_exercice"].values()) - x["litres"]) < 0.02
            for x in d["itemisation"]))

    print("3. Les deux doubles comptages sont retirés une seule fois")
    chk("vendu au verre corrigé = vendu au verre brut moins la contenance des mixtes",
        abs(B["verre_apres"] - (B["verre_avant"] - B["litres"])) < 0.02)
    chk("contenance des articles mixtes = 296,0 L", abs(B["litres"] - 296.0) < 0.05)
    chk("bière servie = pression et pintes pures + bière des quatre mixtes",
        abs(B["biere_servie_l"] - (B["biere_pression_et_pintes_l"]
                                   + B["alcool_reel"]["Fût Affligem"])) < 0.05)
    chk("freinte = 10 % de la bière réellement sortie du fût",
        abs(P["biere"]["litres"] - 0.10 * B["biere_servie_l"]) < 0.005)
    chk("poste offerts = hypothèse journalière moins les offerts déjà enregistrés",
        abs(P["offerts"]["litres"] - (O["poste_avant"] - O["litres_nets"])) < 0.005)
    chk("les offerts nets excluent la contenance des mixtes déjà retirée",
        abs(O["litres_nets"] - (O["litres_bruts"] - O["dont_articles_mixtes"])) < 0.02)

    print("4. Le forfait du service")
    chk("forfait en volume = 15 % des achats", abs(d["forfait_15pct"]["litres"] - 0.15 * 10622) < 0.01)
    chk("forfait en euros = 3 x (8 917,07 + 8 450,72 + 8 253,27) = 76 863,18 €",
        abs(d["forfait_15pct"]["euros_ca_boissons"] - 76863.18) < 0.005)

    print("5. Les pièces publiées portent les mêmes valeurs")
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("  (openpyxl absent : contrôle des pièces sauté)")
        load_workbook = None
    if load_workbook:
        ws = load_workbook(os.path.join(PIECES, "R1-cascade-bilan-matiere.xlsx"),
                           data_only=True)["1. Cascade"]
        lig = {str(r[0]): r for r in ws.iter_rows(values_only=True) if r and r[0]}
        chk("R1-cascade-bilan-matiere.xlsx : total attribué",
            lig["Total attribue a un poste identifie"][1] == d["attribue_arrondi_l"])
        chk("R1-cascade-bilan-matiere.xlsx : résidu",
            lig["Perte pure (casse, evaporation, fonds de verre, rincage)"][1]
            == d["residu_arrondi_l"])

        ws = load_workbook(os.path.join(PIECES, "R1-abattements-double-emploi.xlsx"),
                           data_only=True)["2. Itemisation par exercice"]
        tot = [r for r in ws.iter_rows(values_only=True) if r and str(r[0]).startswith("TOTAL")][0]
        chk("R1-abattements-double-emploi.xlsx : total itémisé en litres",
            abs(tot[4] - A["A"]["litres"]) < 0.06)
        chk("R1-abattements-double-emploi.xlsx : total itémisé en euros",
            abs(tot[8] - A["A"]["euros"]) < 2)

        wb = load_workbook(os.path.join(PIECES, "R1-biere-freinte.xlsx"), data_only=True)
        trouve = any(isinstance(c, (int, float)) and abs(c - P["biere"]["litres"]) < 0.01
                     for w in wb for r in w.iter_rows(values_only=True) for c in r)
        chk("R1-biere-freinte.xlsx : la freinte retenue est celle de la cascade", trouve)

        ws = load_workbook(os.path.join(PIECES, "R1-alcool-cuisine-controles.xlsx"),
                           data_only=True)["5-Solde demandé"]
        tot = [r for r in ws.iter_rows(values_only=True) if r and str(r[0]).strip() == "TOTAL"][0]
        chk("R1-alcool-cuisine-controles.xlsx : cuisine plafonnée = poste de la cascade",
            abs(float(tot[2]) - P["cuisine"]["litres"]) < 0.02)

    print()
    passes = sum(1 for _, c in RESULTATS if c)
    print("%d / %d contrôles passés" % (passes, len(RESULTATS)))
    if passes != len(RESULTATS):
        print("INCOHERENCE : ne rien publier avant de l'avoir levée.")
        sys.exit(1)


if __name__ == "__main__":
    main()
