#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
R1 - Alcool de cuisine : l'origine des doses et la composition des reperes D / E / G / Q.

Produit public/documents/pieces-reponse-1/R1-alcool-cuisine-reperes.xlsx :
  1. Le courriel des dirigeants du 30/03/2026, tel que la proposition du 18/05/2026 le reproduit (p. 37)
  2. Les douze reperes detailles par la proposition (p. 44, 45, 48) et les annexes n 7-1, 7-2 et 7-3
  3. Colonnes du courrier du 04/09/2026 : somme brute des articles contre volume calcule sur le net
  4. Le Feuillete Forestier vendu a la carte, ecarte du repere G
  5. Effet des bornes de la fourchette declaree (Calvados 4/5 cl, macvin 5/4 cl, fondues 10/9 cl)

Les feuilles 1 a 3 sont des transcriptions des documents du service, collationnees sur les images.
Les feuilles 4 et 5 sont recalculees a chaque execution depuis les fichiers de caisse.

Sources (lecture seule) :
  public/documents/caisse-enregistreuse/ANNEXE-D1..D3_synthese-produit_*.xls   (ventes a la carte)
  public/documents/caisse-enregistreuse/ANNEXE-C1..C3_detail-tickets_*.xls     (menus vendus)
  public/documents/vins-boissons/_pipeline_conso_alcool.py                      (composition des menus)
  src/data/calculsBoissons/consoAlcoolMenus.json                                (controle de fidelite)
"""
import ast, collections, json, os
import xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__))) + "/"
CA = BASE + "public/documents/caisse-enregistreuse/"
OUT = BASE + "public/documents/pieces-reponse-1/"
EXOS = ["2022-2023", "2023-2024", "2024-2025"]
LAB = {"2022-2023": "Exercice 1", "2023-2024": "Exercice 2", "2024-2025": "Exercice 3"}

# ==========================================================================
# 1. Le courriel des dirigeants du 30/03/2026 (proposition p. 37)
# ==========================================================================
COURRIEL = [
    ("DESSERTS", "Crème brûlée à l'Absinthe : 8 cl/Litre pour 7 pers. = 1.14 cl d'Absinthe/Pers.",
     "dose unique", "1,14 cl", "Absinthe"),
    ("DESSERTS", "Babas : 4 cl d'alcool (Quelque soit l'alcool)",
     "dose unique", "4 cl", "macvin, marc, liqueur de poire"),
    ("DESSERTS", "Coupes de glaces alcoolisées : 4 cl d'alcool (Quelque soit l'alcool)",
     "dose unique", "4 cl", "crème de cassis, Bailey's"),
    ("DESSERTS", "Crêpes flambées : 4 cl d'alcool (Quelque soit l'alcool)",
     "dose unique", "4 cl", "Grand Marnier, Calvados"),
    ("PLATS", "Sauce Jurassienne : 10 cl de vin jaune pour 1 litre de sauce pour 10 parts = 1 cl de vin jaune/Pers.",
     "dose unique", "1 cl", "vin jaune"),
    ("PLATS", "Sauce Forestière et Sauce Morilles : 10 cl de Porto pour 1 litre de sauce pour 10 parts = 1 cl de Porto/Pers.",
     "dose unique", "1 cl", "porto"),
    ("PLATS", "Camembert, Assiette Franc-Comtoise et Assiette du Père Grégoire : Ce sont des plats flambés "
              "devant les clients, environ 4 à 5 cl d'alcool/Pers. (Macvin ou Calvados)",
     "FOURCHETTE", "5 cl pour le macvin (repère Q), 4 cl pour le Calvados (repère D)", "macvin, Calvados"),
    ("FONDUES", "Fondue Jurassienne : 9 ou10 cl de vin/Pers.",
     "FOURCHETTE", "10 cl", "ravelin"),
    ("FONDUES", "Fondue Vin Jaune et Fondue des Gourmets : 9 ou 10 cl de vin jaune/Pers.",
     "FOURCHETTE", "10 cl", "vin jaune"),
]

# ==========================================================================
# 2. Les douze reperes, les trois exercices
#    Transcription des tableaux de la proposition (p. 44, 45, 48) et des annexes n 7-1 (p. 171
#    a 173), n 7-2 (p. 185 a 188) et n 7-3 (p. 199 a 202), collationnee sur les images.
# ==========================================================================
#    (repere, alcool, exercice, source, article, brut, proportion, net, dose, quantite)
REPERES = [
    ("D", "Calvados", 1, "proposition p. 44, annexe 7-1 p. 171", "Camenbert Rôti", 512, "", 512, 4, 2048),
    ("D", "Calvados", 1, "proposition p. 44, annexe 7-1 p. 171", "Assiette du Père Grégoire", 222, "", 222, 4, 888),
    ("D", "Calvados", 1, "proposition p. 44, annexe 7-1 p. 171", "Menu Végétarien", 263, "35,58 %", 94, 4, 374),
    ("D", "Calvados", 2, "annexe 7-2 p. 185", "Camembert Rôti", 475, "", 475, 4, 1900),
    ("D", "Calvados", 2, "annexe 7-2 p. 185", "Assiette du Père Grégoire", 210, "", 210, 4, 840),
    ("D", "Calvados", 2, "annexe 7-2 p. 185", "Menu Végétarien", 195, "33,60 %", 66, 4, 262),
    ("D", "Calvados", 3, "annexe 7-3 p. 199", "Camembert Rôti", 465, "", 465, 4, 1860),
    ("D", "Calvados", 3, "annexe 7-3 p. 199", "Assiette du Père Grégoire", 216, "", 216, 4, 864),
    ("D", "Calvados", 3, "annexe 7-3 p. 199", "Menu Végétarien", 192, "36,61 %", 70, 4, 281),
    ("E", "Vin jaune", 1, "proposition p. 45, annexe 7-1 p. 171", "Filet Poulet Jurassienne", 955, "", 955, 1, 955),
    ("E", "Vin jaune", 1, "proposition p. 45, annexe 7-1 p. 171", "Sauce Vin jaune", 91, "", 91, 1, 91),
    ("E", "Vin jaune", 1, "proposition p. 45, annexe 7-1 p. 171", "Fondue Vin jaune", 133, "", 133, 10, 1330),
    ("E", "Vin jaune", 1, "proposition p. 45, annexe 7-1 p. 171", "Fondue des Gourmet", 88, "", 88, 10, 880),
    ("E", "Vin jaune", 1, "proposition p. 45, annexe 7-1 p. 171", "Menu Franc Comtois", 748, "61,65 %", 461, 1, 461),
    ("E", "Vin jaune", 2, "annexe 7-2 p. 186", "Filet Poulet Jurassienne", 937, "", 955, 1, 955),
    ("E", "Vin jaune", 2, "annexe 7-2 p. 186", "Sauce Vin jaune", 75, "", 91, 1, 91),
    ("E", "Vin jaune", 2, "annexe 7-2 p. 186", "Fondue Vin jaune", 160, "", 133, 10, 1330),
    ("E", "Vin jaune", 2, "annexe 7-2 p. 186", "Fondue des Gourmet", 87, "", 88, 10, 880),
    ("E", "Vin jaune", 2, "annexe 7-2 p. 186", "Menu Franc Comtois", 721, "60,65 %", 437, 1, 437),
    ("E", "Vin jaune", 3, "annexe 7-3 p. 200", "Filet Poulet Jurassienne", 865, "", 955, 1, 955),
    ("E", "Vin jaune", 3, "annexe 7-3 p. 200", "Sauce Vin jaune", 80, "", 91, 1, 91),
    ("E", "Vin jaune", 3, "annexe 7-3 p. 200", "Fondue Vin jaune", 150, "", 133, 10, 1330),
    ("E", "Vin jaune", 3, "annexe 7-3 p. 200", "Fondue des Gourmet", 39, "", 88, 10, 880),
    ("E", "Vin jaune", 3, "annexe 7-3 p. 200", "Menu Franc Comtois", 803, "61,26 %", 492, 1, 492),
    ("G", "Porto", 1, "proposition p. 45, annexe 7-1 p. 172", "Sauce Forestière", 33, "", 33, 1, 33),
    ("G", "Porto", 1, "proposition p. 45, annexe 7-1 p. 172", "Entrecôte Forestière", 238, "", 238, 1, 238),
    ("G", "Porto", 1, "proposition p. 45, annexe 7-1 p. 172", "Sauce Morille", 37, "", 37, 1, 37),
    ("G", "Porto", 1, "proposition p. 45, annexe 7-1 p. 172", "Entrecôte Morille", 17, "", 17, 1, 17),
    ("G", "Porto", 1, "proposition p. 45, annexe 7-1 p. 172", "Menu Demi-Lune", 153, "47,83 %", 73, 1, 73),
    ("G", "Porto", 2, "annexe 7-2 p. 186", "Sauce Forestière", 50, "", 50, 1, 50),
    ("G", "Porto", 2, "annexe 7-2 p. 186", "Entrecôte Forestière", 226, "", 226, 1, 226),
    ("G", "Porto", 2, "annexe 7-2 p. 186", "Sauce Morille", 25, "", 25, 1, 25),
    ("G", "Porto", 2, "annexe 7-2 p. 186", "Entrecôte Morille", 136, "", 136, 1, 136),
    ("G", "Porto", 2, "annexe 7-2 p. 186", "Menu Demi-Lune", 308, "47,25 %", 146, 1, 146),
    ("G", "Porto", 3, "annexe 7-3 p. 200", "Sauce Forestière", 39, "", 39, 1, 39),
    ("G", "Porto", 3, "annexe 7-3 p. 200", "Entrecôte Forestière", 200, "", 200, 1, 200),
    ("G", "Porto", 3, "annexe 7-3 p. 200", "Sauce Morille", 33, "", 33, 1, 33),
    ("G", "Porto", 3, "annexe 7-3 p. 200", "Entrecôte Morille", 132, "", 132, 1, 132),
    ("G", "Porto", 3, "annexe 7-3 p. 200", "Menu Demi-Lune", 379, "40,06 %", 152, 1, 152),
    ("Q", "Macvin", 1, "proposition p. 48, annexe 7-1 p. 173", "Baba au Macvin", 310, "", 310, 5, 1550),
    ("Q", "Macvin", 1, "proposition p. 48, annexe 7-1 p. 173", "Saucisse flambée au Macvin", 594, "", 594, 5, 2970),
    ("Q", "Macvin", 1, "proposition p. 48, annexe 7-1 p. 173", "Menu Franc Comtois, plat (Poulet/Assiette)", 748, "38,35 %", 287, 5, 1434),
    ("Q", "Macvin", 1, "proposition p. 48, annexe 7-1 p. 173", "Menu Franc Comtois, dessert (Baba / Crème brûlée)", 748, "48,67 %", 364, 5, 1820),
    ("Q", "Macvin", 2, "annexe 7-2 p. 188", "Baba au Macvin", 460, "", 460, 5, 2300),
    ("Q", "Macvin", 2, "annexe 7-2 p. 188", "Saucisse flambée au Macvin", 608, "", 608, 5, 3040),
    ("Q", "Macvin", 2, "annexe 7-2 p. 188", "Menu Franc Comtois, plat (Poulet/Assiette)", 721, "39,35 %", 284, 5, 1419),
    ("Q", "Macvin", 2, "annexe 7-2 p. 188", "Menu Franc Comtois, dessert (Baba / Crème brûlée)", 721, "57,50 %", 415, 5, 2073),
    ("Q", "Macvin", 3, "annexe 7-3 p. 202", "Baba au Macvin", 561, "", 561, 4, 2244),
    ("Q", "Macvin", 3, "annexe 7-3 p. 202", "Saucisse flambée au Macvin", 547, "", 547, 4, 2188),
    ("Q", "Macvin", 3, "annexe 7-3 p. 202", "Menu Franc Comtois, plat (Poulet/Assiette)", 803, "38,74 %", 311, 4, 1244),
    ("Q", "Macvin", 3, "annexe 7-3 p. 202", "Menu Franc Comtois, dessert (Baba / Crème brûlée)", 803, "53,84 %", 432, 4, 1729),
]

# colonnes du courrier du 04/09/2026 (p. 65 a 67) : (page, dose annoncee, nb de plats, volume)
COURRIER = {
    ("D", 1): ("p. 65", "4 cl", "997", 3310),
    ("D", 2): ("p. 65", "4 cl", "880", 3002),
    ("D", 3): ("p. 65", "4 cl", "873", 3005),
    ("E", 1): ("p. 66", "1 cl / 10 cl", "1 507 / 221", 3717),
    ("E", 2): ("p. 66", "1 cl / 10 cl", "1 449 / 247", 3693),
    ("E", 3): ("p. 66", "1 cl / 10 cl", "1 748 / 189", 3748),
    ("G", 1): ("p. 67", "1 cl", "478", 398),
    ("G", 2): ("p. 67", "1 cl", "745", 583),
    ("G", 3): ("p. 67", "1 cl", "783", 556),
    ("Q", 1): ("p. 67", "4 cl", "2 400", 7774),
    ("Q", 2): ("p. 67", "4 cl", "2 510", 8832),
    ("Q", 3): ("p. 67", "4 cl", "2 714", 7406),
}

# ==========================================================================
# Lecture des donnees de caisse
# ==========================================================================
D_FILES = {e: CA + f"ANNEXE-D{i}_synthese-produit_{e}.xls" for i, e in enumerate(EXOS, 1)}
C_FILES = {e: CA + f"ANNEXE-C{i}_detail-tickets_{e}.xls" for i, e in enumerate(EXOS, 1)}

carte = {e: collections.defaultdict(float) for e in EXOS}
for e, f in D_FILES.items():
    sh = xlrd.open_workbook(f).sheet_by_index(0)
    for r in range(sh.nrows):
        lib = str(sh.cell_value(r, 1)).replace("\xa0", " ").strip()
        q = sh.cell_value(r, 4)
        if lib and isinstance(q, float):
            carte[e][lib] += q

# composition des menus : reprise du pipeline du dossier
src = open(BASE + "public/documents/vins-boissons/_pipeline_conso_alcool.py", encoding="utf-8").read()
tree = ast.parse(src)
G = {}
for node in tree.body:
    if isinstance(node, ast.Assign):
        nom = getattr(node.targets[0], "id", "")
        if nom in ("MENUS", "MENU_PRIX", "MENU_COMPO", "PERIODES"):
            try:
                G[nom] = ast.literal_eval(node.value)
            except Exception:
                pass

ACCEPT = {"Menu 'Demi Lune'": {45.0}, "Menu du Dahu": {29.9},
          "Menu Végétarien": {23.0, 25.0, 26.5},
          "Menu Bourguignon": {32.0, 34.0}, "Menu Franc-Comtois": {32.0, 34.0}}
KEY = {"Menu 'Demi Lune'": "Demi Lune", "Menu du Dahu": "Dahu", "Menu Végétarien": "Végétarien",
       "Menu Bourguignon": "Bourguignon", "Menu Franc-Comtois": "Franc-Comtois"}
# (composant, alcool, dose, proba basse, proba haute), par periode de carte
COMP = {
 "A": {"Demi Lune": [("Poulet aux morilles", "Porto", 1, 1, 1), ("Feuillete forestier", "Porto", 1, .33, .60),
                     ("Dessert Baileys", "Baileys", 4, .15, .35), ("Dessert liqueur de poire", "Liqueur de poire", 4, .15, .35)],
       "Dahu": [("Fondue jurassienne", "Ravelin", 10, 1, 1)],
       "Végétarien": [("Assiette pere Gregoire", "Calvados", 4, .35, .65), ("Galette forestiere", "Porto", 1, .35, .65)],
       "Bourguignon": [("Coupe dijonnaise", "Crème de cassis", 4, .35, .65), ("Dessert Baileys", "Baileys", 4, .35, .65)],
       "Franc-Comtois": [("Poulet jurassienne", "Vin jaune", 1, .35, .65), ("Assiette franc-comtoise", "Macvin", 4, .35, .65)]},
 "B": {"Demi Lune": [("Poulet aux morilles", "Porto", 1, 1, 1), ("Feuillete forestier", "Porto", 1, .33, .60)],
       "Dahu": [("Fondue jurassienne", "Ravelin", 10, 1, 1)],
       "Végétarien": [("Assiette pere Gregoire", "Calvados", 4, .35, .65), ("Galette forestiere", "Porto", 1, .35, .65)],
       "Bourguignon": [("Coupe dijonnaise", "Crème de cassis", 4, .35, .65), ("Baba marc de bourgogne", "Marc de bourgogne", 4, .35, .65)],
       "Franc-Comtois": [("Poulet jurassienne", "Vin jaune", 1, .35, .65), ("Assiette franc-comtoise", "Macvin", 4, .35, .65),
                         ("Baba au macvin", "Macvin", 4, .20, .45)]},
 "C": {"Demi Lune": [("Poulet aux morilles", "Porto", 1, 1, 1), ("Feuillete forestier", "Porto", 1, .40, .65)],
       "Dahu": [("Fondue jurassienne", "Ravelin", 10, 1, 1)],
       "Végétarien": [("Assiette pere Gregoire", "Calvados", 4, .35, .65), ("Galette forestiere", "Porto", 1, .35, .65)],
       "Bourguignon": [("Coupe dijonnaise", "Crème de cassis", 4, .35, .65), ("Baba marc de bourgogne", "Marc de bourgogne", 4, .35, .65)],
       "Franc-Comtois": [("Poulet jurassienne", "Vin jaune", 1, .35, .65), ("Assiette franc-comtoise", "Macvin", 4, .35, .65),
                         ("Baba au macvin", "Macvin", 4, .20, .45)]}}


def periode(d):
    return "C" if d < "2022-07-23" else ("B" if d < "2023-04-17" else "A")


comp_port = {e: collections.defaultdict(float) for e in EXOS}
comp_cl = {e: collections.defaultdict(float) for e in EXOS}
for e, f in C_FILES.items():
    sh = xlrd.open_workbook(f).sheet_by_index(0)
    for r in range(1, sh.nrows):
        lib = str(sh.cell_value(r, 10)).strip()
        if lib not in ACCEPT:
            continue
        prix = round(float(sh.cell_value(r, 13)), 2)
        if prix not in ACCEPT[lib]:
            continue
        q = float(sh.cell_value(r, 11))
        p = periode(str(sh.cell_value(r, 0))[:10])
        for (nom, alc, dose, pb, ph) in COMP[p][KEY[lib]]:
            pm = (pb + ph) / 2
            comp_port[e][(nom, alc)] += q * pm
            comp_cl[e][(nom, alc)] += q * dose * pm

# controle de fidelite : les menus recalcules doivent reproduire consoAlcoolMenus.json
ref = json.load(open(BASE + "src/data/calculsBoissons/consoAlcoolMenus.json", encoding="utf-8"))
ctrl_ok, ctrl_detail = True, []
for e in EXOS:
    agg = collections.defaultdict(float)
    for (nom, alc), v in comp_cl[e].items():
        agg[alc] += v
    r = {x["alcool"]: x["cl_moyen"] for x in ref["exercices"][e]["alcools"]}
    for alc, v in agg.items():
        att = r.get(alc)
        if att is None or abs(att - v) > 0.5:
            ctrl_ok = False
        ctrl_detail.append((e, alc, round(v, 1), att))

# ==========================================================================
# 4. Feuillete Forestier vendu a la carte
# ==========================================================================
LIB_FEUILLETE = "Feuilleté Forestier"
LIB_CHEVRE = "Chèvre  chaud"
feuillete = {e: carte[e][LIB_FEUILLETE] for e in EXOS}
chevre_ex1 = carte[EXOS[0]][LIB_CHEVRE]
prop_service = 170.50 / (170.50 + 186.00) * 100

# ==========================================================================
# 5. Bornes de la fourchette declaree
# ==========================================================================
BORNES = [
    ("Calvados, plats flambés", ["CAMEMBERT ROTI", "ASSIETTE PERE GREGOI"], ["Assiette pere Gregoire"],
     4, 5, "« environ 4 à 5 cl » (courriel du 30/03/2026) ; le service retient 4 cl (repère D)"),
    ("Macvin, plats flambés", ["ASSIETTE FRANC-COMTO"], ["Assiette franc-comtoise"],
     5, 4, "« environ 4 à 5 cl » (courriel du 30/03/2026) ; le service retient 5 cl (repère Q)"),
    ("Ravelin, fondues", ["Fondue Jurassienne", "Fondue Forestière"], ["Fondue jurassienne"],
     10, 9, "« 9 ou 10 cl » (courriel du 30/03/2026) ; notre chiffrage retient 10 cl"),
    ("Vin jaune, fondues", ["FONDUE VIN JAUNE", "Fondue des Gourmets"], [],
     10, 9, "« 9 ou 10 cl » (courriel du 30/03/2026) ; le service applique 10 cl (repère E)"),
]

bornes_rows = []
for nom, libs, comps, d_ref, d_alt, obs in BORNES:
    ligne = {"nom": nom, "d_ref": d_ref, "d_alt": d_alt, "obs": obs,
             "carte": {}, "menus": {}, "tot": {}}
    for e in EXOS:
        qc = sum(carte[e][x] for x in libs)
        qm = sum(p for (n, a), p in comp_port[e].items() if n in comps)
        ligne["carte"][e], ligne["menus"][e], ligne["tot"][e] = qc, qm, qc + qm
    bornes_rows.append(ligne)

# ==========================================================================
# Ecriture du classeur
# ==========================================================================
TH = Font(bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F3A5F")
OKF = PatternFill("solid", fgColor="E3F5EA")
BADF = PatternFill("solid", fgColor="FBE4E4")
B = Border(*[Side(style="thin", color="D0D0D0")] * 4)


def sheet(wb, nom, titre, entetes, largeurs, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = nom
    ws.cell(1, 1, titre).font = Font(bold=True, size=12)
    for j, (h, w) in enumerate(zip(entetes, largeurs), 1):
        c = ws.cell(3, j, h)
        c.font, c.fill, c.border = TH, HF, B
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A4"
    return ws


def put(ws, r, vals, fmts=None, fill=None, bold=False):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.border = B
        if bold:
            c.font = Font(bold=True)
        if fmts and fmts[j - 1]:
            c.number_format = fmts[j - 1]
        if fill and fill(j):
            c.fill = fill(j)


wb = Workbook()

# --- feuille 1 : le courriel
ws = sheet(wb, "1-Courriel 30-03-2026",
           "1. Le courriel des dirigeants du 30/03/2026, tel que la proposition de rectifications du 18/05/2026 le reproduit page 37",
           ["Rubrique", "Texte du courriel (transcription littérale)", "Forme de la dose",
            "Dose retenue par le service", "Alcools concernés"],
           [12, 86, 14, 40, 30], first=True)
r = 4
for rub, txt, forme, retenue, alcools in COURRIEL:
    put(ws, r, [rub, txt, forme, retenue, alcools],
        fill=(lambda j: BADF if j == 3 else None) if forme == "FOURCHETTE" else (lambda j: OKF if j == 3 else None))
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(r, 4).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
ws.cell(r + 1, 1, "Six lignes sur neuf donnent une dose unique au centilitre. Trois donnent une fourchette, et ce sont celles des plats flambés devant le client et des fondues.")
ws.cell(r + 2, 1, "Proposition de rectifications du 18/05/2026, page 32 : « Pour ces deux alcools, les dirigeants ont indiqué que chaque portion, suivant le plat concerné, nécessitait 4 à 5 centilitres de ces alcools. "
                  "Le vérificateur a d'abord retenu 5 centilitres pour les deux. Puis, compte tenu de l'importance de l'incohérence, a retenu 5 centilitres pour le Macvin, et seulement 4 pour le Calvados. »")
ws.cell(r + 3, 1, "Même page : « En effet, en matière de calvados, il s'est heurté à une trop importante surconsommation de calvados en retenant 5 centilitres. Il a donc plutôt retenu 4 centilitres. "
                  "Mais même dans ces conditions, une très nette surconsommation de calvados, par rapport aux achats disponibles, demeurait encore, et ceci sur tous les exercices : "
                  "210 centilitres sur 2022-2023 (soit environ 8 % des achats disponibles), 802 centilitres sur 2023-2024 (soit environ 40 %), 1 005 centilitres sur 2024-2025 (soit environ 50 %). »")

# --- feuille 2 : les reperes des annexes n 7
ws = sheet(wb, "2-Repères annexes 7",
           "2. Les douze repères D, E, G et Q tels que la proposition du 18/05/2026 et ses annexes n° 7-1, 7-2 et 7-3 les détaillent",
           ["Repère", "Alcool", "Exercice", "Source", "Article", "Nb articles caisse (brut)",
            "Proportion appliquée", "Nb articles net", "Dose (cl)", "Quantité (cl)"],
           [8, 12, 10, 34, 46, 16, 14, 13, 9, 12])
r = 4
for rep in ["D", "E", "G", "Q"]:
    for ex in (1, 2, 3):
        lignes = [x for x in REPERES if x[0] == rep and x[2] == ex]
        for (_, alc, _e, source, art, brut, prop, net, dose, qte) in lignes:
            put(ws, r, [rep, alc, LAB[EXOS[ex - 1]], source, art, brut, prop, net, dose, qte],
                [None] * 5 + ["#,##0", None, "#,##0", "0", "#,##0"])
            r += 1
        sb, sn, sq = sum(x[5] for x in lignes), sum(x[7] for x in lignes), sum(x[9] for x in lignes)
        put(ws, r, [rep, lignes[0][1], LAB[EXOS[ex - 1]], "", "TOTAL du repère", sb, "", sn, "", sq],
            [None] * 5 + ["#,##0", None, "#,##0", None, "#,##0"], bold=True)
        r += 1
ws.cell(r + 1, 1, "Transcription collationnée sur les images, libellés du service compris (« Camenbert Rôti », « Fondue des Gourmet »). Les quantités sont celles imprimées : "
                  "le service calcule sur le net non arrondi (263 x 35,58 % = 93,57 d'où 374 cl, et 748 x 38,35 % = 286,86 d'où 1 434 cl).")
ws.cell(r + 2, 1, "Repère E, exercices 2 et 3 : la colonne « Nb articles net » des quatre articles vendus à la carte n'a pas été recalculée, elle reproduit les effectifs de l'exercice 1 (955, 91, 133, 88) "
                  "alors qu'aucune proportion ne s'y applique et que la colonne caisse porte 937, 75, 160, 87 puis 865, 80, 150, 39. Seule la ligne du menu Franc-Comtois est recalculée.")
ws.cell(r + 3, 1, "Repère G : la colonne de dose est intitulée « Vin Jaune En cl » sur les trois exercices, dans un tableau qui ne chiffre que du porto.")
ws.cell(r + 5, 1, "Les totaux de cette feuille sont la somme des quantités imprimées ligne à ligne : le repère Q de l'exercice 3 y ressort à 7 405 cl quand l'annexe imprime 7 406, le service calculant sur le net non arrondi.")
ws.cell(r + 4, 1, "Repère Q : la dose appliquée est de 5 cl sur les exercices 1 et 2, et de 4 cl sur l'exercice 3. Les trois annexes portent la mention « LE MACVIN ROSE N'A PAS ETE PRIS EN COMPTE ».")

# --- feuille 3 : brut contre net
ws = sheet(wb, "3-Brut contre net",
           "3. Ce que les colonnes du courrier du 04/09/2026 rapprochent : une somme brute d'articles et un volume calculé sur le net",
           ["Repère", "Alcool", "Exercice", "Page du courrier", "Dose annoncée en en-tête",
            "Nombre de plats imprimé par le courrier", "Somme brute des articles (annexes n° 7)",
            "Somme nette des articles (annexes n° 7)", "Volume imprimé par le courrier (cl)",
            "Total du repère (annexes n° 7, cl)", "Dose réellement appliquée",
            "Colonne « nombre de plats » du courrier"],
           [8, 12, 10, 13, 16, 18, 18, 18, 16, 18, 14, 22])
r = 4
for rep in ["D", "E", "G", "Q"]:
    for ex in (1, 2, 3):
        lignes = [x for x in REPERES if x[0] == rep and x[2] == ex]
        page, dose_ann, nb_courrier, vol_courrier = COURRIER[(rep, ex)]
        sb, sn, sq = sum(x[5] for x in lignes), sum(x[7] for x in lignes), sum(x[9] for x in lignes)
        doses = sorted({x[8] for x in lignes})
        if rep == "E":
            nature = {1: "net", 2: "net, effectifs réels", 3: "menus comptés en entier"}[ex]
        else:
            nature = "somme brute"
        ok = (rep != "E" and str(sb) == nb_courrier.replace(" ", ""))
        put(ws, r, [rep, lignes[0][1], LAB[EXOS[ex - 1]], page, dose_ann, nb_courrier, sb, sn,
                    vol_courrier, sq, " et ".join(f"{d} cl" for d in doses), nature],
            [None] * 6 + ["#,##0", "#,##0", "#,##0", "#,##0", None, None],
            fill=(lambda j: BADF if j == 12 else None) if ok else None)
        r += 1
ws.cell(r + 1, 1, "Lecture : pour le Calvados, le porto et le macvin, sur les trois exercices sans exception, le nombre de plats imprimé par le courrier (997, 880, 873 ; 478, 745, 783 ; 2 400, 2 510, 2 714) "
                  "est la somme BRUTE de la colonne « nombre d'articles caisse », menus comptés en entier, alors que le volume en regard est calculé sur la colonne « nombre d'articles NET ».")
ws.cell(r + 2, 1, "Aucun rapport entre les deux colonnes du courrier ne constitue donc une dose : 3 310 / 997 = 3,32 cl est le quotient d'un volume net par un effectif brut.")
ws.cell(r + 3, 1, "Vin jaune : refait sur les effectifs de caisse que les annexes impriment elles-mêmes, le repère E donnerait 3 717, puis 3 919, puis 3 327 cl, soit 109,63 L sur trois exercices, "
                  "contre 111,58 L retranchés. L'écart joue dans les deux sens et, au total, légèrement en faveur de la société.")
ws.cell(r + 4, 1, "Macvin : si l'exercice 3 avait été traité comme les deux précédents, à 5 cl, le repère Q y aurait retranché 9 255 cl au lieu de 7 406, soit 18,5 L de plus.")

# --- feuille 4 : feuillete forestier
ws = sheet(wb, "4-Feuilleté écarté",
           "4. Le Feuilleté Forestier vendu à la carte : compté par le service pour bâtir sa proportion, absent de son décompte de porto",
           ["Exercice", "Feuilleté Forestier vendu à la carte (portions)", "Porto à 1 cl (L)",
            "Retenu dans le repère G ?", "Source"],
           [14, 22, 14, 18, 46])
r = 4
for e in EXOS:
    put(ws, r, [LAB[e], round(feuillete[e], 2), round(feuillete[e] / 100, 3), "non",
                os.path.basename(D_FILES[e])],
        [None, "#,##0.00", "#,##0.000", None, None],
        fill=lambda j: BADF if j == 4 else None)
    r += 1
tot_f = sum(feuillete.values())
put(ws, r, ["3 exercices", round(tot_f, 2), round(tot_f / 100, 3), "non", ""],
    [None, "#,##0.00", "#,##0.000", None, None], bold=True)
r += 2
ws.cell(r, 1, "Contrôle de la clé du service : la proposition écrit, page 45, « 170,50 Feuilleté Forestier 47,83 % Entrée avec sauce Forestière » et « 186,00 Salade de Chèvre Chaud 52,17 % ».")
r += 1
ws.cell(r, 1, f"Feuilleté Forestier lu en caisse sur l'exercice 1 : {feuillete[EXOS[0]]:.2f} portions. "
              f"Chèvre chaud lu en caisse sur l'exercice 1 : {chevre_ex1:.2f} portions. "
              f"170,50 / (170,50 + 186,00) = {prop_service:.2f} %, soit exactement la proportion du service.")
r += 1
ws.cell(r, 1, "Le service utilise donc ce décompte pour répartir les menus Demi-Lune, et il écrit page 46 qu'il y a eu « 73 Feuilleté Forestier nécessitant la sauce Forestière préparée à base de Porto ». "
              "Mais la ligne « Feuilleté Forestier » vendu à la carte ne figure pas dans le tableau qui produit le total du repère G : seules y figurent Sauce Forestière, Entrecôte Forestière, Sauce Morille, Entrecôte Morille et Menu Demi-Lune.")
r += 1
ws.cell(r, 1, "Les feuilletés servis à l'intérieur des menus Demi-Lune sont comptés à part dans notre chiffrage, au titre du porto des menus, et ne sont pas inclus dans les portions ci-dessus.")

# --- feuille 5 : bornes de la fourchette
ws = sheet(wb, "5-Bornes de fourchette",
           "5. Effet des deux bornes de la fourchette déclarée par les dirigeants le 30/03/2026",
           ["Poste", "Exercice", "Portions à la carte", "Portions dans les menus (espérance)",
            "Portions totales", "Dose de référence (cl)", "Volume à la dose de référence (L)",
            "Dose alternative (cl)", "Volume à la dose alternative (L)", "Écart (L)", "Observation"],
           [24, 12, 14, 18, 13, 13, 16, 13, 16, 11, 62])
r = 4
for lg in bornes_rows:
    for e in EXOS:
        v1 = lg["tot"][e] * lg["d_ref"] / 100
        v2 = lg["tot"][e] * lg["d_alt"] / 100
        put(ws, r, [lg["nom"], LAB[e], round(lg["carte"][e], 2), round(lg["menus"][e], 2),
                    round(lg["tot"][e], 2), lg["d_ref"], round(v1, 2), lg["d_alt"], round(v2, 2),
                    round(v2 - v1, 2), ""],
            [None, None, "#,##0.00", "#,##0.00", "#,##0.00", "0", "#,##0.00", "0", "#,##0.00", "#,##0.00", None])
        r += 1
    t = sum(lg["tot"].values())
    put(ws, r, [lg["nom"], "3 exercices", round(sum(lg["carte"].values()), 2),
                round(sum(lg["menus"].values()), 2), round(t, 2), lg["d_ref"],
                round(t * lg["d_ref"] / 100, 2), lg["d_alt"], round(t * lg["d_alt"] / 100, 2),
                round(t * (lg["d_alt"] - lg["d_ref"]) / 100, 2), lg["obs"]],
        [None, None, "#,##0.00", "#,##0.00", "#,##0.00", "0", "#,##0.00", "0", "#,##0.00", "#,##0.00", None],
        bold=True)
    ws.cell(r, 11).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(r, 11).font = Font(bold=False)
    r += 1
ws.cell(r + 1, 1, "Périmètre strict : la fourchette « environ 4 à 5 cl » du courriel ne vise que les trois plats flambés devant le client (Camembert, Assiette Franc-Comtoise, Assiette du Père Grégoire). "
                  "Les babas, coupes glacées et crêpes flambées relèvent d'une autre ligne du même courriel, qui fixe 4 cl sans fourchette.")
ws.cell(r + 2, 1, "La part « menus » est une espérance (nombre de menus vendus x probabilité que le plat soit choisi, lue en caisse), et non un comptage : les tickets n'enregistrent pas l'option choisie à l'intérieur d'un menu. "
                  "C'est la convention du service lui-même, celle de la colonne « nombre d'articles net » de ses annexes n° 7.")
ws.cell(r + 3, 1, "La fondue du Menu du Dahu est l'unique plat de ce menu : sa probabilité est de 1, et le ravelin des menus n'est donc pas affecté par cette incertitude.")
ws.cell(r + 4, 1, "Contrôle de fidélité des menus recalculés contre src/data/calculsBoissons/consoAlcoolMenus.json : "
                  + ("concordant sur les trois exercices et tous les alcools." if ctrl_ok else "DIVERGENT, voir la console."))

os.makedirs(OUT, exist_ok=True)
wb.save(OUT + "R1-alcool-cuisine-reperes.xlsx")
print("-> public/documents/pieces-reponse-1/R1-alcool-cuisine-reperes.xlsx")

# ==========================================================================
# recapitulatif console
# ==========================================================================
print("\n== 2/3. les douze reperes ==")
for rep in ["D", "E", "G", "Q"]:
    for ex in (1, 2, 3):
        lignes = [x for x in REPERES if x[0] == rep and x[2] == ex]
        sb, sn, sq = sum(x[5] for x in lignes), sum(x[7] for x in lignes), sum(x[9] for x in lignes)
        page, dose_ann, nb_c, vol_c = COURRIER[(rep, ex)]
        exact = sum(((x[5] * float(x[6].replace(" %", "").replace(",", ".")) / 100) if x[6] else x[5]) * x[8]
                    for x in lignes)
        print(f"  {rep} {lignes[0][1]:10} ex{ex}  brut {sb:5} net {sn:5} total {sq:6} (recalcul {exact:8.1f})"
              f"  |  courrier {page}: plats {nb_c:>12} volume {vol_c:6} dose annoncee {dose_ann}")

print("\n== 4. feuillete forestier a la carte ==")
for e in EXOS:
    print(f"  {e} : {feuillete[e]:8.2f} portions -> {feuillete[e]/100:6.3f} L de porto a 1 cl")
print(f"  TOTAL : {tot_f:.2f} portions -> {tot_f/100:.3f} L")
print(f"  proportion 170,50 / (170,50 + 186,00) = {prop_service:.2f} %  (chevre chaud lu exo1 : {chevre_ex1:.2f})")

print("\n== 5. bornes de la fourchette ==")
for lg in bornes_rows:
    t = sum(lg["tot"].values())
    print(f"  {lg['nom']:26} {t:9.2f} portions  {lg['d_ref']} cl -> {t*lg['d_ref']/100:7.2f} L  |  "
          f"{lg['d_alt']} cl -> {t*lg['d_alt']/100:7.2f} L  |  ecart {t*(lg['d_alt']-lg['d_ref'])/100:+7.2f} L")

print("\n== controle de fidelite des menus ==")
print("  ", "OK" if ctrl_ok else "DIVERGENT")
if not ctrl_ok:
    for x in ctrl_detail:
        print("   ", x)
